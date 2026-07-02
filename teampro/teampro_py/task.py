
import frappe
@frappe.whitelist()
def validate_et_vs_at(name,alloc):
    task = frappe.get_doc("Task", name)
    if task.service=="IT-SW":
        employee = frappe.get_value("Employee", {"user_id": alloc}, "name")

        if not employee:
            frappe.throw(f"Employee record for the current user ({frappe.session.user}) not found.")

        timesheet_data = frappe.db.sql("""
            SELECT SUM(td.hours) AS total_hours
            FROM `tabTimesheet Detail` td
            INNER JOIN `tabTimesheet` ti ON td.parent = ti.name
            WHERE td.task = %s AND ti.employee = %s
        """, (task.name, employee), as_dict=True)

        actual_time = timesheet_data[0].total_hours if timesheet_data and timesheet_data[0].total_hours else 0.0
        estimated_time = frappe.db.get_value(
        "Task",
        {"name": task.name, "status": "Working"},
        "expected_time"
        ) or 0.0

        if not isinstance(actual_time, (int, float)) or not isinstance(estimated_time, (int, float)):
            frappe.throw("Invalid time values detected.")

        if actual_time > estimated_time and not task.custom_et_vs_at_remark:
            frappe.msgprint(f"ET vs AT Remark field is mandatory when Actual Time {round(actual_time, 2)} exceeds Estimated Time hours {round(estimated_time,2)}.")
            return True
        else:
            return False
        
import frappe
import os
import requests
from frappe.utils import get_files_path
from pdf2image import convert_from_path

@frappe.whitelist()
def convert_pdf_to_images_jd(file_url):
    if not file_url:
        frappe.throw("No file URL provided")

    pdf_url = f"https://erp.teamproit.com{file_url}"
    
    frappe.logger().error(f"Fetching PDF from: {pdf_url}")

    output_dir = get_files_path("pdf_images", is_private=False)
    os.makedirs(output_dir, exist_ok=True)
    
    pdf_filename = os.path.basename(file_url).replace(' ', '_')
    pdf_path = os.path.join(output_dir, pdf_filename)

    try:
        response = requests.get(pdf_url, stream=True)
        response.raise_for_status()

        with open(pdf_path, "wb") as f:
            for chunk in response.iter_content(1024):
                f.write(chunk)

    except requests.exceptions.RequestException as e:
        frappe.throw(f"Failed to download PDF: {str(e)}")

    try:
        images = convert_from_path(pdf_path, dpi=150)
    except Exception as e:
        frappe.throw(f"PDF conversion error: {str(e)}")

    image_urls = []
    for i, img in enumerate(images):
        img_filename = f"{pdf_filename}_page_{i + 1}.png"
        img_path = os.path.join(output_dir, img_filename)
        img.save(img_path, "PNG")

        image_urls.append(f"/files/pdf_images/{img_filename}")

    return image_urls

@frappe.whitelist()
def candidate_details_html(task):
    candidates = frappe.db.get_all("Candidate",filters={"task":task},fields=["name","given_name","passport_number","mobile_number","pending_for"])
    if candidates:
        return candidates
    
from io import BytesIO
from openpyxl.styles import GradientFill, PatternFill
from openpyxl.styles import Font, Alignment, Border, Side
import openpyxl

@frappe.whitelist()
def task_candidate_download():
    filename = "Candidate Details"
    build_xlsx_response_test(filename)

def build_xlsx_response_test(filename):
    xlsx_file = make_xlsx_test(filename)
    frappe.response['filename'] = filename + '.xlsx'
    frappe.response['filecontent'] = xlsx_file.getvalue()
    frappe.response['type'] = 'binary'
 

def make_xlsx_test(data, sheet_name=None, wb=None, column_widths=None):
    args = frappe.local.form_dict
    if args.get("statuses"):
        args["statuses"] = frappe.parse_json(args["statuses"])
    column_widths = column_widths or []
    if wb is None:
        wb = openpyxl.Workbook()
    ws = wb.create_sheet(sheet_name, 0)
    
    column_widths = {
        
        "A": 5,  
        "B": 15,  
        "C": 15,  
        "D": 40,  
        "E": 15,  
        "F": 15,  
        "G": 10,  
        "H": 10,  
        "I":10 , 
        "J": 36,  
        "K": 18,  
        "L":18 ,  
        "M": 18,
        "N":15,
        "O":70   
    }

    for col, width in column_widths.items():
        ws.column_dimensions[col].width = width

    sub_headers = [
        "", "", "", "", "", "",
        "India", "Gulf", "Total",
        "",
        "Current", "Expected",
        "", "", ""
    ]

    # Define styles
    position_fill = PatternFill(start_color="0f1568", end_color="0f1568", fill_type="solid")
    position_font = Font(color="FFFFFF", bold=True)
    header_fill = PatternFill(start_color="00b1f0", end_color="00b1f0", fill_type="solid")
    header_font = Font(bold=True)
    black_border = Border(
        left=Side(border_style="thin", color="000000"),
        right=Side(border_style="thin", color="000000"),
        top=Side(border_style="thin", color="000000"),
        bottom=Side(border_style="thin", color="000000")
    )
    even_row_fill = PatternFill(
    start_color="e6f2f1",
    end_color="e6f2f1",
    fill_type="solid"
    )
    # Fetch candidate data grouped by positions
    position_candidates = get_data_grouped_by_position(args)
    
   

    for position, position_data in position_candidates.items():
        
        candidates = position_data["rows"]
        currency = position_data["currency"]

        headers = [
            "S.No", "CD ID", "PP Number", "Candidate Name", "Qualification",
            "Specialization", "Experience", "", "",
            "Current Employer",
            f"Salary ({currency})", "",
            "Current Location", "Notice Period", "Remarks"
        ]
        # Add position row
        position_row = ws.max_row + 1
        ws.merge_cells(start_row=position_row, start_column=1, end_row=position_row, end_column=15)
        cell = ws.cell(row=position_row, column=1)
        cell.value = f"{position}"
        cell.fill = position_fill
        cell.font = position_font
        cell.alignment = Alignment(horizontal="center", vertical="center")
        cell.border = black_border
        header_row_1 = ws.max_row + 1
        header_row_2 = header_row_1 + 1
        # First header row
        for col_num, header in enumerate(headers, start=1):
            cell = ws.cell(row=header_row_1, column=col_num)
            cell.value = header
            cell.fill = header_fill
            cell.font = header_font
            cell.alignment = Alignment(horizontal="center", vertical="center")
            cell.border = black_border
        # Second header row (sub headers)
        for col_num, header in enumerate(sub_headers, start=1):
            cell = ws.cell(row=header_row_2, column=col_num)
            cell.value = header
            cell.fill = header_fill
            cell.font = header_font
            cell.alignment = Alignment(horizontal="center", vertical="center")
            cell.border = black_border
            
        rowspan_columns = [1,2,3,4,5,6,10,13,14,15]
        for col in rowspan_columns:
            ws.merge_cells(
                start_row=header_row_1,
                start_column=col,
                end_row=header_row_2,
                end_column=col
            )
    
        ws.merge_cells(
        start_row=header_row_1,
        start_column=7,
        end_row=header_row_1,
        end_column=9
        )
        
        ws.merge_cells(
        start_row=header_row_1,
        start_column=11,
        end_row=header_row_1,
        end_column=12
        )

        row_num = ws.max_row + 1



        # Add details for the position
        for idx, candidate in enumerate(candidates, start=1):
            # ws.append(candidate)
            row_num = ws.max_row + 1
            is_even_row = idx % 2 == 0
            
            cell = ws.cell(row=row_num, column=1)
            cell.value = idx
            cell.alignment = Alignment(horizontal="center", vertical="center")
            cell.border = black_border
            if is_even_row:
                cell.fill = even_row_fill
            
            
            
            for col_num, value in enumerate(candidate, start=2):
                cell = ws.cell(row=row_num, column=col_num)
                cell.value = value
                cell.border = black_border  # Apply border to each cell
                
                if col_num in[1, 2,3,  5, 6,7,8,9,10,14]:
                    
                    cell.alignment = Alignment(horizontal="center", vertical="center")
                if col_num in[11,12]:
                    
                    cell.alignment = Alignment(horizontal="right", vertical="center")    
                if col_num == 15:
                    cell.alignment = Alignment(wrap_text=True, vertical="top")    
                    
                if is_even_row:
                    cell.fill = even_row_fill

        # Add an empty row for separation
        ws.append([])

    xlsx_file = BytesIO()
    wb.save(xlsx_file)
    xlsx_file.seek(0)
    return xlsx_file

def get_data_grouped_by_position(args):
    data = {}
    filters = {"task": args.task}

    if args.statuses:
        filters["pending_for"] = ["in", args.statuses]
    candidates = frappe.get_all(
        "Candidate",
        filters= filters,
        fields=["name", "passport_number", "given_name", "highest_degree","specialization","india_experience",
                "overseas_experience","total_experience",  "current_employer",
                "current_ctc", "expected_ctc", "location", "notice_period_months",
                "remarks_1", "position","currency_ctc"]
    )
    
    for candidate in candidates:
        position = candidate.get("position", "")
        currency=candidate.currency_ctc
        formatted_ctc = f"{currency} {candidate.current_ctc}" if candidate.current_ctc else "0"
        formatted_expected_ctc = f"{currency} {candidate.expected_ctc}" if candidate.expected_ctc else "0"
        if position not in data:
            data[position] = {
                "currency": currency,
                "rows": []
            }
        data[position]["rows"].append([
            candidate.name, candidate.passport_number, candidate.given_name,
            candidate.highest_degree,  candidate.specialization, candidate.india_experience,candidate.overseas_experience, 
            candidate.total_experience, candidate.current_employer,
            formatted_ctc, formatted_expected_ctc, candidate.location, 
            candidate.notice_period_months, candidate.remarks_1
        ])

    return data

import json
@frappe.whitelist()
def assign_task_to_users(tasks, users,rc):
    tasks = json.loads(tasks) 
    users = json.loads(users) 
    rc=json.loads(rc)
    for task in tasks:
        doc = frappe.get_doc("Task", task)
        doc.set("custom_assign_task", [])
        for user in users:
            new_user=user.get('assigned_to')
            doc.append("custom_assign_task", {"assigned_to": new_user})
        doc.custom_required_count=rc
        # doc.custom_production_date=p_date
        doc.save()
        frappe.db.commit()

@frappe.whitelist()
def task_age_calculation():
    tasks = frappe.db.get_all("Task",filters={"service":"IT-SW", "status":["not in",["Hold","Completed","Cancelled"]]},fields=["name","creation"])
    
    for task in tasks:
        age_days = (frappe.utils.now_datetime() - task.creation).days

        frappe.db.set_value("Task",task.name,"custom_age",age_days)

    frappe.db.commit()


@frappe.whitelist()
def update_candidate_tcount():
    tasks = frappe.get_all('Task',{'service':('in',('REC-I', 'REC-D')),"status":("in", ['Working', 'Open', 'Overdue', 'Pending Review'])},["name"])
    # tasks = frappe.get_all("Task",filters={"name": "TS20840"},fields=["name"])
    for task in tasks:
        # task = task.name
#         print(task)
        submit_spoc = frappe.db.count(
            'Candidate', {'task': task.name, 'pending_for': 'Submit(SPOC)'}) or 0
        submit_client = frappe.db.count(
            'Candidate', {'task': task.name, 'pending_for': 'Submitted(Client)'}) or 0
        psl = frappe.db.count('Candidate', {'task': task.name, 'pending_for': (
            'in', ('Client Offered', 'Proposed PSL'))}) or 0
        shortlisted = frappe.db.count(
            'Candidate', {'task': task.name, 'pending_for':'Shortlisted'}) or 0
        linedup = frappe.db.count(
            'Candidate', {'task': task.name, 'pending_for':('in', ('Linedup','Linedup Confirmed'))}) or 0
        interviewed = frappe.db.count(
            'Candidate', {'task': task.name, 'pending_for': 'Interviewed'}) or 0
        result_pending =frappe.db.count('Candidate',{'task':task.name,'pending_for':'Result Pending'}) or 0

        frappe.db.set_value('Task', task.name, 'psl', psl)
        frappe.db.set_value('Task', task.name, 'fp',(submit_spoc + interviewed + submit_client))
        frappe.db.set_value('Task',task.name,'custom_rp',result_pending)
        frappe.db.set_value('Task', task.name, 'sl', shortlisted)
        frappe.db.set_value('Task', task.name, 'custom_lp',linedup)

        task_status = frappe.db.get_value('Task', task.name, 'status')

        if task_status in ('Completed', 'Cancelled'):
            frappe.db.set_value('Task', task.name, 'sp', 0)
            frappe.db.set_value('Task', task.name, 'custom_lp', 0)
        else:
            vac = frappe.db.get_value('Task', task.name, 'vac')
            prop = frappe.db.get_value('Task', task.name, 'prop')
            pps = (vac - psl) * prop - (submit_spoc + submit_client+
                                        interviewed + shortlisted +linedup)
            frappe.db.set_value('Task', task.name, 'sp', pps)





@frappe.whitelist()
def update_task_positions_count_hourly():
    frappe.enqueue(
        update_candidate_tcount,
        queue="long",
        timeout=36000,
        is_async=True,
        now=False,
        job_name='Task Update',
        enqueue_after_commit=False,
    )



@frappe.whitelist()
def update_task_from_candidate1():
    job = frappe.db.exists('Scheduled Job Type','update_task_positions_count_hourly')
    if not job:
        task = frappe.new_doc("Scheduled Job Type")
        task.update({
            "method": 'teampro.teampro_py.task.update_task_positions_count_hourly',
            "frequency": 'Cron',
            "cron_format": '0 * * * *'
        })
        task.save(ignore_permissions=True)


@frappe.whitelist()
def update_sprint_task(task_id,dev_team=None,sprint=None,project=None):
    if not dev_team or not sprint:
        return
    sprint_doc = frappe.get_doc("Sprint", {"team":dev_team,"sprint_id":sprint})
    status=frappe.db.get_value('Task',task_id,'status')
    if not any(row.task == task_id for row in sprint_doc.sprint_task):
        sprint_doc.append("sprint_task", {
            "task": task_id,
            "project":project,
            "status":status
        })
        sprint_doc.save(ignore_permissions=True)

# import frappe

# def update_prd_sprint_task(doc, method):
#     if not doc.custom_production_date:
#         return
#     if not doc.custom_sprint or not doc.custom_dev_team:
#         return
#     if doc.kt_confirmed == 0:
#         return
#     sprint_doc = frappe.get_doc("Sprint", {"team":doc.custom_dev_team,"sprint_id":doc.custom_sprint})
#     if any(row.task == doc.name for row in sprint_doc.sprint_task):
#         return

#     sprint_doc.append("sprint_task", {
#         "task": doc.name,
#         "project":doc.project,
#         "status":doc.status,
#         "created_on":doc.creation
#     })

#     sprint_doc.save(ignore_permissions=True)



import frappe
from frappe.utils import flt

def update_prd_sprint_task(doc, method):
    if not doc.custom_production_date:
        return

    if not doc.custom_sprint or not doc.custom_dev_team:
        return

    if doc.kt_confirmed == 0:
        return

    sprint_doc = frappe.get_doc(
        "Sprint",
        {
            "team": doc.custom_dev_team,
            "sprint_id": doc.custom_sprint
        }
    )

    for row in sprint_doc.sprint_task:
        if row.task == doc.name:
            if flt(row.rt) != flt(doc.rt):
                row.rt = doc.rt
                sprint_doc.save(ignore_permissions=True)

            return

    sprint_doc.append("sprint_task", {
        "task": doc.name,
        "project": doc.project,
        "status": doc.status,
        "created_on": doc.creation,
        "rt": doc.rt
    })

    sprint_doc.save(ignore_permissions=True)



import frappe
from frappe.utils import flt
@frappe.whitelist()
def update_prd_sprint_task(doc, method):

    sprint_name = frappe.db.exists(
        "Sprint",
        {
            "team": doc.custom_dev_team,
            "sprint_id": doc.custom_sprint
        }
    )

    if not sprint_name:
        return

    # Directly update the child table row instead of loading and saving the entire Sprint doc.
    # This avoids triggering Sprint's validate hooks (update_sprint_hours, update_allocated_hrs,
    # update_sprint_status) which perform O(n) SQL queries per row.
    child_row = frappe.db.get_value(
        "Sprint Task",
        {"parent": sprint_name, "task": doc.name},
        ["name", "rt"],
        as_dict=True,
    )

    if child_row and flt(child_row.rt) != flt(doc.rt):
        frappe.db.set_value("Sprint Task", child_row.name, "rt", doc.rt)