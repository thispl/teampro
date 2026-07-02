import frappe
from frappe.utils.csvutils import read_csv_content
from frappe.utils import get_first_day, get_last_day, format_datetime, get_url_to_form
from frappe.utils import cint
from frappe.utils.data import date_diff, now_datetime, nowdate, today, add_days
import datetime
from frappe import _
from frappe import _
from frappe.model.document import Document
from frappe.utils import getdate, nowdate
from frappe import throw, msgprint
import frappe
from datetime import date
from frappe import throw, _
from frappe.utils import getdate, today
today = date.today()
from frappe.model.document import Document
import datetime
import frappe,erpnext
from frappe.utils import cint
from frappe.utils import validate_email_address
import json
from frappe.utils import date_diff, add_months,today,add_days,add_years,nowdate,flt
from frappe.model.mapper import get_mapped_doc
from frappe.utils.file_manager import get_file
from frappe.utils.csvutils import UnicodeWriter, read_csv_content
import datetime
from datetime import date,datetime,timedelta
import openpyxl
from openpyxl import Workbook
import openpyxl
import xlrd
import re
from frappe.utils import today
from openpyxl import Workbook
from openpyxl.styles import PatternFill, Font, Alignment
from openpyxl.utils import get_column_letter
from datetime import datetime
from io import BytesIO
from openpyxl.styles import Font, Alignment, Border, Side
from openpyxl import load_workbook
from openpyxl.utils import get_column_letter
from openpyxl.styles import GradientFill, PatternFill
import pandas as pd
from frappe.utils import formatdate
from frappe.utils import now
from erpnext.setup.utils import get_exchange_rate
from datetime import date
from six import BytesIO, string_types
from frappe.utils import time_diff
from frappe.utils.csvutils import read_csv_content
from erpnext.buying.doctype.purchase_order.purchase_order import update_status
from frappe.utils.file_manager import get_file
from urllib.parse import urlencode
from frappe.model.rename_doc import rename_doc
from datetime import datetime
from frappe.utils import today

@frappe.whitelist()
def set_to_clarification(docname):
    doc = frappe.get_doc("Project", docname)
    doc.status = "Clarification"
    doc.save(ignore_permissions=True)
    frappe.db.commit()
    return "Status Updated"

import json
import frappe

@frappe.whitelist()
def update_criteria_changes(name=None, criteria=None):
    if not name:
        frappe.throw("Project name is required.")

    if isinstance(criteria, str):
        criteria = json.loads(criteria)

    if not criteria:
        return "No criteria provided."
    tasks = frappe.get_all("Task", {"project": name}, ["name"])

    for task_info in tasks:
        task = frappe.get_doc("Task", task_info.name)

        if not task.custom_criteria_table:
            continue

        for row in criteria:
            task.append("custom_criteria_table", {
                "scheduling_criteria": row.get("scheduling_criteria"),
                "scheduling_parameter": row.get("scheduling_parameter")
            })

        task.save()

    return "Ok"

@frappe.whitelist()
def update_profile_submission_project(project):
    tasks=frappe.db.get_all("Task",{"project":project},["name","subject"])
    return tasks


@frappe.whitelist()
def update_task_fields(project):
    tasks = frappe.get_all('Task', {'project': project}, [
                           'name', 'project_manager', 'account_manager', 'service'])
    proj = frappe.get_doc('Project', {'name': project})
    for t in tasks:
        if t.project_manager != proj.project_manager:
            frappe.db.set_value(
                'Task', t.name, 'project_manager', proj.project_manager)
        if t.account_manager != proj.account_manager:
            frappe.db.set_value(
                'Task', t.name, 'account_manager', proj.account_manager)
        if t.service != proj.service:
            frappe.db.set_value('Task', t.name, 'service', proj.service)

@frappe.whitelist()
def add_project_id(project):
    projects = frappe.db.sql(
        """select project_id from `tabProject` where project_id is not null order by creation""", as_dict=True)
    project_id = projects[-1].project_id
    return 'PRO' + str(int(project_id.strip('PRO'))+1)

@frappe.whitelist()
def update_cdr_for_tl(name,alloc,sub,proj,priority):
    is_tl=frappe.db.get_value('Employee',{'user_id':alloc},['custom_is_tl'])
    is_sub_tl=frappe.db.get_value('Employee',{'user_id':alloc},['custom_is_sub_tl'])
    if is_tl==1 and is_sub_tl==0:
        return
        # team=frappe.db.get_value('Employee',{'user_id':alloc},['custom_dev_team'])
        # cdr_reviewer=frappe.db.get_value('Dev Team',{'name':team},['code_reviewer'])
        # if cdr_reviewer:
        #     cdr=frappe.db.get_value('Employee',{'user_id':cdr_reviewer},['name'])
    else:
        cdr=frappe.db.get_value('Employee',{'user_id':alloc},['custom_tl'])
    if cdr:
        if frappe.db.exists('Timesheet',{'employee':cdr,'start_date':today(),'status':'Submitted'}):
            timesheet=frappe.get_doc('Timesheet',{'employee':cdr,'start_date':today(),'status':'Submitted'})
            added_tasks=[]
            for t in timesheet.timesheet_summary:
                added_tasks.append(t.id)
            if name not in added_tasks:
                cb=frappe.db.get_value('Employee',{'user_id':alloc},['short_code'])
                data = {
                    "document": 'Task',
                    "id": name,
                    'project': proj,
                    'subject':sub,
                    'status': 'Code Review',
                    'cb': cb,
                    'priority': priority
                }
                timesheet.append("timesheet_summary", data)
                timesheet.save()
                frappe.db.commit()
        else:
            pass

@frappe.whitelist()
def reverse_revision(user):
    role = frappe.db.sql("""
        SELECT `tabUser`.name as name
        FROM `tabUser`
        LEFT JOIN `tabHas Role` ON `tabHas Role`.parent = `tabUser`.name
        WHERE `tabHas Role`.role = 'Customer Executive'
        AND `tabUser`.enabled = 1
        AND `tabUser`.name = %s
    """, (user,), as_dict=True)

    if role:
        return role[0].get('name')
    else:
        return None
    
@frappe.whitelist()
def reverse_revision_nc(name):
    reason = frappe.db.get_value("Energy Point And Non Conformity", {'reason_of_ep': ['like', '%'+name+'%']}, "name")

    if reason:
        nc = frappe.get_doc("Energy Point And Non Conformity", reason)
        nc.delete()

@frappe.whitelist()
def update_candidate_list(candidate,project,customer,task):
    can = json.loads(candidate)
    for c in can:
        cand = frappe.get_doc("Candidate",(c["candidate_id"]))
        cand.update({
            "pending_for": c["candidate_status"],
            "degree" : c.get("degree"),
            # "specialization" : c.get("specialization"),
            # "current_ctc" :c.get("current_ctc"),
            # "current_ctc" :c.get("current_ctc"),
            "indian_experience" : c.get("indian_experience"),
            "gulf_experience" : c.get("gulf_experience"),
            "currency_type" : c.get("currency_type"),
            "expected_ctc" : c.get("expected_ctc"),
            "passport_no" : c.get("passport_no"),
            "expiry_date" : c.get("expiry_date"),
            "ecr_status" : c.get("ecr_status"),
            "current_location" : c.get("current_location"),
            "mobile" : c.get("mobile"),
            "associate_name" : c.get("associate"),
            "user" : c.get("user"),
        })
        cand.db_update()
        frappe.db.commit()

@frappe.whitelist()
def load_candidates(task):
    candidates = frappe.get_all("Candidate", {"task": task}, ["*"], order_by="given_name asc")
    return candidates

import frappe

@frappe.whitelist()
def dev_team(user):
    emp = frappe.db.get_value(
        'Employee',
        {'status': 'Active', 'user_id': user},
        'custom_dev_team'
    )

    return emp if emp else None

@frappe.whitelist()
def create_cust(name):
    if not frappe.db.exists("Existing Customer",name):
        cust=frappe.new_doc("Existing Customer")
        cust.customer_id=name
        cust.save()
        frappe.db.commit()

@frappe.whitelist()
def update_dnc(name):
    sfp_name = frappe.db.get_all("Sales Follow Up", {"party_name":name}, "name")
    if sfp_name:
        for i in sfp_name:
            sfp_doc = frappe.get_doc("Sales Follow Up", i.name)
            sfp_doc.status = "Do Not Contact"
            sfp_doc.save()
            sfp_doc.reload()
            frappe.db.commit()

@frappe.whitelist()
def update_dnc_converted(name):
    sfp_name = frappe.db.get_all("Sales Follow Up", {"party_name":name}, "name")
    if sfp_name:
        for i in sfp_name:
            sfp_doc = frappe.get_doc("Sales Follow Up", i.name)
            sfp_doc.status = "Converted"
            sfp_doc.save()
            sfp_doc.reload()
            frappe.db.commit()

@frappe.whitelist()
def update_sla_details(name, service, sla_from_date, sla_to_date, sla_type, status,attach):
    try:
        customer = frappe.get_doc("Customer", name)
        if customer.custom_sla_details:
            for i in customer.custom_sla_details:
                customer.append("custom_sla_history", {
                    "service": i.service,
                    "sla_from_date": i.sla_from_date,
                    "sla_to_date": i.sla_to_date,
                    "sla_type": i.sla_type,
                    "status": i.status,
                    "attach":i.attach,
                })
            customer.set("custom_sla_details", [])
        customer.append("custom_sla_details", {
            "service": service,
            "sla_from_date": sla_from_date,
            "sla_to_date": sla_to_date,
            "sla_type": sla_type,
            "status": status,
            "attach":attach,
        })

        customer.save(ignore_permissions=True)
        frappe.db.commit()
        return {"status": "success", "message": "SLA details updated successfully"}
    except Exception as e:
        frappe.log_error(frappe.get_traceback(), "Update SLA Details")
        return {"status": "error", "message": str(e)}

#For downloading Timesheet status as excel
@frappe.whitelist()
def make_time_sheet():
    args = frappe.local.form_dict
    filename = args.name
    test = build_xlsx_response(filename)

def make_xlsx_timesheet(data, sheet_name=None, wb=None, column_widths=None):
    args = frappe.local.form_dict
    column_widths = column_widths or []
    if wb is None:
        wb = openpyxl.Workbook()
    ws = wb.create_sheet(sheet_name, 0)
    doc = frappe.get_doc("Timesheet",args.name)
    
    if doc:
        ws.append(["Document Name","Document ID","Subject","Project Name","CB","Status","ET","AT"])
        cb=''
        for i in doc.timesheet_summary:
            if i.document=="Task":
                cb = frappe.db.get_value("Task",{"name":i.id},["cb"])
                status =frappe.db.get_value("Task",{"name":i.id},["status"])
                type =frappe.db.get_value("Employee",{"name":doc.employee},["custom_dept_type"])
                if type == "OPS":
                    et = frappe.db.get_value("Task",{"name":i.id},["expected_time"])
                elif type == "CS":
                    et = frappe.db.get_value("Task",{"name":i.id},["pr_expected_time"])
                else:
                    et = 0
            else:
                status = frappe.db.get_value("Issue",{"name":i.id},["custom_issue_status"])
                et = 0
            ws.append([i.document,i.id,i.subject,i.project,cb,status,round(et,2),round(i.tu,2)])
    xlsx_file = BytesIO()
    wb.save(xlsx_file)
    return xlsx_file

def build_xlsx_response(filename):
    xlsx_file =make_xlsx_timesheet(filename)
    frappe.response['filename'] = filename + '.xlsx'
    frappe.response['filecontent'] = xlsx_file.getvalue()
    frappe.response['type'] = 'binary' 

@frappe.whitelist()
def return_detailed_ts(timesheet):
    tdoc=frappe.get_doc('Timesheet',timesheet)
    alloc=frappe.db.get_value('Employee',{'name':tdoc.employee},['user_id'])
    task_query ="""
        SELECT `tabTimesheet Detail`.task,
               `tabTimesheet Detail`.subject,
               `tabTimesheet Detail`.project,
               SUM(`tabTimesheet Detail`.hours) AS hours,
               `tabTimesheet Detail`.task_status,
               GROUP_CONCAT(`tabTimesheet Detail`.description SEPARATOR ', ') AS description
        FROM `tabTimesheet`
        LEFT JOIN `tabTimesheet Detail`
        ON `tabTimesheet`.name = `tabTimesheet Detail`.parent
        WHERE `tabTimesheet Detail`.task IS NOT NULL
          AND `tabTimesheet`.name = %s
        GROUP BY `tabTimesheet Detail`.task
    """
    task = frappe.db.sql(task_query, (timesheet,), as_dict=True)
    task_list = [t["task"] for t in task]

    emp = frappe.db.get_value('Timesheet', {'name': timesheet}, 'employee')
    start_date = frappe.db.get_value('Timesheet', {'name': timesheet}, 'start_date')
    alloc = frappe.db.get_value('Employee', {'employee': emp}, 'user_id')

    additional_task_query = """
        SELECT
            task.name,
            task.subject,
            task.project,
            task.status
        FROM `tabTask` task
        WHERE task.custom_allocated_to = %s
        AND task.custom_production_date = %s
        AND task.name NOT IN (
            SELECT `tabTimesheet Detail`.task
            FROM `tabTimesheet`
            LEFT JOIN `tabTimesheet Detail`
            ON `tabTimesheet`.name = `tabTimesheet Detail`.parent
            WHERE `tabTimesheet Detail`.task IS NOT NULL
                AND `tabTimesheet`.name = %s
        )
    """
    additional_tasks = frappe.db.sql(additional_task_query, (alloc, start_date, timesheet), as_dict=True)

    for t in additional_tasks:
        tstatus=frappe.db.get_all('Task',{'name':t['name']},['status'])
        if t['name'] not in task_list:
            task.append({
                'task': t['name'],
                'subject': t.get('subject', ''),
                'project': t.get('project', ''),
                'hours': 0,
                'task_status': t.get('status', ''),
                'description': ''
            })
    cdr_list = []
    is_tl=frappe.db.get_value('Employee', {'name': emp}, ['custom_is_tl'])
    sub_tl=frappe.db.get_value('Employee', {'name': emp}, ['custom_is_sub_tl'])
    # if is_tl==1:
    #     allocated_persons=[]
    #     if sub_tl==0:
    #         user=frappe.db.get_value('Employee', {'name': emp}, ['user_id'])
    #         team=frappe.db.get_value('Dev Team', {'code_reviewer': user}, ['name'])
    #         if team:
    #             team_tl=frappe.db.get_value('Employee', {'status':'Active','custom_is_tl': 1,'custom_is_sub_tl':0,'custom_dev_team':team,"department":"IT. Development - THIS"}, ['user_id'])
    #             if team_tl:
    #                 allocated_persons.append(team_tl)

    #     cdr_employees = frappe.db.get_all('Employee', {'custom_tl': emp}, ['user_id'])
    #     for cdr in cdr_employees:
    #         allocated_persons.append(cdr.user_id)
    #     cdr_tasks = frappe.db.get_all(
    #         'Task',
    #         {
    #             'custom_allocated_to': ('in',(allocated_persons)),
    #             'custom_new_pr_date': start_date,
    #             'status':'Pending Review'
    #         },
    #         ['name', 'subject', 'project', 'status']
    #     )
    #     for tsk in cdr_tasks:
    #         if tsk['name'] not in task_list:
    #             task.append({
    #                     'task': tsk['name'],
    #                     'subject': tsk.get('subject', ''),
    #                     'project': tsk.get('project', ''),
    #                     'hours': 0,
    #                     'task_status': tsk.get('status', ''),
    #                     'description': ''
    #                 })

    meeting_query = """
        SELECT `tabTimesheet Detail`.custom_meeting,
               `tabTimesheet Detail`.custom_subject_meeting,
               `tabTimesheet Detail`.project,
               SUM(`tabTimesheet Detail`.hours) AS hours
        FROM `tabTimesheet`
        LEFT JOIN `tabTimesheet Detail`
        ON `tabTimesheet`.name = `tabTimesheet Detail`.parent
        WHERE `tabTimesheet Detail`.custom_meeting IS NOT NULL
          AND `tabTimesheet`.name = %s
        GROUP BY `tabTimesheet Detail`.custom_meeting
    """
    meeting = frappe.db.sql(meeting_query, (timesheet,), as_dict=True)

    issue_query = """
        SELECT `tabTimesheet Detail`.custom_issue,
               `tabTimesheet Detail`.custom_subject_issue,
               `tabTimesheet Detail`.project,
               SUM(`tabTimesheet Detail`.hours) AS hours
        FROM `tabTimesheet`
        LEFT JOIN `tabTimesheet Detail`
        ON `tabTimesheet`.name = `tabTimesheet Detail`.parent
        WHERE `tabTimesheet Detail`.custom_issue IS NOT NULL
          AND `tabTimesheet`.name = %s
        GROUP BY `tabTimesheet Detail`.custom_issue
    """
    issue = frappe.db.sql(issue_query, (timesheet,), as_dict=True)
    return task, meeting, issue

@frappe.whitelist()
def get_task_name(task):
    task_name = frappe.db.sql(
        """select name from `tabTask` where status = 'Open' and name = '%s'  """ % (task), as_dict=1)[0]
    return task_name['name']

from frappe.utils import formatdate

import frappe

@frappe.whitelist()
def get_tasks_by_date_and_employee(employee, date, custom_dev_team):
    # has_sub_tl = frappe.db.exists(
    #     "Employee",
    #     {
    #         "custom_dev_team": custom_dev_team,
    #         "custom_is_sub_tl": 1
    #     }
    # )
    emp = frappe.db.get_value("Employee", employee, ["user_id"])
    tasks = frappe.get_all(
        "Task",
        filters={
            "custom_production_date": date,
            "custom_allocated_to": emp
        },
        fields=["name", "status", "rt", "project", "subject", "priority",'custom_dev_team','cb']
    )

    priority_order = {"Urgent": 1, "High": 2, "Medium": 3, "Low": 4}
    tasks.sort(key=lambda x: priority_order.get(x.get("priority") or "Low", 5))
    cdr_list = []
    is_tl = frappe.db.get_value('Employee', {'name': employee}, ['custom_is_tl'])
    is_sub_tl = frappe.db.get_value('Employee', {'name': employee}, ['custom_is_sub_tl'])
    allocated_persons = []
    # if not has_sub_tl:
    # if is_tl == 1 and is_sub_tl==0:
    #     user = frappe.db.get_value('Employee', {'name': employee}, ['user_id'])
    #     team = frappe.db.get_value('Dev Team', {'code_reviewer': user}, ['name'])
    #     if team:
    #         team_tl = frappe.db.get_value('Employee', {
    #             'status': 'Active',
    #             'custom_is_tl': 1,
    #             "custom_is_sub_tl":0,
    #             'custom_dev_team': team,
    #             'department':'IT. Development - THIS'
    #         }, ['user_id'])
    #         if team_tl:
    #             allocated_persons.append(team_tl)
    #     cdr_employees = frappe.db.get_all(
    #         'Employee',
    #         {'custom_tl': employee},
    #         ['user_id']
    #     )
    #     for cdr in cdr_employees:
    #         allocated_persons.append(cdr.user_id)
    #     cdr_tasks = frappe.db.get_all(
    #         'Task',
    #         filters={
    #             'custom_allocated_to': ('in', allocated_persons),
    #             'custom_new_pr_date': date,
    #             'status':'Pending Review'
    #         },
    #         fields=['name', 'subject', 'project', 'status', 'priority','custom_dev_team','cb']
    #     )

    #     frappe.errprint(cdr_tasks)


    #     cdr_tasks.sort(key=lambda x: priority_order.get(x.get("priority") or "Low", 5))
    #     for tsk in cdr_tasks:
    #         cdr_list.append({
    #             'task': tsk['name'],
    #             'subject': tsk.get('subject', ''),
    #             'project': tsk.get('project', ''),
    #             'hours': 0,
    #             'task_status': tsk.get('status', ''),
    #             'priority':tsk.get('priority', ''),
    #             'description': ''
    #         })
    # if is_tl == 1 and is_sub_tl==1:
        
    #     cdr_employees = frappe.db.get_all(
    #         'Employee',
    #         {'custom_tl': employee},
    #         ['user_id']
    #     )
    #     for cdr in cdr_employees:
    #         allocated_persons.append(cdr.user_id)
    #     cdr_tasks = frappe.db.get_all(
    #         'Task',
    #         filters={
    #             'custom_allocated_to': ('in', allocated_persons),
    #             'custom_new_pr_date': date,
    #             'status':'Pending Review'
    #         },
    #         fields=['name', 'subject', 'project', 'status', 'priority','custom_dev_team','cb']
    #     )

    #     frappe.errprint(cdr_tasks)


    #     cdr_tasks.sort(key=lambda x: priority_order.get(x.get("priority") or "Low", 5))
    #     for tsk in cdr_tasks:
    #         cdr_list.append({
    #             'task': tsk['name'],
    #             'subject': tsk.get('subject', ''),
    #             'project': tsk.get('project', ''),
    #             'hours': 0,
    #             'task_status': tsk.get('status', ''),
    #             'priority':tsk.get('priority', ''),
    #             'description': ''
    #         })





    if not tasks and not cdr_list:
        return "<div>No tasks found for the selected date and employee.</div>"

    html = """
    <style>
        .task-table, .task-table th, .task-table td {
            border: 1px solid black;
            border-collapse: collapse;
        }
    </style>
    <table class="table table-bordered">
        <thead>
            <tr style='background-color:#0f1568;color:white;text-align:center'>
                <th>Task ID</th>
                <th>Project</th>
                <th>Subject</th>
                <th>Status</th>
                <th>Priority</th>
                <th>RT</th>
                <th>Today RT</th>
            </tr>
        </thead>
        <tbody>
    """
    for task in tasks:
        today_rt = get_today_rt_from_child(task["name"], date,task["custom_dev_team"],task["cb"])

        html += f"""
            <tr>
                <td>{task.name}</td>
                <td>{task.project}</td>
                <td>{task.subject}</td>
                <td>{task.status}</td>
                <td>{task.priority}</td>
                <td style='text-align:right'>{task.rt or ''}</td>
                <td style='text-align:right'>{today_rt or task.rt}</td>
            </tr>
        """

    for cdr in cdr_list:
        html += f"""
            <tr style="background-color:#d3e8f2;">
                <td>{cdr['task']}</td>
                <td>{cdr['project']}</td>
                <td>{cdr['subject']}</td>
                <td>{cdr['task_status']}</td>
                <td>{cdr['priority']}</td>
                <td style='text-align:right'>0.5</td>
                <td style='text-align:right'>0.5</td>
            </tr>
        """


    html += "</tbody></table>"
    return html

def get_today_rt_from_child(task_id, date,team,cb=None):
    daily_monitors = frappe.get_all(
        "Daily Monitor",
        filters={"date": date,"dev_team":team},
        fields=["name"]
    )
    for dm in daily_monitors:
        child = frappe.db.get_value(
            "Allocated Tasks",
            {"parent": dm.name, "id": task_id,"cb":cb},
            "today_rt"
        )
        if child:
            return child
    return 0

@frappe.whitelist()
def update_visit_status_sfp(name,visit_status):
    sfp_name = frappe.db.get_all("Sales Follow Up", {"party_name": name}, "name")
    if sfp_name:
        for i in sfp_name:
            sfp_doc = frappe.get_doc("Sales Follow Up", i.name)
            sfp_doc.visit_status = visit_status
            sfp_doc.save()
            sfp_doc.reload()
            frappe.db.commit()


@frappe.whitelist()
def update_qualification_status(name,qualification_status):
    sfp_name = frappe.db.get_all("Sales Follow Up", {"party_name": name}, "name")
    if sfp_name:
        for i in sfp_name:
            sfp_doc = frappe.get_doc("Sales Follow Up", i.name)
            sfp_doc.qualification_status = qualification_status
            sfp_doc.save()
            sfp_doc.reload()
            frappe.db.commit()

@frappe.whitelist()
def update_company_name(name):
    organiation_name=frappe.db.get_value("Existing Leads",name,"lead_name")
    return organiation_name

# @frappe.whitelist()
# def update_territory_sfp(name,territory):
#     sfp_name = frappe.db.get_all("Sales Follow Up", {"party_name": name}, "name")
#     if sfp_name:
#         for i in sfp_name:
#             sfp_doc = frappe.get_doc("Sales Follow Up", i.name)
#             sfp_doc.sfp_territory = territory
#             sfp_doc.save()
#             sfp_doc.reload()
#             frappe.db.commit()

@frappe.whitelist()
def update_territory_sfp(name, territory):
    sfp_list = frappe.get_all("Sales Follow Up",{"party_name": name},["name"])
    for row in sfp_list:
        frappe.db.set_value("Sales Follow Up",row.name,"sfp_territory",territory)
    frappe.db.commit()
    return "Updated"

@frappe.whitelist()
def address(lead):
    if frappe.db.exists('Address',{'address_title':lead}):
        ad = frappe.get_doc('Address',{'address_title':lead})
        ad.address_title = lead
        ad.address_type = frappe.db.get_value('Lead',{'name':lead},['address_type'])
        ad.address_line1 = frappe.db.get_value('Lead',{'name':lead},['address_line_1'])
        ad.address_line2 = frappe.db.get_value('Lead',{'name':lead},['address_line_2'])
        ad.city = frappe.db.get_value('Lead',{'name':lead},['city_town'])
        ad.state = frappe.db.get_value('Lead',{'name':lead},['custom_state_list'])
        ad.country = frappe.db.get_value('Lead',{'name':lead},['country__'])
        ad.pincode = frappe.db.get_value('Lead',{'name':lead},['postal_code'])
        ad.save(ignore_permissions=True)
        frappe.db.commit()
    else:
        ad = frappe.new_doc('Address')
        ad.address_title = lead
        ad.address_type = frappe.db.get_value('Lead',{'name':lead},['address_type'])
        ad.address_line1 = frappe.db.get_value('Lead',{'name':lead},['address_line_1'])
        ad.address_line2 = frappe.db.get_value('Lead',{'name':lead},['address_line_2'])
        ad.city = frappe.db.get_value('Lead',{'name':lead},['city_town'])
        ad.state = frappe.db.get_value('Lead',{'name':lead},['custom_state_list'])
        ad.country = frappe.db.get_value('Lead',{'name':lead},['country__'])
        ad.pincode = frappe.db.get_value('Lead',{'name':lead},['postal_code'])
        ad.save(ignore_permissions=True)
        frappe.db.commit()
    frappe.msgprint('Address Updated Successfully')

@frappe.whitelist()
def update_market_segment_sfp(name,market_segment):
    sfp_name = frappe.db.get_all("Sales Follow Up", {"party_name": name}, "name")
    if sfp_name:
        for i in sfp_name:
            sfp_doc = frappe.get_doc("Sales Follow Up", i.name)
            sfp_doc.market_segment = market_segment
            sfp_doc.save()
            sfp_doc.reload()
            frappe.db.commit()