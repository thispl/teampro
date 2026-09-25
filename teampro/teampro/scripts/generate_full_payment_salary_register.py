# Copyright (c) 2026, TeamPRO and contributors
# For license information, please see license.txt

import os
from datetime import datetime

import frappe
from frappe.utils import flt
from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter


@frappe.whitelist()
def generate():
    """Generate a salary register Excel for June 2026 with full (30-day) payment values."""
    site_path = frappe.get_site_path()
    output_dir = os.path.join(site_path, "public", "files")
    if not os.path.exists(output_dir):
        os.makedirs(output_dir)

    filename = "Salary_Register_June_2026_Full_Payment_Days.xlsx"
    output_path = os.path.join(output_dir, filename)

    slips = frappe.get_all(
        "Salary Slip",
        filters={
            "start_date": [">=", "2026-06-01"],
            "end_date": ["<=", "2026-06-30"],
            "docstatus": ["!=", 2],
        },
        fields=[
            "name",
            "employee",
            "employee_name",
            "designation",
            "department",
            "start_date",
            "end_date",
            "total_working_days",
            "payment_days",
            "leave_without_pay",
            "absent_days",
            "gross_pay",
            "total_deduction",
            "net_pay",
        ],
        order_by="employee",
    )

    # Fetch DOJ from Employee in a single query
    employees = {s.employee for s in slips}
    doj_map = {}
    if employees:
        doj_map = {
            e.name: e.date_of_joining
            for e in frappe.get_all(
                "Employee",
                filters={"name": ["in", list(employees)]},
                fields=["name", "date_of_joining"],
            )
        }

    wb = Workbook()
    ws = wb.active
    ws.title = "June 2026 Salary Register"

    headers = [
        "S.No",
        "Salary Slip",
        "Employee Code",
        "Employee Name",
        "Department",
        "Designation",
        "DOJ",
        "Total Working Days",
        "Payment Days (Actual)",
        "Leave Without Pay",
        "Absent Days",
        "Gross Pay (Actual)",
        "Total Deduction (Actual)",
        "Net Pay (Actual)",
        "Full Payment Days",
        "Gross Pay (30 Days)",
        "Total Deduction (30 Days)",
        "Net Pay (30 Days)",
    ]
    ws.append(headers)

    header_fill = PatternFill(start_color="366092", end_color="366092", fill_type="solid")
    header_font = Font(color="FFFFFF", bold=True)
    thin_border = Border(
        left=Side(style="thin"),
        right=Side(style="thin"),
        top=Side(style="thin"),
        bottom=Side(style="thin"),
    )

    for col in range(1, len(headers) + 1):
        cell = ws.cell(row=1, column=col)
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = Alignment(horizontal="center", vertical="center")
        cell.border = thin_border

    totals = {
        "gross_actual": 0.0,
        "deduction_actual": 0.0,
        "net_actual": 0.0,
        "gross_full": 0.0,
        "deduction_full": 0.0,
        "net_full": 0.0,
    }

    for idx, slip in enumerate(slips, start=1):
        payment_days = flt(slip.payment_days)
        full_days = 30

        gross_actual = flt(slip.gross_pay)
        deduction_actual = flt(slip.total_deduction)
        net_actual = flt(slip.net_pay)

        if payment_days:
            gross_full = round(gross_actual / payment_days * full_days, 2)
            deduction_full = round(deduction_actual / payment_days * full_days, 2)
            net_full = round(net_actual / payment_days * full_days, 2)
        else:
            gross_full = deduction_full = net_full = 0.0

        totals["gross_actual"] += gross_actual
        totals["deduction_actual"] += deduction_actual
        totals["net_actual"] += net_actual
        totals["gross_full"] += gross_full
        totals["deduction_full"] += deduction_full
        totals["net_full"] += net_full

        gross_actual = round(gross_actual, 2)
        deduction_actual = round(deduction_actual, 2)
        net_actual = round(net_actual, 2)

        row = [
            idx,
            slip.name,
            slip.employee,
            slip.employee_name,
            slip.department,
            slip.designation,
            doj_map.get(slip.employee),
            flt(slip.total_working_days),
            payment_days,
            flt(slip.leave_without_pay),
            flt(slip.absent_days),
            gross_actual,
            deduction_actual,
            net_actual,
            full_days,
            gross_full,
            deduction_full,
            net_full,
        ]
        ws.append(row)

        for col in range(1, len(headers) + 1):
            cell = ws.cell(row=idx + 1, column=col)
            cell.border = thin_border
            if col >= 8:
                cell.alignment = Alignment(horizontal="right")
            else:
                cell.alignment = Alignment(horizontal="left")

    total_row = [
        "",
        "TOTAL",
        "",
        "",
        "",
        "",
        "",
        "",
        "",
        "",
        "",
        round(totals["gross_actual"], 2),
        round(totals["deduction_actual"], 2),
        round(totals["net_actual"], 2),
        "",
        round(totals["gross_full"], 2),
        round(totals["deduction_full"], 2),
        round(totals["net_full"], 2),
    ]
    ws.append(total_row)

    total_font = Font(bold=True)
    last_row = ws.max_row
    for col in range(1, len(headers) + 1):
        cell = ws.cell(row=last_row, column=col)
        cell.font = total_font
        cell.border = thin_border
        cell.alignment = Alignment(horizontal="right" if col >= 8 else "left")

    # Column widths
    widths = [6, 22, 16, 26, 22, 22, 12, 12, 12, 12, 12, 16, 18, 16, 14, 18, 22, 18]
    for i, width in enumerate(widths, start=1):
        ws.column_dimensions[get_column_letter(i)].width = width

    wb.save(output_path)
    public_url = frappe.utils.get_url() + "/files/" + filename
    frappe.msgprint(
        f"Salary register generated successfully.<br>File: <b>{filename}</b><br>Path: {output_path}<br>URL: {public_url}"
    )
    return {"file_path": output_path, "file_url": public_url, "record_count": len(slips)}
