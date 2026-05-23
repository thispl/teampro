
import frappe
import openpyxl
from six import BytesIO
from frappe import _
from frappe.utils.file_manager import get_file
from frappe.utils import get_url
import requests
import xlrd
from datetime import datetime, date


@frappe.whitelist()
def create_bank_transactions(docname, bank_account, company):
    file_url = frappe.db.get_value(
        "File",
        {"attached_to_doctype": "Bank Statement Import", "attached_to_name": docname},
        "file_url"
    )
    frappe.errprint(file_url)

    if not file_url:
        frappe.throw(_("No file attached for this document."))

    try:
        # in_memory_file = download_external_file_url(file_url)
        # wb = openpyxl.load_workbook(in_memory_file)
        # ws = wb.active
        # pps = list(ws.values)[7:]
        in_memory_file = download_external_file_url(file_url)
        
        # Determine file type
        if file_url.endswith(".xlsx"):
            wb = openpyxl.load_workbook(in_memory_file)
            ws = wb.active
            pps = list(ws.values)[7:]
        elif file_url.endswith(".xls"):
            book = xlrd.open_workbook(file_contents=in_memory_file.read())
            sheet = book.sheet_by_index(0)
            pps = [sheet.row_values(i) for i in range(7, sheet.nrows)]
        else:
            frappe.throw(_("Unsupported file format. Please upload .xlsx or .xls files."))
    except Exception as e:
        frappe.throw(f"Failed to read Excel file: {str(e)}")

    created_docs=[]

    for pp in pps:
        if not pp[1]: 
            continue

        bt = frappe.new_doc("Bank Transaction")
        bt.bank_account = bank_account
        bt.company = company
        # bt.date = pp[2]
        bt.description = pp[5]
        bt.reference_number = pp[4]
        excel_date = pp[2]
        date_value = None

        if isinstance(excel_date, (float, int)):  
            date_value = datetime(*xlrd.xldate_as_tuple(excel_date, 0)).date()
        else:
            try:
                date_value = datetime.strptime(str(excel_date).strip(), "%d/%m/%Y").date()
            except ValueError:
                try:
                    date_value = datetime.strptime(str(excel_date).strip(), "%Y-%m-%d").date()
                except ValueError:
                    frappe.throw(f"Unrecognized date format: {excel_date}")

        bt.date = date_value

        if str(pp[6]).strip().upper() == "CR":
            bt.deposit = pp[7]
            bt.withdrawal = 0
        elif str(pp[6]).strip().upper() == "DR":
            bt.withdrawal = pp[7]
            bt.deposit = 0
        else:
            bt.deposit = 0
            bt.withdrawal = 0

        bt.insert(ignore_permissions=True)
        created_docs.append(bt.name)  


    if created_docs:
        transactions_text = "\n".join(created_docs)
        frappe.db.set_value("Bank Statement Import", docname, "custom_transaction_list", transactions_text)
    
    return "success"


@frappe.whitelist()
def create_bank_transactions_new(docname, bank_account, company):
    file_url = frappe.db.get_value(
        "File",
        {"attached_to_doctype": "Bank Statement Import", "attached_to_name": docname},
        "file_url"
    )
    frappe.errprint(file_url)

    if not file_url:
        frappe.throw(_("No file attached for this document."))

    try:
        # in_memory_file = download_external_file_url(file_url)
        # wb = openpyxl.load_workbook(in_memory_file)
        # ws = wb.active
        # pps = list(ws.values)[22:]
        in_memory_file = download_external_file_url(file_url)
        
        # Determine file type
        if file_url.endswith(".xlsx"):
            wb = openpyxl.load_workbook(in_memory_file)
            ws = wb.active
            pps = list(ws.values)[22:]
        elif file_url.endswith(".xls"):
            book = xlrd.open_workbook(file_contents=in_memory_file.read())
            sheet = book.sheet_by_index(0)
            pps = [sheet.row_values(i) for i in range(22, sheet.nrows)]
        else:
            frappe.throw(_("Unsupported file format. Please upload .xlsx or .xls files."))
    except Exception as e:
        frappe.throw(f"Failed to read Excel file: {str(e)}")

    created_docs=[]

    for pp in pps:
        if not pp[1]: 
            continue

        bt = frappe.new_doc("Bank Transaction")
        bt.bank_account = bank_account
        bt.company = company
        # bt.date = pp[0]
        bt.description = pp[1]
        bt.reference_number = pp[2] or ''
        bt.deposit = pp[5]
        bt.withdrawal = pp[4]
        excel_date = pp[0]
        date_value = None

        if isinstance(excel_date, (float, int)):  
            date_value = datetime(*xlrd.xldate_as_tuple(excel_date, 0)).date()
        else:
            str_date = str(excel_date).strip()

            for fmt in ("%d/%m/%Y", "%Y-%m-%d", "%d-%m-%Y", "%m/%d/%Y", "%d/%m/%y"):
                try:
                    date_value = datetime.strptime(str_date, fmt).date()
                    break
                except ValueError:
                    continue
        if not date_value:
            frappe.log_error(f"Unrecognized date format: {excel_date}, defaulting to today.")
            date_value = date.today()

        bt.date = date_value

        bt.insert(ignore_permissions=True)
        created_docs.append(bt.name)  

    if created_docs:
        transactions_text = "\n".join(created_docs)
        frappe.db.set_value("Bank Statement Import", docname, "custom_transaction_list", transactions_text)
    
    return "success"

    

def download_external_file_url(file_url):
    if not file_url.startswith("http"):
        file_url = get_url() + file_url
    response = requests.get(file_url, stream=True)
    if response.status_code != 200:
        frappe.throw(_("Failed to download file from external storage."))

    return BytesIO(response.content)


import frappe
from io import BytesIO
from openpyxl import Workbook

@frappe.whitelist()
def download_created_items(docname):
    doc = frappe.get_doc("Bank Statement Import", docname)
    created_items = [i.strip() for i in (doc.custom_transaction_list or "").splitlines() if i.strip()]

    if not created_items:
        frappe.throw("No Transactions were created during this upload.")

    items = frappe.get_all(
        "Bank Transaction",
        filters={"name": ["in", created_items]},
        fields=["company","bank_account","date","description","reference_number","deposit","withdrawal"]
    )

    all_titles = set()
    item_parameters = {}


    sorted_titles = sorted(all_titles)  

    wb = Workbook()
    ws = wb.active
    ws.title = "Uploaded Transactions"

    headers = ["Company", "Bank Account", "Date", "Narration", "Reference Number","Deposit", "Withdrawal"] 
    ws.append(headers)

    for item in items:
        base_row = [
            item.company or "",
            item.bank_account or "",
            item.date or "",
            item.description or "",
            item.reference_number or "",
            item.deposit or "",
            item.withdrawal or ""
        ]

        ws.append(base_row)

    output = BytesIO()
    wb.save(output)
    output.seek(0)

    frappe.response['filename'] = "Uploaded Transactions.xlsx"
    frappe.response['filecontent'] = output.read()
    frappe.response['type'] = 'binary'