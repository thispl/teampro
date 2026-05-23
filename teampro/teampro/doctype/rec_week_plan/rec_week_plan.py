# Copyright (c) 2025, TeamPRO and contributors
# For license information, please see license.txt

import frappe
from frappe.model.document import Document
from frappe import _
from collections import defaultdict
import calendar


class RECWeekPlan(Document):
	# pass
    def validate(self):
        if self.has_value_changed("workflow_state"):
            self.append("status_transition",{
                "user":frappe.session.user,
                "status":self.workflow_state,
                "date":now_datetime()
            })
        # previous_doc=frappe.db.get_value("REC Week Plan",{"workflow_state":"Planned"},["name"])
        # if previous_doc:
        #     frappe.throw("Previous Week Plan is not submitted.Kindly submit it before create new Week Plan")
        processed_projects = set()
        indexes_to_remove = []

        for alloc_row in self.allocation:
            project = alloc_row.project
            if project and project not in processed_projects:
                for i, detail_row in enumerate(self.project_details):
                    if detail_row.project == project:
                        self.append("details", {
                            "project": project,
                            "project_name":detail_row.project_name,
                            "vac":detail_row.vac,
                            "sp":detail_row.sp,
                            "fp":detail_row.fp,
                            "sl":detail_row.sl,
                            "psl":detail_row.psl,
                            "status":detail_row.status,
                        })
                        indexes_to_remove.append(i)
                        processed_projects.add(project)
                        break

        for i in sorted(indexes_to_remove, reverse=True):
            self.project_details.pop(i)

        for i, row in enumerate(self.details, start=1):
            row.idx = i
        for i, row in enumerate(self.project_details, start=1):
            row.idx = i
        self.validate_duplicate_allocation()
        self.validate_week_plan()
        self.restrict_planned_transition_on_weekdays()
        # if frappe.db.exists("REC Week Plan",{"workflow_state":"Draft","name":("!=",self.name)}):
        #     frappe.throw("Not allowed to create two Week Plan in Draft state.")
    def validate_duplicate_allocation(self):
        seen = set()
        for row in self.allocation:
            key = (row.exe, row.task, row.date)
            if key in seen:
                frappe.throw(_(f"Duplicate allocation found for EXE: <b>{row.exe}</b>, Task: <b>{row.task}</b>, Date: <b>{row.date}</b>"))
            seen.add(key)

    def validate_week_plan(self):
        current_date = getdate(today())
        weekday = current_date.weekday() 
        days_to_prev_monday = weekday + 7
        prev_week_start = add_days(current_date, -days_to_prev_monday)
        prev_week_end = add_days(prev_week_start, 6)

        plans_prev_week = frappe.get_all(
            "REC Week Plan",
            filters=[
                ["workflow_state", "=", "Planned"],
                ["start_date", ">=", prev_week_start],
                ["start_date", "<=", prev_week_end],
            ],
            fields=["name", "start_date"]
        )
        for i in plans_prev_week:
            if i.name:
                frappe.throw("Previous Week Plan is not submitted.Kindly submit it before create new Week Plan")
                
   

    def validate_rc_vs_vac(self):
        rc_totals = defaultdict(int)

        for row in self.allocation:
            if not row.task:
                continue
            rc_totals[row.task] += row.rc or 0

        for task, total_rc in rc_totals.items():
            vac = frappe.db.get_value("REC Task Details", {"task": task}, "sp")
            if vac is not None and total_rc > vac:
                frappe.throw(_(f"Total RC ({total_rc}) for Task <b>{task}</b> exceeds SP target ({vac})."))

    def restrict_planned_transition_on_weekdays(self):
        if self.workflow_state != "Planned":
            return

        previous_state = frappe.db.get_value(self.doctype, self.name, "workflow_state")
        if previous_state != "Proposed":
            return

        user_roles = frappe.get_roles(frappe.session.user)
        allowed_roles = {"CEO", "MD"}
        if set(user_roles) & allowed_roles:
            return

        today = frappe.utils.getdate()
        day_name = calendar.day_name[today.weekday()]

        if day_name not in ["Saturday", "Sunday"]:
            frappe.throw(_("Only CEO or MD can move to <b>Planned</b> on weekdays."))


@frappe.whitelist()
def get_teampro_holidays(start_date, end_date):
    return frappe.get_all("Holiday", 
        filters={
            "parent": "TEAMPRO 2023",
            "holiday_date": ["between", [start_date, end_date]]
        },
        fields=["holiday_date"]
    )

@frappe.whitelist()
def update_week_plan_ac_by_today(candidate):
    from frappe.utils import today, getdate

    today_date = getdate(today())

    week_plan = frappe.db.get_value("REC Week Plan", {
        "start_date": ["<=", today_date],
        "end_date": [">=", today_date]
    }, "name")

    if week_plan:
        return updated_dsr_data(week_plan)

@frappe.whitelist()
def update_week_plan_ac_by_cron():
    from frappe.utils import today, getdate

    today_date = getdate(today())

    week_plan = frappe.db.get_value("REC Week Plan", {
        "start_date": ["<=", today_date],
        "end_date": [">=", today_date]
    }, "name")

    if week_plan:
        return updated_dsr_data(week_plan)


@frappe.whitelist()
def update_allocation_actual_count(docname):
    doc = frappe.get_doc("REC Week Plan", docname)

    for row in doc.allocation: 
        task = row.task
        exe = row.exe
        date = row.date

        if task and exe and date:
            count = frappe.db.sql("""
                SELECT COUNT(DISTINCT c.name) AS status_count
                FROM `tabCandidate` c
                INNER JOIN `tabCandidate status` cs ON c.name = cs.parent
                WHERE DATE(cs.sourced_date) = %s
                AND cs.sourced_by = %s
                AND c.candidate_created_by = %s
                AND cs.task = %s
                AND cs.status IN (%s)
            """, (date, exe, exe, task, "Submit(SPOC)"))

            row.ac = count[0][0] if count else 0

    doc.save()
    frappe.db.commit()
    return "AC updated successfully"


@frappe.whitelist()
def task_mail_notification_status():
    job = frappe.db.exists('Scheduled Job Type', 'create_nc_for_weekplan')
    if not job:
        task = frappe.new_doc("Scheduled Job Type")
        task.update({
            "method": 'teampro.teampro.doctype.rec_week_plan.rec_week_plan.create_nc_for_weekplan',
            "frequency": 'Cron',
            # Minute Hour Day Month DayOfWeek
            "cron_format": '50 23 * * 0'   # Every Sunday at 23:50
        })
        task.save(ignore_permissions=True)



from frappe.utils import nowdate, getdate
import frappe

@frappe.whitelist()
def run_week_monitor_rec_dpr():
    date=getdate(nowdate())
    today = add_days(date,1)

    week_plan = frappe.get_all(
        "REC Week Plan",
        filters={
            "start_date": ["<=", today],
            "end_date": [">=", today]
        },
        fields=["name", "start_date", "end_date"]
    )

    for dm in week_plan:
        try:
            # update_rec_dpr(
            #     name=dm.name,
            # )
            rec_dpr_mail(name=dm.name)
        except Exception as e:
            frappe.log_error(frappe.get_traceback(), f"REC Week Plan Monitor Auto Update Failed for {dm.name}")

from datetime import datetime
from frappe.utils.data import date_diff, now_datetime, nowdate, today, add_days
import frappe
@frappe.whitelist()
def run_week_monitor_dsr():
    today = nowdate()
    week_plan = frappe.get_all(
        "REC Week Plan",
        filters={
            "start_date": ["<=", today],
            "end_date": [">=", today]
        },
        fields=["name", "start_date", "end_date"]
    )

    for dm in week_plan:
        try:
            # update_rec_dsr(
            #     name=dm.name,
            # )
            rec_dsr_mail(name=dm.name)
        except Exception as e:
            frappe.log_error(frappe.get_traceback(), f"REC Week Plan Monitor Auto Update Failed for {dm.name}")



from datetime import datetime
from frappe.utils.data import date_diff, now_datetime, nowdate, today, add_days
import frappe

@frappe.whitelist()
def update_rec_dpr(name):
    date=nowdate()
    today = add_days(date,1)
    # today = nowdate()
    date_obj = datetime.strptime(today, '%Y-%m-%d')
    formatted_date = date_obj.strftime('%d/%m/%Y')

    emp = frappe.db.get_all("Employee", {
        'status': 'Active',
        'reports_to': 'TI00003',
        'user_id': ('not in', ['keerthana.k@groupteampro.com'])
    }, ['user_id'])

    recievers = [i.user_id for i in emp]
    recievers.append('sangeetha.a@groupteampro.com')

    parent_doc = frappe.get_doc("REC Week Plan", name)
    count = 1
    data = '<table border="1" width="100%" style="border-collapse: collapse;">'
    data += '''
    <tr style="background-color: #0f1568 ;text-align:center;color: white;"><b>
        <td style='width:5%'><b>SI NO</b></td>
        <td style='width:10%'><b>ID</b></td>
        <td style='width:15%'><b>Project</b></td>
        <td style='width:20%'><b>Subject</b></td>
        <td style='width:13%'><b>Allocated To</b></td>
        <td style='width:13%'><b>RC</b></td>
    </b></tr>
    '''

    task_group = {}

    for i in parent_doc.allocation:
        if str(i.date) == today:  
            data += '<tr><td>{}</td><td>{}</td><td>{}</td><td>{}</td><td>{}</td><td>{}</td></tr>'.format(
                count, i.task, i.project or '-', i.subject, i.exe, i.rc or '-'
            )
            count += 1

            if i.exe not in task_group:
                task_group[i.exe] = []
            task_group[i.exe].append(i)

    data += '</table>'

    frappe.sendmail(
        sender='sangeetha.a@groupteampro.com',
        recipients=['lokeshkumar.a@groupteampro.com','aruna.g@groupteampro.com','sangeetha.a@groupteampro.com'],
        # recipients='divya.p@groupteampro.com',
        subject=f'REC-DPR {formatted_date} - Reg',
        message=f"""
            <b>Dear Team,</b><br><br>
            Please find the below DPR for {formatted_date} for your kind reference and action.<br><br>
            {data}<br><br>
            Thanks & Regards,<br>TEAM ERP<br>
            <i>This email has been automatically generated. Please do not reply</i>
        """
    )

    for allocated_to, tasks in task_group.items():
        task_data = '<table border="1" width="100%" style="border-collapse: collapse;">'
        task_data += '''
        <tr style="background-color: #0f1568; text-align:center; color:white;">
            <td style='width:5%'><b>SI NO</b></td>
            <td style='width:15%'><b>Project</b></td>
            <td style='width:10%'><b>ID</b></td>
            <td style='width:20%'><b>Subject</b></td>
            <td style='width:20%'><b>Allocated To</b></td>
            <td style='width:13%'><b>RC</b></td>
        </tr>
        '''
        individual_count = 1
        for j in tasks:
            task_data += '<tr><td>{}</td><td>{}</td><td>{}</td><td>{}</td><td>{}</td><td>{}</td></tr>'.format(
                individual_count,j.project or '-', j.task, j.subject, j.exe,j.rc or '0'
            )
            individual_count += 1
        task_data += '</table>'

        if allocated_to:
            frappe.sendmail(
                sender='sangeetha.a@groupteampro.com',
                # recipients='divya.p@groupteampro.com',
                recipients=allocated_to,
                subject=f'Task DPR {formatted_date} - Reg',
                message=f"""
                    <b>Dear Team</b>,<br><br>
                    Please find your assigned task for {formatted_date}.<br><br>
                    {task_data}<br><br>
                    Thanks & Regards,<br>TEAM ERP<br>
                    <i>This email has been automatically generated. Please do not reply</i>
                """
            )

    frappe.msgprint("DPR mail has been successfully sent")

from datetime import datetime
from frappe.utils.data import getdate, nowdate
import frappe

@frappe.whitelist()
def update_rec_dsr(name):
    today = nowdate()
    date_obj = datetime.strptime(today, '%Y-%m-%d')
    formatted_date = date_obj.strftime('%d/%m/%Y')

    emp = frappe.db.get_all("Employee", {
        'status': 'Active',
        'reports_to': 'TI00003',
        'user_id': ('not in', ['keerthana.k@groupteampro.com'])
    }, ['user_id'])

    recievers = [i.user_id for i in emp]
    recievers.append('sangeetha.a@groupteampro.com')

    parent_doc = frappe.get_doc("REC Week Plan", name)
    count = 1
    data = '<table border="1" width="100%" style="border-collapse: collapse;">'
    data += '''
    <tr style="background-color: #0f1568 ;text-align:center;color: white;">
        <td style='width:5%'><b>SI NO</b></td>
        <td style='width:15%'><b>Project</b></td>
        <td style='width:10%'><b>ID</b></td>
        <td style='width:20%'><b>Subject</b></td>
        <td style='width:15%'><b>Allocated To</b></td>
        <td style='width:10%'><b>RC</b></td>
        <td style='width:10%'><b>AC</b></td>
    </tr>
    '''

    task_group = {}

    for i in parent_doc.allocation:
        if str(i.date) == today:
            data += '<tr><td>{}</td><td>{}</td><td>{}</td><td>{}</td><td>{}</td><td>{}</td><td style="text-align:right">{}</td></tr>'.format(
                count, i.project or '-', i.task, i.subject, i.exe, i.rc or '0', i.ac or '0'
            )
            count += 1

            if i.exe not in task_group:
                task_group[i.exe] = []
            task_group[i.exe].append(i)

    data += '</table>'

    frappe.sendmail(
        sender='sangeetha.a@groupteampro.com',
        recipients=['lokeshkumar.a@groupteampro.com','aruna.g@groupteampro.com','sangeetha.a@groupteampro.com'],
        # recipients='divya.p@groupteampro.com',
        subject=f'REC-DSR {formatted_date} - RC & AC Summary',
        message=f"""
            <b>Dear Team,</b><br><br>
            Please find the below DSR (RC & AC) for {formatted_date} for your kind reference.<br><br>
            {data}<br><br>
            Thanks & Regards,<br>TEAM ERP<br>
            <i>This email has been automatically generated. Please do not reply</i>
        """
    )

    for allocated_to, tasks in task_group.items():
        task_data = '<table border="1" width="100%" style="border-collapse: collapse;">'
        task_data += '''
        <tr style="background-color: #0f1568; text-align:center; color:white;">
            <td style='width:5%'><b>SI NO</b></td>
             <td style='width:15%'><b>Project</b></td>
            <td style='width:10%'><b>ID</b></td>
            <td style='width:25%'><b>Subject</b></td>
            <td style='width:10%'><b>RC</b></td>
            <td style='width:10%'><b>AC</b></td>
        </tr>
        '''
        individual_count = 1
        for j in tasks:
            task_data += '<tr><td>{}</td><td>{}</td><td>{}</td><td>{}</td><td style="text-align:right">{}</td><td style="text-align:right">{}</td></tr>'.format(
                individual_count,j.project or '-', j.task, j.subject, j.rc or '0', j.ac or '0'
            )
            individual_count += 1
        task_data += '</table>'

        if allocated_to:
            frappe.sendmail(
                sender='sangeetha.a@groupteampro.com',
                recipients=allocated_to,
                subject=f'Your Task DSR (RC & AC) - {formatted_date}',
                message=f"""
                    <b>Dear Team</b>,<br><br>
                    Please find your task update for {formatted_date}.<br><br>
                    {task_data}<br><br>
                    Thanks & Regards,<br>TEAM ERP<br>
                    <i>This email has been automatically generated. Please do not reply</i>
                """
            )

    frappe.msgprint("DSR mail has been successfully sent")

@frappe.whitelist()
def get_dpr_collapsible_html(name, start_date=None, end_date=None,executive=None):
    import frappe
    from frappe.utils import getdate
    parent_doc = frappe.get_doc("REC Week Plan", name)

    # Parse the dates
    start_date = getdate(start_date)
    end_date = getdate(end_date)

    exe_tasks = {}
    for row in parent_doc.allocation:
        row_date = getdate(row.date)
        if start_date <= row_date <= end_date and (not executive or row.exe == executive):
            exe_tasks.setdefault(row.exe, []).append({
                "task": row.task,
                "subject":row.subject,
                "rc": row.rc or 0
            })

    return exe_tasks


@frappe.whitelist()
def get_dsr_collapsible_html(name, start_date=None, end_date=None, executive=None, team_type=None):
    import frappe
    from frappe.utils import getdate, add_days, formatdate

    parent_doc = frappe.get_doc("REC Week Plan", name)

    start_date = getdate(start_date)
    end_date = getdate(end_date)

    date_list = []
    current = start_date
    while current <= end_date:
        date_list.append(current)
        current = add_days(current, 1)

    result = {}

    def init_task_entry(subject, project,project_name=None):
        return {
            "subject": subject,
            "project": project,
            "project_name": project_name,
            "dates": {str(d): {"rt": 0, "ac": 0} for d in date_list}
        }

    for row in parent_doc.allocation:
        row_date = getdate(row.date)
        if start_date <= row_date <= end_date and ((not executive or row.exe == executive) and ( not team_type or row.team == team_type) ):
        # if start_date <= row_date <= end_date and (not executive or row.exe == executive):
            # result.setdefault(row.exe, {})
            # task_map = result[row.exe]
            employee_name = frappe.db.get_value(
                "Employee",
                {"user_id": row.exe},
                "employee_name"
            ) or row.exe

            result.setdefault(employee_name, {})
            task_map = result[employee_name]
            if row.task not in task_map:
                project_name = frappe.db.get_value("Project", row.project, "project_name") if row.project else ""
                task_map[row.task] = init_task_entry(row.subject, row.project, project_name)
            task_map[row.task]["dates"][str(row_date)]["rt"] = row.rc or 0
            task_map[row.task]["dates"][str(row_date)]["ac"] += row.ac or 0

    for row in parent_doc.dsr:
        row_date = getdate(row.date)
        if start_date <= row_date <= end_date and ((not executive or row.exe == executive) and ( not team_type or row.team == team_type) ):
        # if start_date <= row_date <= end_date and (not executive or row.exe == executive):
            result.setdefault(row.exe, {})
            task_map = result[row.exe]
            if row.task not in task_map:
                project_name = frappe.db.get_value("Project", row.project, "project_name") if row.project else ""
                task_map[row.task] = init_task_entry(row.subject, row.project, project_name)

            task_map[row.task]["dates"][str(row_date)]["ac"] += row.ac or 0
            task_map[row.task]["dates"][str(row_date)]["rt"] += row.rc or 0

    return {
        "date_headers": [formatdate(d, "d MMM") for d in date_list],
        "raw_dates": [str(d) for d in date_list],
        "data": result
    }


@frappe.whitelist()
def get_dsr_collapsible_html_sams(name, start_date=None, end_date=None, sams=None):
    import frappe
    from frappe.utils import getdate, add_days, formatdate

    parent_doc = frappe.get_doc("REC Week Plan", name)

    start_date = getdate(start_date)
    end_date = getdate(end_date)

    date_list = []
    current = start_date
    while current <= end_date:
        date_list.append(current)
        current = add_days(current, 1)

    result = {}

    def init_task_entry(subject, project, project_name=None):
        return {
            "subject": subject,
            "project": project,
            "project_name": project_name,
            "dates": {str(d): {"rt": 0, "ac": 0} for d in date_list}
        }

    # Allocation Table
    for row in parent_doc.allocation_agent:
        row_date = getdate(row.date)
        if start_date <= row_date <= end_date and (not sams or row.sams == sams):
            sams_id = row.sams or "Unknown"
            result.setdefault(sams_id, {})
            task_map = result[sams_id]

            if row.task not in task_map:
                project_name = frappe.db.get_value("Project", row.project, "project_name") if row.project else ""
                task_map[row.task] = init_task_entry(row.subject, row.project, project_name)

            task_map[row.task]["dates"][str(row_date)]["rt"] += row.rc or 0
            task_map[row.task]["dates"][str(row_date)]["ac"] += row.ac or 0


    return {
        "date_headers": [formatdate(d, "d MMM") for d in date_list],
        "raw_dates": [str(d) for d in date_list],
        "data": result
    }





@frappe.whitelist()
def update_dsr_data(name):
    from frappe.utils import today
    import frappe

    doc = frappe.get_doc("REC Week Plan", name)
    task_ids = set()
    
    from frappe.utils import today, getdate

    date_today = getdate(today())

    for row in doc.allocation:
        if getdate(row.date) != date_today:
            continue 
        candidate_count = frappe.db.sql("""
            SELECT COUNT(DISTINCT c.name)
            FROM `tabCandidate` c
            INNER JOIN `tabCandidate status` cs ON c.name = cs.parent
            WHERE DATE(cs.sourced_date) = %s
              AND cs.sourced_by = %s
              AND c.candidate_created_by = %s
              AND cs.task = %s
              AND cs.status IN (%s)
        """, (row.date, row.exe, row.exe, row.task, "Submit(SPOC)"))

        row.ac = candidate_count[0][0] if candidate_count else 0
        task_ids.add((row.task, row.exe, str(row.date)))

    candidate_tasks = frappe.db.sql("""
        SELECT cs.task, c.candidate_created_by, DATE(cs.sourced_date)
        FROM `tabCandidate` c
        INNER JOIN `tabCandidate status` cs ON c.name = cs.parent
        WHERE DATE(cs.sourced_date) = %s
          AND cs.status IN (%s)
    """, (date_today, "Submit(SPOC)"))

    for task_id, owner, sourced_date in candidate_tasks:
        key = (task_id, owner, str(sourced_date))
        if key not in task_ids:
            add_count = frappe.db.sql("""
                SELECT COUNT(DISTINCT c.name)
                FROM `tabCandidate` c
                INNER JOIN `tabCandidate status` cs ON c.name = cs.parent
                WHERE DATE(cs.sourced_date) = %s
                  AND cs.sourced_by = %s
                  AND c.candidate_created_by = %s
                  AND cs.task = %s
                  AND cs.status IN (%s)
            """, (sourced_date, owner, owner, task_id, "Submit(SPOC)"))

            count = add_count[0][0] if add_count else 0
            if count > 0:
                existing_row = next(
                    (d for d in doc.allocation if d.task == task_id and d.exe == owner and str(d.date) == str(sourced_date)),
                    None
                )
                if existing_row:
                    existing_row.ac = count
                else:
                    doc.append("allocation", {
                        "task": task_id,
                        "exe": owner,
                        "ac": count,
                        "date": sourced_date
                    })
                task_ids.add(key)

    doc.save()
    frappe.db.commit()


import frappe
from frappe.utils import formatdate, escape_html
from frappe import _

@frappe.whitelist()
def get_project_html(project_name):
    if not project_name:
        return "<p>No project selected.</p>"

    project = frappe.db.get_value("Project", project_name, [
       'name','territory','sourcing_statu','tvac','tsp','tfp','tsl','tpsl','project_name'
    ], as_dict=True)

    if not project:
        return "<p>Project not found.</p>"

    html = f"""
<h4>Project Details</h4>
<table class="table table-bordered">
    <thead style="background-color: #007BFF; color: white;">
        <tr>
            <td>Project Name</td>
            <td>Sourcing Status</td>
            <td>Territory</td>
            <td>VAC</td>
            <td>#SP</td>
            <td>#FP</td>
            <td>#SL</td>
            <td>#PSL</td>
        </tr>
    </thead>
    <tbody>
        <tr>
            <td>{escape_html(project.project_name or "")}</td>
            <td>{escape_html(project.sourcing_statu or "N/A")}</td>
            <td>{escape_html(project.territory or "N/A")}</td>
            <td>{escape_html(project.tvac or "")}</td>
            <td>{escape_html(project.tsp or "")}</td>
            <td>{escape_html(project.tfp or "")}</td>
            <td>{escape_html(project.tsl or "")}</td>
            <td>{escape_html(project.tpsl or "")}</td>
        </tr>
    </tbody>
</table>
"""

    return html


from frappe.utils import today, add_days, get_first_day, get_last_day, getdate
@frappe.whitelist()
def create_nc_for_weekplan():
    current_date = getdate(today())
    weekday = current_date.weekday()
    days_to_next_monday = (7 - weekday) % 7 or 7
    next_week_start = add_days(current_date, days_to_next_monday)
    next_week_end = add_days(next_week_start, 6)
    plans_next_week = frappe.get_all("REC Week Plan",filters=[["workflow_state", "=", "Planned"],["start_date", ">=", next_week_start],["start_date", "<=", next_week_end],],fields=["name", "start_date"])
    all_plans = frappe.get_all("REC Week Plan",filters={"workflow_state": "Planned"},fields=["name", "start_date"])
    remaining = [p for p in all_plans if p not in plans_next_week]
    for i in remaining:
        print("Remaining:", i.name, i.start_date)
        frappe.db.set_value("REC Week Plan",i.name,"workflow_state","Completed")
        nc=frappe.new_doc("Energy Point And Non Conformity")
        nc.emp="TI00003"
        nc.action="Non Conformity(NC)"
        nc.class_proposed="Critical"
        nc.reason_of_ep=f"{i.name} Week Plan is not submitted"
        nc.save()
        nc.submit()


from datetime import datetime
from frappe.utils.data import date_diff, now_datetime, nowdate, today, add_days
import frappe

@frappe.whitelist()
def rec_dpr_mail(name):
    date=nowdate()
    today = add_days(date,1)
    date_obj = datetime.strptime(today, '%Y-%m-%d')
    formatted_date = date_obj.strftime('%d/%m/%Y')

    emp = frappe.db.get_all("Employee", {
        'status': 'Active',
        'reports_to': 'TI00003',
        'user_id': ('not in', ['keerthana.k@groupteampro.com'])
    }, ['user_id'])

    recievers = [i.user_id for i in emp]
    recievers.append('sangeetha.a@groupteampro.com')

    parent_doc = frappe.get_doc("REC Week Plan", name)
    count = 1
    s_count=1
    data = '<table border="1" width="100%" style="border-collapse: collapse;">'
    data += '''
    <tr style="background-color: #0f1568 ;text-align:center;color: white;"><b>
        <td style='width:5%'><b>SI NO</b></td>
        <td style='width:10%'><b>ID</b></td>
        <td style='width:15%'><b>Project</b></td>
        <td style='width:20%'><b>Subject</b></td>
        <td style='width:13%'><b>Allocated To</b></td>
        <td style='width:13%'><b>RC</b></td>
    </b></tr>
    '''
    table='<table border="1" width="100%" style="border-collapse: collapse;">'
    table += '''
    <tr style="background-color: #0f1568 ;text-align:center;color: white;"><b>
        <td style='width:5%'><b>SI NO</b></td>
        <td style='width:10%'><b>ID</b></td>
        <td style='width:15%'><b>Project</b></td>
        <td style='width:20%'><b>Subject</b></td>
        <td style='width:13%'><b>Agent</b></td>
         <td style='width:13%'><b>Allocated To</b></td>
        <td style='width:13%'><b>RC</b></td>
    </b></tr>
    '''
    task_group = {}
    task_groups={}
    for i in parent_doc.allocation:
        if str(i.date) == today:  
            data += '<tr><td>{}</td><td>{}</td><td>{}</td><td>{}</td><td>{}</td><td>{}</td></tr>'.format(
                count, i.task, i.project or '-', i.subject, i.exe, i.rc or '-'
            )
            count += 1

            if i.exe not in task_group:
                task_group[i.exe] = []
            task_group[i.exe].append(i)

    data += '</table>'
    for j in parent_doc.allocation_agent:
        if str(j.date) == today:  
            table += '<tr><td>{}</td><td>{}</td><td>{}</td><td>{}</td><td>{}</td><td>{}</td><td>{}</td></tr>'.format(
                s_count, j.task, j.project or '-', j.subject, j.sams,j.exe or '-', i.rc or '-'
            )
            s_count += 1

            if j.exe not in task_groups:
                task_groups[j.exe] = []
            task_groups[j.exe].append(j)

    table += '</table>'

    frappe.sendmail(
        sender='sangeetha.a@groupteampro.com',
        recipients=['lokeshkumar.a@groupteampro.com','aruna.g@groupteampro.com','sangeetha.a@groupteampro.com'],
        # recipients='bhuvaneswari.a@groupteampro.com',
        subject=f'REC-DPR {formatted_date} - Reg',
        message=f"""
            <b>Dear Team,</b><br><br>
            Please find the below DPR for {formatted_date} for your kind reference and action.<br><br>
            {data if data else ''}<br><br>
            {table if table else ''}<br><br>
            Thanks & Regards,<br>TEAM ERP<br>
            <i>This email has been automatically generated. Please do not reply</i>
        """
    )

    exe_emails = {}
    for exe in set(list(task_group.keys()) + list(task_groups.keys())):
        exe_email = frappe.db.get_value("Employee", {"user_id": exe}, "user_id")
        if exe_email:
            exe_emails[exe] = exe_email
    for allocated_to in set(list(task_group.keys()) + list(task_groups.keys())):
        exe_email = exe_emails.get(allocated_to)
        if not exe_email:
            continue
        task_data = ""
        if allocated_to in task_group:
            task_data += '<table border="1" width="100%" style="border-collapse: collapse;">'
            task_data += '''
            <tr style="background-color: #0f1568; text-align:center; color:white;">
                <td style='width:5%'><b>SI NO</b></td>
                <td style='width:15%'><b>Project</b></td>
                <td style='width:10%'><b>ID</b></td>
                <td style='width:20%'><b>Subject</b></td>
                <td style='width:20%'><b>Allocated To</b></td>
                <td style='width:13%'><b>RC</b></td>
            </tr>
            '''
            count = 1
            for j in task_group[allocated_to]:
                task_data += '<tr><td>{}</td><td>{}</td><td>{}</td><td>{}</td><td>{}</td><td>{}</td></tr>'.format(
                    count, j.project or '-', j.task, j.subject, j.exe, j.rc or '0'
                )
                count += 1
            task_data += '</table><br>'

        sams_data = ""
        if allocated_to in task_groups:
            sams_data += '<table border="1" width="100%" style="border-collapse: collapse;">'
            sams_data += '''
            <tr style="background-color: #0f1568; text-align:center; color:white;">
                <td style='width:5%'><b>SI NO</b></td>
                <td style='width:15%'><b>Project</b></td>
                <td style='width:10%'><b>ID</b></td>
                <td style='width:20%'><b>Subject</b></td>
                <td style='width:20%'><b>Agent</b></td>
                <td style='width:20%'><b>Allocated To</b></td>
                <td style='width:13%'><b>RC</b></td>
            </tr>
            '''
            s_count = 1
            for k in task_groups[allocated_to]:
                sams_data += '<tr><td>{}</td><td>{}</td><td>{}</td><td>{}</td><td>{}</td><td>{}</td><td>{}</td></tr>'.format(
                    s_count, k.project or '-', k.task, k.subject, k.sams or '', k.exe, k.rc or '0'
                )
                s_count += 1
            sams_data += '</table><br>'


        if allocated_to:

            frappe.sendmail(
                recipients=exe_email,
                # recipients='bhuvaneswari.a@groupteampro.com',
                subject=f'Task DPR {formatted_date} - Reg',
                message=f"""
                    <b>Dear Team</b>,<br><br>
                    Please find your assigned task for {formatted_date}.<br><br>
                    {task_data if task_data else ''}<br>
                    {sams_data if sams_data else ''}<br>
                    Thanks & Regards,<br>TEAM ERP<br>
                    <i>This email has been automatically generated. Please do not reply</i>
                """
            )

    frappe.msgprint("DPR mail has been successfully sent")

from datetime import datetime
from frappe.utils.data import date_diff, now_datetime, nowdate, today, add_days
import frappe

@frappe.whitelist()
def rec_dsr_mail(name):
    today = nowdate()
    date_obj = datetime.strptime(today, '%Y-%m-%d')
    formatted_date = date_obj.strftime('%d/%m/%Y')

    emp = frappe.db.get_all("Employee", {
        'status': 'Active',
        'reports_to': 'TI00003',
        'user_id': ('not in', ['keerthana.k@groupteampro.com'])
    }, ['user_id'])

    recievers = [i.user_id for i in emp]
    recievers.append('sangeetha.a@groupteampro.com')

    parent_doc = frappe.get_doc("REC Week Plan", name)
    count = 1
    s_count=1
    data = '<table border="1" width="100%" style="border-collapse: collapse;">'
    data += '''
    <tr style="background-color: #0f1568 ;text-align:center;color: white;"><b>
        <td style='width:5%'><b>SI NO</b></td>
        <td style='width:10%'><b>ID</b></td>
        <td style='width:15%'><b>Project</b></td>
        <td style='width:20%'><b>Subject</b></td>
        <td style='width:13%'><b>Allocated To</b></td>
        <td style='width:13%'><b>RC</b></td>
        <td style='width:13%'><b>AC</b></td>
    </b></tr>
    '''
    table='<table border="1" width="100%" style="border-collapse: collapse;">'
    table += '''
    <tr style="background-color: #0f1568 ;text-align:center;color: white;"><b>
        <td style='width:5%'><b>SI NO</b></td>
        <td style='width:10%'><b>ID</b></td>
        <td style='width:15%'><b>Project</b></td>
        <td style='width:20%'><b>Subject</b></td>
        <td style='width:13%'><b>Agent</b></td>
         <td style='width:13%'><b>Allocated To</b></td>
        <td style='width:13%'><b>RC</b></td>
        <td style='width:13%'><b>AC</b></td>
    </b></tr>
    '''
    task_group = {}
    task_groups={}
    for i in parent_doc.allocation:
        if str(i.date) == today:  
            data += '<tr><td>{}</td><td>{}</td><td>{}</td><td>{}</td><td>{}</td><td>{}</td><td>{}</td></tr>'.format(
                count, i.task, i.project or '-', i.subject, i.exe, i.rc or '-',i.ac or '-'
            )
            count += 1

            if i.exe not in task_group:
                task_group[i.exe] = []
            task_group[i.exe].append(i)

    data += '</table>'
    for j in parent_doc.allocation_agent:
        if str(j.date) == today:  
            table += '<tr><td>{}</td><td>{}</td><td>{}</td><td>{}</td><td>{}</td><td>{}</td><td>{}</td><td>{}</td></tr>'.format(
                s_count, j.task, j.project or '-', j.subject, j.sams,j.exe or '-', j.rc or '-',j.ac or '0'
            )
            s_count += 1

            if j.exe not in task_groups:
                task_groups[j.exe] = []
            task_groups[j.exe].append(j)

    table += '</table>'

    frappe.sendmail(
        sender='sangeetha.a@groupteampro.com',
        recipients=['lokeshkumar.a@groupteampro.com','aruna.g@groupteampro.com','sangeetha.a@groupteampro.com','bhuvaneswari.a@groupteampro.com'],
        # recipients='divya.p@groupteampro.com',
        subject=f'REC-DSR {formatted_date} - RC & AC Summary',
        message=f"""
            <b>Dear Team,</b><br><br>
            Please find the below DSR (RC & AC) for {formatted_date} for your kind reference.<br><br>
            {data if data else ""}<br>
            {table if table else ""}<br>
            Thanks & Regards,<br>TEAM ERP<br>
            <i>This email has been automatically generated. Please do not reply</i>
        """
    )
    exe_emails = {}
    for exe in set(list(task_group.keys()) + list(task_groups.keys())):
        exe_email = frappe.db.get_value("Employee", {"user_id": exe}, "user_id")
        if exe_email:
            exe_emails[exe] = exe_email
    for allocated_to in set(list(task_group.keys()) + list(task_groups.keys())):
        exe_email = exe_emails.get(allocated_to)
        if not exe_email:
            continue
        task_data = ""
        if allocated_to in task_group:
            task_data += '<table border="1" width="100%" style="border-collapse: collapse;">'
            task_data += '''
            <tr style="background-color: #0f1568; text-align:center; color:white;">
                <td style='width:5%'><b>SI NO</b></td>
                <td style='width:15%'><b>Project</b></td>
                <td style='width:10%'><b>ID</b></td>
                <td style='width:20%'><b>Subject</b></td>
                <td style='width:20%'><b>Allocated To</b></td>
                <td style='width:13%'><b>RC</b></td>
                <td style='width:13%'><b>AC</b></td>
            </tr>
            '''
            count = 1
            for j in task_group[allocated_to]:
                task_data += '<tr><td>{}</td><td>{}</td><td>{}</td><td>{}</td><td>{}</td><td>{}</td><td>{}</td></tr>'.format(
                    count, j.project or '-', j.task, j.subject, j.exe, j.rc or '0',j.ac or '0'
                )
                count += 1
            task_data += '</table><br>'

        sams_data = ""
        if allocated_to in task_groups:
            sams_data += '<table border="1" width="100%" style="border-collapse: collapse;">'
            sams_data += '''
            <tr style="background-color: #0f1568; text-align:center; color:white;">
                <td style='width:5%'><b>SI NO</b></td>
                <td style='width:15%'><b>Project</b></td>
                <td style='width:10%'><b>ID</b></td>
                <td style='width:20%'><b>Subject</b></td>
                <td style='width:20%'><b>Agent</b></td>
                <td style='width:20%'><b>Allocated To</b></td>
                <td style='width:13%'><b>RC</b></td>
                <td style='width:13%'><b>AC</b></td>
            </tr>
            '''
            s_count = 1
            for k in task_groups[allocated_to]:
                sams_data += '<tr><td>{}</td><td>{}</td><td>{}</td><td>{}</td><td>{}</td><td>{}</td><td>{}</td><td>{}</td></tr>'.format(
                    s_count, k.project or '-', k.task, k.subject, k.sams or '', k.exe, k.rc or '0',k.ac or '0'
                )
                s_count += 1
            sams_data += '</table><br>'

        if allocated_to:
            frappe.sendmail(
                sender='sangeetha.a@groupteampro.com',
                recipients=exe_emails,
                # recipients='divya.p@groupteampro.com',
                subject=f'Your Task DSR (RC & AC) - {formatted_date}',
                message=f"""
                    <b>Dear Team</b>,<br><br>
                    Please find your task update for {formatted_date}.<br><br>
                    {task_data if task_data else ""}<br>
                    {sams_data if sams_data else ""}<br>
                    Thanks & Regards,<br>TEAM ERP<br>
                    <i>This email has been automatically generated. Please do not reply</i>
                """
            )

    frappe.msgprint("DSR mail has been successfully sent")



@frappe.whitelist()
def updated_dsr_data(name):
    from frappe.utils import today
    import frappe

    doc = frappe.get_doc("REC Week Plan", name)
    task_ids = set()
    sams_task=set()
    from frappe.utils import today, getdate

    date_today = getdate(today())
    # todays=getdate(today())
    # date_today = add_days(todays, -1)
    for row in doc.allocation:
        if getdate(row.date) != date_today:
            continue 
        employee = frappe.db.get_value(
            "Employee",
            {"user_id": row.exe},
            ["name", "custom_dev_team"], 
            as_dict=True
        )

        
        candidate_count = frappe.db.sql("""
            SELECT COUNT(DISTINCT c.name)
            FROM `tabCandidate` c
            INNER JOIN `tabCandidate status` cs ON c.name = cs.parent
            WHERE DATE(cs.sourced_date) = %s
              AND c.candidate_created_by = %s
              AND cs.task = %s
              AND cs.status IN (%s)
        """, (row.date,row.exe, row.task, "Submit(SPOC)"))

        row.ac = candidate_count[0][0] if candidate_count else 0
        if employee:
            row.team = employee.custom_dev_team
        task_ids.add((row.task, row.exe, str(row.date)))
    for i in doc.allocation_agent:
        if getdate(i.date) != date_today:
            continue 
        employee = frappe.db.get_value(
                "Employee",
                {"user_id": i.exe},
                ["name", "custom_dev_team"], 
                as_dict=True
            )
        candidate_sams_count = frappe.db.sql("""
            SELECT COUNT(DISTINCT c.name)
            FROM `tabCandidate` c
            INNER JOIN `tabCandidate status` cs ON c.name = cs.parent
            WHERE DATE(cs.sourced_date) = %s
              AND c.candidate_created_by = %s
              AND cs.task = %s
              AND c.sa_agent=%s
              AND cs.status IN (%s)
        """, (i.date,i.exe, i.task,i.sams, "Submit(SPOC)"))

        i.ac = candidate_sams_count[0][0] if candidate_sams_count else 0
        if employee:
            i.team = employee.custom_dev_team
        sams_task.add((i.task, i.exe, str(i.date),i.sams))
    candidate_tasks = frappe.db.sql("""
        SELECT cs.task, c.candidate_created_by, DATE(cs.sourced_date)
        FROM `tabCandidate` c
        INNER JOIN `tabCandidate status` cs ON c.name = cs.parent
        WHERE DATE(cs.sourced_date) = %s
          AND cs.status IN (%s)
    """, (date_today, "Submit(SPOC)"))
    candidate_sams_tasks = frappe.db.sql("""
        SELECT cs.task, c.candidate_created_by, DATE(cs.sourced_date),c.sa_agent
        FROM `tabCandidate` c
        INNER JOIN `tabCandidate status` cs ON c.name = cs.parent
        WHERE DATE(cs.sourced_date) = %s
          AND cs.status IN (%s)
    """, (date_today, "Submit(SPOC)"))

    for task_id, owner, sourced_date in candidate_tasks:
        key = (task_id, owner, str(sourced_date))
        if key not in task_ids:
            employee = frappe.db.get_value(
                    "Employee",
                    {"user_id": owner},
                    ["name", "custom_dev_team"],  
                    as_dict=True
                )
            add_count = frappe.db.sql("""
                SELECT COUNT(DISTINCT c.name)
                FROM `tabCandidate` c
                INNER JOIN `tabCandidate status` cs ON c.name = cs.parent
                WHERE DATE(cs.sourced_date) = %s
                  AND c.candidate_created_by = %s
                  AND cs.task = %s
                  AND cs.status IN (%s)
            """, (sourced_date, owner, task_id, "Submit(SPOC)"))

            count = add_count[0][0] if add_count else 0
            if employee and employee.custom_dev_team:
                team = employee.custom_dev_team
            else:
                team = ""
            if count > 0:
                existing_row = next(
                    (d for d in doc.allocation if d.task == task_id and d.exe == owner and str(d.date) == str(sourced_date)),
                    None
                )
                if existing_row:
                    existing_row.ac = count
                else:
                    doc.append("allocation", {
                        "task": task_id,
                        "exe": owner,
                        "ac": count,
                        "date": sourced_date,
                        "team":team
                    })
                task_ids.add(key)
    # sams
    for task,agent, owners, sourced_date in candidate_sams_tasks:
        key = (task, owners, str(sourced_date))
        if key not in sams_task:
            employee = frappe.db.get_value(
                "Employee",
                {"user_id": owner},
                ["name", "employee_name", "custom_dev_team"],  
                as_dict=True
            )

            add_count = frappe.db.sql("""
                SELECT COUNT(DISTINCT c.name)
                FROM `tabCandidate` c
                INNER JOIN `tabCandidate status` cs ON c.name = cs.parent
                WHERE DATE(cs.sourced_date) = %s
                  AND c.candidate_created_by = %s
                  AND c.sa_agent=%s
                  AND cs.task = %s
                  AND cs.status IN (%s)
            """, (sourced_date, owners,agent, task, "Submit(SPOC)"))

            count = add_count[0][0] if add_count else 0
            team = employee.custom_dev_team
            if count > 0:
                existing_row = next(
                    (d for d in doc.allocation_agent if d.task == task and d.exe == owners and str(d.date) == str(sourced_date) and d.sams==agent),
                    None
                )
                if existing_row:
                    existing_row.ac = count
                else:
                    doc.append("allocation_agent", {
                        "task": task_id,
                        "sams":agent,
                        "exe": owners,
                        "ac": count,
                        "team":team,
                        "date": sourced_date
                    })
                sams_task.add(key)

    doc.save(ignore_permissions=True)
    frappe.db.commit()

@frappe.whitelist()
def get_rec_task_data():
    task_data = frappe.db.sql("""
        SELECT
            t.project,
            p.project_name,
            p.customer,
            t.territory,
            t.subject,
            t.priority,
            t.spoc,
            t.vac,
            t.sp,
            t.fp,
            t.sl,
            t.custom_lp as lp,
            t.psl,
            t.name as task,
            t.custom_task_sourcing_status as src_s,
            t.custom_sourcing_method as src,
            t.mode_of_interview as moi
        FROM `tabTask` t
        INNER JOIN `tabProject` p ON p.name = t.project
        WHERE t.status IN ('Open','Working','Pending Review','Overdue')
          AND t.service IN ('REC-I','REC-D')
          AND p.status NOT IN ('Draft','Enquiry','Hold','Completed','Cancelled')
        ORDER BY p.project_name ASC,t.custom_task_sourcing_status ASC
    """, as_dict=True)

    for row in task_data:
        if row.get("spoc"):
            row["spoc_short_code"] = frappe.db.get_value(
                "Employee",
                {"user_id": row.get("spoc")},
                "short_code"
            ) or ""

    return task_data

@frappe.whitelist()
def get_dev_team_logos(service="REC-I"):
    return frappe.db.sql("""
        SELECT
            d.logo,
            d.name,
            s.rec_capacity_per_day       
        FROM `tabDev Team` d
        INNER JOIN `tabTeam Services` s ON s.parent = d.name
        WHERE s.service = %s
          AND d.logo IS NOT NULL
        ORDER BY d.team_name
    """, service, as_dict=True)

@frappe.whitelist()
def update_task_det_rec(name):
    if not name:
        return

    frappe.enqueue(
        method="teampro.teampro.doctype.rec_week_plan.rec_week_plan.enqueue_update_task_det_rec",
        queue="long",  
        timeout=30000,
        name=name
    )

    return "Task update queued successfully"

def enqueue_update_task_det_rec(name):
    doc = frappe.get_doc("REC Week Plan", name)
    if not doc.rec_task_planner:
        return
    for row in doc.rec_task_planner:
        if not row.task_id:
            continue
        task = frappe.get_doc("Task", row.task_id)
        task.priority = row.priority
        task.custom_task_sourcing_status = row.src_s
        task.custom_sourcing_method = row.src
        task.save(ignore_permissions=True)
    frappe.db.commit()

@frappe.whitelist()
def update_task_src(tasks):
    """
    Update SRC and SRC_S in Task
    """
    import json

    if isinstance(tasks, str):
        tasks = json.loads(tasks)

    for row in tasks:
        if not row.get("task_id"):
            continue
        frappe.db.set_value("Task",row["task_id"],"custom_sourcing_method",row.get("src"))
        frappe.db.set_value("Task",row["task_id"],"custom_task_sourcing_status",row.get("src_s"))
        # frappe.db.set_value("Task",row["task_id"],{"custom_sourcing_method": row.get("src"),"custom_task_sourcing_status": row.get("src_s")},update_modified=False)

    frappe.db.commit()

import frappe
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side

@frappe.whitelist()
def download_master_excel(name):

    doc = frappe.get_doc("REC Week Plan", name)

    selected_customer = doc.customer
    selected_project = doc.project_collapsible
    selected_task = doc.task_collapsible
    selected_src = doc.src_s

    wb = Workbook()
    ws = wb.active
    ws.title = "Project Planner"

    # ===== STYLES =====
    header_fill = PatternFill("solid", fgColor="2B177A")
    client_fill = PatternFill("solid", fgColor="CFD8DC")
    project_fill = PatternFill("solid", fgColor="D1C4E9")
    row1_fill = PatternFill("solid", fgColor="F3F3F3")
    row2_fill = PatternFill("solid", fgColor="E7D2BF")

    white_font = Font(color="FFFFFF", bold=True)
    bold_font = Font(bold=True)

    center = Alignment(horizontal="center", vertical="center")
    left = Alignment(horizontal="left", vertical="center")

    border = Border(
        left=Side(style="thin"),
        right=Side(style="thin"),
        top=Side(style="thin"),
        bottom=Side(style="thin")
    )

    # ===== COLUMN WIDTH =====
    widths = [6, 35, 12, 10, 8, 8, 8, 8, 8, 8, 8]
    for i, w in enumerate(widths, 1):
        ws.column_dimensions[chr(64+i)].width = w

    row_no = 1

    # ===== HEADER =====
    headers = ["S.No","Customer / Project","Territory","SRC_S","VAC","FP","SP","SL","LP","PSL","CC"]

    for col, h in enumerate(headers, 1):
        cell = ws.cell(row=row_no, column=col, value=h)
        cell.fill = header_fill
        cell.font = white_font
        cell.alignment = center
        cell.border = border

    row_no += 1

    # ===== GROUPING =====
    grouped = {}

    for r in doc.rec_task_planner:

        if selected_customer and r.client != selected_customer:
            continue
        if selected_project and r.project_name != selected_project:
            continue
        if selected_task and r.task != selected_task:
            continue
        if selected_src and r.src_s != selected_src:
            continue

        client = r.client or "No Client"
        project = r.project_name or "No Project"

        grouped.setdefault(client, {}).setdefault(project, []).append(r)

    client_idx = 0

    for client, projects in grouped.items():
        client_idx += 1

        # ===== CLIENT TOTAL =====
        totals = dict(vac=0, fp=0, sp=0, sl=0, lp=0, psl=0)

        for p in projects.values():
            for r in p:
                totals["vac"] += r.vac or 0
                totals["fp"] += r.fp or 0
                totals["sp"] += r.sp or 0
                totals["sl"] += r.sl or 0
                totals["lp"] += r.lp or 0
                totals["psl"] += r.psl or 0

        row = [
            client_idx, f"{client}", "", "",
            totals["vac"], totals["fp"], totals["sp"],
            totals["sl"], totals["lp"], totals["psl"], ""
        ]

        for col, val in enumerate(row, 1):
            cell = ws.cell(row=row_no, column=col, value=val)
            cell.fill = client_fill
            cell.font = bold_font
            cell.alignment = left if col == 2 else center
            cell.border = border

        row_no += 1

        proj_idx = 0

        for project, tasks in projects.items():
            proj_idx += 1

            # ===== PROJECT TOTAL =====
            totals = dict(vac=0, fp=0, sp=0, sl=0, lp=0, psl=0)

            for r in tasks:
                totals["vac"] += r.vac or 0
                totals["fp"] += r.fp or 0
                totals["sp"] += r.sp or 0
                totals["sl"] += r.sl or 0
                totals["lp"] += r.lp or 0
                totals["psl"] += r.psl or 0

            row = [
                proj_idx, f"   ↳ {project}", "", "",
                totals["vac"], totals["fp"], totals["sp"],
                totals["sl"], totals["lp"], totals["psl"], ""
            ]

            for col, val in enumerate(row, 1):
                cell = ws.cell(row=row_no, column=col, value=val)
                cell.fill = project_fill
                cell.font = bold_font
                cell.alignment = left if col == 2 else center
                cell.border = border

            row_no += 1

            # ===== TASK HEADER =====
            task_headers = ["S.No","Task","Territory","SRC_S","VAC","FP","SP","SL","LP","PSL","CC"]

            for col, h in enumerate(task_headers, 1):
                cell = ws.cell(row=row_no, column=col, value=h)
                cell.fill = header_fill
                cell.font = white_font
                cell.alignment = center
                cell.border = border

            row_no += 1

            # ===== TASK ROWS =====
            for i, r in enumerate(tasks, 1):

                fill = row1_fill if i % 2 == 0 else row2_fill

                row = [
                    i, r.task, r.territory, r.src_s,
                    r.vac, r.fp, r.sp, r.sl, r.lp, r.psl, r.cc
                ]

                for col, val in enumerate(row, 1):
                    cell = ws.cell(row=row_no, column=col, value=val)
                    cell.fill = fill
                    cell.alignment = left if col == 2 else center
                    cell.border = border

                row_no += 1

    # ===== SAVE FILE =====
    file_path = f"/tmp/{name}_Project_Planner.xlsx"
    wb.save(file_path)

    with open(file_path, "rb") as f:
        frappe.response.filename = f"{name}_Project_Planner.xlsx"
        frappe.response.filecontent = f.read()
        frappe.response.type = "download"




import frappe
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from io import BytesIO
from teampro.teampro.doctype.rec_week_plan.rec_week_plan import get_dsr_collapsible_html
from openpyxl.utils import get_column_letter

@frappe.whitelist()
def download_dsr_excel(name, start_date=None, end_date=None, executive=None, team_type=None):
    dsr_data = get_dsr_collapsible_html(name, start_date, end_date, executive, team_type)

    wb = Workbook()
    ws = wb.active
    ws.title = "DSR"

    header_fill = PatternFill(start_color='001F5B', end_color='001F5B', fill_type='solid')  # Dark blue
    subheader_rc_fill = PatternFill(start_color='BBDEFB', end_color='BBDEFB', fill_type='solid')  # Light blue
    subheader_ac_fill = PatternFill(start_color='E1F5FE', end_color='E1F5FE', fill_type='solid')  # Lighter blue
    bold_font = Font(bold=True)
    white_bold_font = Font(bold=True, color="FFFFFF")  
    center_align = Alignment(horizontal='center', vertical='center')
    left_align = Alignment(horizontal='left', vertical='center')

    thin_border = Border(left=Side(style='thin'),
                         right=Side(style='thin'),
                         top=Side(style='thin'),
                         bottom=Side(style='thin'))

    dates = dsr_data['raw_dates']
    executives_data = dsr_data['data']

    ws.cell(row=1, column=1, value='S.No').font = white_bold_font
    ws.cell(row=1, column=1).fill = header_fill
    ws.cell(row=1, column=1).alignment = center_align
    ws.merge_cells(start_row=1, start_column=1, end_row=2, end_column=1)

    ws.cell(row=1, column=2, value='EXE').font = white_bold_font
    ws.cell(row=1, column=2).fill = header_fill
    ws.cell(row=1, column=2).alignment = center_align
    ws.merge_cells(start_row=1, start_column=2, end_row=2, end_column=2)

    ws.cell(row=1, column=3, value='Task').font = white_bold_font
    ws.cell(row=1, column=3).fill = header_fill
    ws.cell(row=1, column=3).alignment = center_align
    ws.merge_cells(start_row=1, start_column=3, end_row=2, end_column=3)

    ws.cell(row=1, column=4, value='Subject').font = white_bold_font
    ws.cell(row=1, column=4).fill = header_fill
    ws.cell(row=1, column=4).alignment = center_align
    ws.merge_cells(start_row=1, start_column=4, end_row=2, end_column=4)

    col = 5
    ws.cell(row=1, column=col, value='Total RC').font = white_bold_font
    ws.cell(row=1, column=col).fill = header_fill
    ws.cell(row=1, column=col).alignment = center_align
    ws.merge_cells(start_row=1, start_column=col, end_row=2, end_column=col)
    col += 1

    ws.cell(row=1, column=col, value='Total AC').font = white_bold_font
    ws.cell(row=1, column=col).fill = header_fill
    ws.cell(row=1, column=col).alignment = center_align
    ws.merge_cells(start_row=1, start_column=col, end_row=2, end_column=col)
    col += 1


    for date in dates:
        if isinstance(date, str):
            date_obj = datetime.strptime(date, "%Y-%m-%d")
        else:
            date_obj = date
        date_str = date_obj.strftime("%d %b")
        ws.cell(row=1, column=col, value=date_str).font = white_bold_font
        ws.cell(row=1, column=col).fill = header_fill
        ws.cell(row=1, column=col).alignment = center_align
        ws.merge_cells(start_row=1, start_column=col, end_row=1, end_column=col+1)

        ws.cell(row=2, column=col, value='RC').fill = subheader_rc_fill
        ws.cell(row=2, column=col).font = bold_font
        ws.cell(row=2, column=col).alignment = center_align
        ws.cell(row=2, column=col+1, value='AC').fill = subheader_ac_fill
        ws.cell(row=2, column=col+1).alignment = center_align
        ws.cell(row=2, column=col+1).font = bold_font

        col += 2

    current_row = 3
    s_no = 1

    overall_total_rc = 0
    overall_total_ac = 0
    daywise_rc_totals = {d: 0 for d in dates}
    daywise_ac_totals = {d: 0 for d in dates}

    for exe_email, tasks in executives_data.items():
        num_tasks = len(tasks)
        exe_total_rc = sum(sum(task['dates'][d]['rt'] for d in dates) for task in tasks.values())
        exe_total_ac = sum(sum(task['dates'][d]['ac'] for d in dates) for task in tasks.values())
        ws.cell(row=current_row, column=1, value='').alignment = center_align
        ws.cell(row=current_row, column=2, value=exe_email).font = bold_font
        ws.cell(row=current_row, column=2).alignment = left_align


        ws.cell(row=current_row, column=3, value=num_tasks).font = bold_font
        ws.cell(row=current_row, column=3).alignment = center_align

        ws.cell(row=current_row, column=5, value=exe_total_rc).font = bold_font
        ws.cell(row=current_row, column=5).alignment = center_align
        ws.cell(row=current_row, column=6, value=exe_total_ac).font = bold_font
        ws.cell(row=current_row, column=6).alignment = center_align

        col_offset = 7
        for date in dates:
            rc_sum = sum(task['dates'][date]['rt'] for task in tasks.values())
            ac_sum = sum(task['dates'][date]['ac'] for task in tasks.values())
            ws.cell(row=current_row, column=col_offset, value=rc_sum).font = bold_font
            ws.cell(row=current_row, column=col_offset).alignment = center_align
            ws.cell(row=current_row, column=col_offset+1, value=ac_sum).font = bold_font
            ws.cell(row=current_row, column=col_offset+1).alignment = center_align
            col_offset += 2

            daywise_rc_totals[date] += rc_sum
            daywise_ac_totals[date] += ac_sum

        overall_total_rc += exe_total_rc
        overall_total_ac += exe_total_ac

        current_row += 1
        

        for task_code, task_info in tasks.items():
            ws.cell(row=current_row, column=1, value=s_no).font = bold_font
            ws.cell(row=current_row, column=1).alignment = center_align
            ws.cell(row=current_row, column=2, value=f"{task_info['project_name'] or task_info['project']}").alignment = left_align
            ws.cell(row=current_row, column=3, value=task_code).alignment = left_align
            ws.cell(row=current_row, column=4, value=task_info['subject']).alignment = left_align

            ws.cell(row=current_row, column=5, value=sum(task_info['dates'][d]['rt'] for d in dates)).alignment = center_align
            ws.cell(row=current_row, column=6, value=sum(task_info['dates'][d]['ac'] for d in dates)).alignment = center_align

            col_offset = 7
            for date in dates:
                ws.cell(row=current_row, column=col_offset, value=task_info['dates'][date]['rt']).fill = subheader_rc_fill
                ws.cell(row=current_row, column=col_offset).alignment = center_align
                ws.cell(row=current_row, column=col_offset+1, value=task_info['dates'][date]['ac']).fill = subheader_ac_fill
                ws.cell(row=current_row, column=col_offset+1).alignment = center_align
                col_offset += 2

            current_row += 1
            s_no+=1

    ws.cell(row=current_row, column=1, value='Overall Totals').font = bold_font
    ws.merge_cells(start_row=current_row, start_column=1, end_row=current_row, end_column=4)
    ws.cell(row=current_row, column=1).alignment = center_align
    ws.cell(row=current_row, column=5, value=overall_total_rc).font = bold_font
    ws.cell(row=current_row, column=5).alignment = center_align
    ws.cell(row=current_row, column=6, value=overall_total_ac).font = bold_font
    ws.cell(row=current_row, column=6).alignment = center_align

    col_offset = 7
    for date in dates:
        ws.cell(row=current_row, column=col_offset, value=daywise_rc_totals[date]).font = bold_font
        ws.cell(row=current_row, column=col_offset).alignment = center_align
        ws.cell(row=current_row, column=col_offset+1, value=daywise_ac_totals[date]).font = bold_font
        ws.cell(row=current_row, column=col_offset+1).alignment = center_align
        col_offset += 2

    max_col = col_offset - 1
    for row in ws.iter_rows(min_row=1, max_row=current_row, min_col=1, max_col=max_col):
        for cell in row:
            cell.border = thin_border

    column_widths = {
        2: 35,  
        4: 35, 
    }

    for col, width in column_widths.items():
        ws.column_dimensions[get_column_letter(col)].width = width

    output = BytesIO()
    wb.save(output)
    output.seek(0)

    frappe.local.response.filename = "DSR_Report.xlsx"
    frappe.local.response.filecontent = output.getvalue()
    frappe.local.response.type = "download"



# import frappe
# from openpyxl import Workbook
# from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
# from io import BytesIO
# from teampro.teampro.doctype.rec_week_plan.rec_week_plan import get_dsr_collapsible_html
# from openpyxl.utils import get_column_letter
# from datetime import datetime

# @frappe.whitelist()
# def download_dsr_excel(name, start_date=None, end_date=None, executive=None, team_type=None):
#     dsr_data = get_dsr_collapsible_html(name, start_date, end_date, executive, team_type)

#     wb = Workbook()
#     ws = wb.active
#     ws.title = "DSR"

#     # Styles
#     header_fill = PatternFill(start_color='001F5B', end_color='001F5B', fill_type='solid')
#     subheader_rc_fill = PatternFill(start_color='BBDEFB', end_color='BBDEFB', fill_type='solid')
#     subheader_ac_fill = PatternFill(start_color='E1F5FE', end_color='E1F5FE', fill_type='solid')
#     bold_font = Font(bold=True)
#     white_bold_font = Font(bold=True, color="FFFFFF")
#     center_align = Alignment(horizontal='center', vertical='center')
#     left_align = Alignment(horizontal='left', vertical='center')
#     thin_border = Border(left=Side(style='thin'),
#                          right=Side(style='thin'),
#                          top=Side(style='thin'),
#                          bottom=Side(style='thin'))

#     dates = dsr_data['raw_dates']
#     executives_data = dsr_data['data']

#     # ------------------- HEADER -------------------
#     ws.cell(row=1, column=1, value='S.No').font = white_bold_font
#     ws.cell(row=1, column=1).fill = header_fill
#     ws.cell(row=1, column=1).alignment = center_align

#     ws.cell(row=1, column=2, value='EXE').font = white_bold_font
#     ws.cell(row=1, column=2).fill = header_fill
#     ws.cell(row=1, column=2).alignment = center_align

#     ws.cell(row=1, column=3, value='Task').font = white_bold_font
#     ws.cell(row=1, column=3).fill = header_fill
#     ws.cell(row=1, column=3).alignment = center_align

#     ws.cell(row=1, column=4, value='Subject').font = white_bold_font
#     ws.cell(row=1, column=4).fill = header_fill
#     ws.cell(row=1, column=4).alignment = center_align

#     ws.cell(row=1, column=5, value='Total RC').font = white_bold_font
#     ws.cell(row=1, column=5).fill = header_fill
#     ws.cell(row=1, column=5).alignment = center_align

#     ws.cell(row=1, column=6, value='Total AC').font = white_bold_font
#     ws.cell(row=1, column=6).fill = header_fill
#     ws.cell(row=1, column=6).alignment = center_align

#     # Dates start from column 7
#     col = 7
#     for date in dates:
#         if isinstance(date, str):
#             date_obj = datetime.strptime(date, "%Y-%m-%d")
#         else:
#             date_obj = date
#         date_str = date_obj.strftime("%d %b")
#         ws.cell(row=1, column=col, value=date_str).font = white_bold_font
#         ws.cell(row=1, column=col).fill = header_fill
#         ws.cell(row=1, column=col).alignment = center_align
#         ws.merge_cells(start_row=1, start_column=col, end_row=1, end_column=col+1)

#         ws.cell(row=2, column=col, value='RC').fill = subheader_rc_fill
#         ws.cell(row=2, column=col).font = bold_font
#         ws.cell(row=2, column=col).alignment = center_align

#         ws.cell(row=2, column=col+1, value='AC').fill = subheader_ac_fill
#         ws.cell(row=2, column=col+1).font = bold_font
#         ws.cell(row=2, column=col+1).alignment = center_align

#         col += 2

#     # ------------------- DATA ROWS -------------------
#     current_row = 3
#     s_no = 1

#     for exe_email, tasks in executives_data.items():
#         # Executive summary row (S.No empty)
#         ws.cell(row=current_row, column=1, value='')  # S.No empty
#         ws.cell(row=current_row, column=2, value=exe_email).font = bold_font
#         ws.cell(row=current_row, column=2).alignment = left_align
#         # You can add totals here if needed, e.g., total RC/AC per executive
#         current_row += 1

#         # Project/task rows (S.No filled)
#         for task_code, task_info in tasks.items():
#             ws.cell(row=current_row, column=1, value=s_no).alignment = center_align
#             ws.cell(row=current_row, column=2, value=f"{task_info['project_name'] or task_info['project']}").alignment = left_align
#             ws.cell(row=current_row, column=3, value=task_code).alignment = left_align
#             ws.cell(row=current_row, column=4, value=task_info['subject']).alignment = left_align

#             # Total RC/AC for task
#             ws.cell(row=current_row, column=5, value=sum(task_info['dates'][d]['rt'] for d in dates)).alignment = center_align
#             ws.cell(row=current_row, column=6, value=sum(task_info['dates'][d]['ac'] for d in dates)).alignment = center_align

#             # Date-wise RC/AC
#             col_offset = 7
#             for date in dates:
#                 ws.cell(row=current_row, column=col_offset, value=task_info['dates'][date]['rt']).fill = subheader_rc_fill
#                 ws.cell(row=current_row, column=col_offset).alignment = center_align
#                 ws.cell(row=current_row, column=col_offset+1, value=task_info['dates'][date]['ac']).fill = subheader_ac_fill
#                 ws.cell(row=current_row, column=col_offset+1).alignment = center_align
#                 col_offset += 2

#             current_row += 1
#             s_no += 1 

#     ws.cell(row=current_row, column=1, value='Overall Totals').font = bold_font
#     ws.merge_cells(start_row=current_row, start_column=1, end_row=current_row, end_column=3)
#     ws.cell(row=current_row, column=1).alignment = center_align
#     ws.cell(row=current_row, column=4, value=overall_total_rc).font = bold_font
#     ws.cell(row=current_row, column=4).alignment = center_align
#     ws.cell(row=current_row, column=5, value=overall_total_ac).font = bold_font
#     ws.cell(row=current_row, column=5).alignment = center_align

#     col_offset = 6
#     for date in dates:
#         ws.cell(row=current_row, column=col_offset, value=daywise_rc_totals[date]).font = bold_font
#         ws.cell(row=current_row, column=col_offset).alignment = center_align
#         ws.cell(row=current_row, column=col_offset+1, value=daywise_ac_totals[date]).font = bold_font
#         ws.cell(row=current_row, column=col_offset+1).alignment = center_align
#         col_offset += 2

#     max_col = col_offset - 1
#     for row in ws.iter_rows(min_row=1, max_row=current_row, min_col=1, max_col=max_col):
#         for cell in row:
#             cell.border = thin_border

#     column_widths = {
#         1: 35,  
#         3: 35, 
#     }

#     for col, width in column_widths.items():
#         ws.column_dimensions[get_column_letter(col)].width = width

#     # Save to response
#     output = BytesIO()
#     wb.save(output)
#     output.seek(0)

#     frappe.local.response.filename = "DSR_Report.xlsx"
#     frappe.local.response.filecontent = output.getvalue()
#     frappe.local.response.type = "download"