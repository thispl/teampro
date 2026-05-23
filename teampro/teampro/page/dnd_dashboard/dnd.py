import frappe
from frappe import _
from frappe.utils import getdate, flt, today, nowdate
from frappe.utils import now_datetime
from datetime import datetime
from openpyxl.styles import PatternFill
from openpyxl.utils import get_column_letter
from openpyxl.styles import PatternFill, Font, Alignment, Border, Side
from openpyxl.utils import get_column_letter

import json

import json
import frappe
from openpyxl import Workbook
from openpyxl.styles import Font, Alignment, Border, Side
from openpyxl.drawing.image import Image
import frappe, io
import os
from frappe.utils.file_manager import get_file_path
from openpyxl.drawing.spreadsheet_drawing import AnchorMarker, OneCellAnchor
import base64


@frappe.whitelist()
def get_ptsr_combined_data():

    # =====================================================
    # INTERNAL
    # =====================================================

    internal_status = [
        "PSL",
        "Emigration",
        "Ticket",
        "Onboarding"
    ]

    internal_filters = {
        "status": ["in", internal_status],
        "nationality": "Indian"
    }

    internal_closures = frappe.db.get_all(
        "Closure",
        filters=internal_filters,
        fields=[
            "name",
            "given_name",
            "passport_no",
            "customer",
            "territory",
            "status",
            "remark",
            "last_updated_on",
            "sa_name",
            "sa_mobile_number",
            "associate",
            "mobile",
            "std_remarks",
            "custom_next_follow_up_on",
            "task",
            "project"
        ],
        order_by="last_updated_on desc"
    )

    # =====================================================
    # CANDIDATE
    # =====================================================

    candidate_status = [
        "Signed Offer Letter",
        "Premedical",
        "PCC",
        "Final Medical"
    ]

    candidate_filters = {
        "status": ["in", candidate_status],
        "nationality": "Indian",
        "sa_name": ["is", "not set"]
    }

    candidate_closures = frappe.db.get_all(
        "Closure",
        filters=candidate_filters,
        fields=[
            "name",
            "given_name",
            "passport_no",
            "customer",
            "territory",
            "status",
            "remark",
            "last_updated_on",
            "sa_name",
            "sa_mobile_number",
            "associate",
            "mobile",
            "std_remarks",
            "custom_next_follow_up_on",
            "task",
            "project"
        ],
        order_by="last_updated_on desc"
    )

    # =====================================================
    # COMMON FUNCTION
    # =====================================================

    project_cache = {}

    def process_rows(rows):

        for c in rows:

            # HISTORY
            history = frappe.get_all(
                "Closure Status History",
                filters={
                    "parent": c["name"],
                    "parenttype": "Closure",
                    "parentfield": "custom_history"
                },
                fields=["date"]
            )

            c["custom_history"] = history

            # PROJECT
            project = c.get("project")

            if project:

                if project not in project_cache:

                    project_cache[project] = frappe.db.get_value(
                        "Project",
                        project,
                        [
                            "name",
                            "project_name",
                            "tvac",
                            "tsp",
                            "tfp",
                            "tsl",
                            "custom_t_lp",
                            "custom_spoc_remark"
                        ],
                        as_dict=True
                    )

                project_data = project_cache.get(project) or {}

                c["project_id"] = (
                    project_data.get("name")
                    or project
                )

                c["project_name"] = (
                    project_data.get("project_name")
                    or project
                )

                c["tvac"] = project_data.get("tvac") or 0
                c["tsp"] = project_data.get("tsp") or 0
                c["tfp"] = project_data.get("tfp") or 0
                c["tsl"] = project_data.get("tsl") or 0
                c["custom_t_lp"] = (
                    project_data.get("custom_t_lp")
                    or 0
                )

                c["custom_spoc_remark"] = (
                    project_data.get("custom_spoc_remark")
                    or "-"
                )

            else:

                c["project_id"] = ""
                c["project_name"] = "No Project"

                c["tvac"] = 0
                c["tsp"] = 0
                c["tfp"] = 0
                c["tsl"] = 0
                c["custom_t_lp"] = 0
                c["custom_spoc_remark"] = "-"

            # TASK SUBJECT
            task = c.get("task")

            if task:

                c["task_subject"] = frappe.db.get_value(
                    "Task",
                    task,
                    "subject"
                ) or "No Position"

            else:

                c["task_subject"] = "No Position"

        return rows

    internal_closures = process_rows(internal_closures)
    candidate_closures = process_rows(candidate_closures)

    return {
        "internal": internal_closures,
        "candidate": candidate_closures
    }
