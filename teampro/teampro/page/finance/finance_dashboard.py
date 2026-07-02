import frappe
import math


#card-1
# @frappe.whitelist()
# def card(from_date=None, to_date=None):
#     conditions = ""
#     if from_date and to_date:
#         conditions = "AND posting_date BETWEEN %(from_date)s AND %(to_date)s"

#     data = frappe.db.sql(f"""
#         SELECT SUM(outstanding_amount) AS Outstanding
#         FROM `tabSales Invoice`
#         WHERE status NOT IN ('Return', 'Credit Note Issued', 'Paid', 'Cancelled')
#         AND docstatus = 1
#         {conditions}
#     """, {"from_date": from_date, "to_date": to_date}, as_dict=True)

#     return data[0].Outstanding or 0

@frappe.whitelist()
def card(from_date=None, to_date=None, overall_service=None):
    conditions = ["status NOT IN ('Return', 'Credit Note Issued', 'Paid', 'Cancelled')", "docstatus = 1"]

    if from_date and to_date:
        conditions.append("posting_date BETWEEN %(from_date)s AND %(to_date)s")
    if overall_service:
        conditions.append("services = %(overall_service)s")

    query = f"""
        SELECT services, SUM(outstanding_amount) as total
        FROM `tabSales Invoice`
        WHERE {' AND '.join(conditions)}
        GROUP BY services
    """

    rows = frappe.db.sql(query, {
        'from_date': from_date,
        'to_date': to_date,
        'overall_service': overall_service
    }, as_dict=True)

    group_map = {
        "BCS": "HRS", "REC-I": "HRS", "Payroll": "HRS", "SEP": "HRS", "REC-D": "HRS",
        "IT-SW": "ITS", "IT-IS": "ITS",
        "R&S": "CMN", "TGT": "CMN", "EMS": "CMN", "CMN": "CMN", "NL": "CMN",
        "TFP": "TFP",
        "HRIT": "HRIT"
    }

    group_totals = {"HRS": 0, "ITS": 0, "CMN": 0, "TFP": 0, "HRIT": 0}

    for row in rows:
        group = group_map.get(row.services)
        if group:
            group_totals[group] += row.total or 0

    overall_total = sum(group_totals.values())

    return {
        "total": overall_total,
        "groups": group_totals
    }


# card -2
# @frappe.whitelist()
# def card_1(from_date=None, to_date=None):
#     conditions = ""
#     if from_date and to_date:
#         conditions = "AND transaction_date BETWEEN %s AND %s"

#     card = frappe.db.sql(f"""
#         SELECT
#             SUM(base_grand_total - (base_grand_total * (per_billed / 100))) AS billed
#         FROM `tabSales Order`
#         WHERE status NOT IN ('On Hold', 'To Deliver', 'Closed', 'Cancelled', 'Completed')
#             AND docstatus = 1
#             {conditions}
#     """, (from_date, to_date) if from_date and to_date else (), as_dict=True)

#     return card[0].billed or 0


@frappe.whitelist()
def card_1(from_date=None, to_date=None, overall_service=None):
    conditions = ["status NOT IN ('On Hold', 'To Deliver', 'Closed', 'Cancelled', 'Completed')", "docstatus = 1"]

    if from_date and to_date:
        conditions.append("transaction_date BETWEEN %(from_date)s AND %(to_date)s")
    if overall_service:
        conditions.append("service = %(overall_service)s")

    query = f"""
        SELECT service, 
            SUM(base_grand_total - (base_grand_total * (per_billed / 100))) AS total
        FROM `tabSales Order`
        WHERE {' AND '.join(conditions)}
        GROUP BY service
    """

    rows = frappe.db.sql(query, {
        'from_date': from_date,
        'to_date': to_date,
        'overall_service': overall_service
    }, as_dict=True)

    group_map = {
        "BCS": "HRS", "REC-I": "HRS", "Payroll": "HRS", "SEP": "HRS", "REC-D": "HRS",
        "IT-SW": "ITS", "IT-IS": "ITS",
        "R&S": "CMN", "TGT": "CMN", "EMS": "CMN", "CMN": "CMN", "NL": "CMN",
        "TFP": "TFP",
        "HRIT": "HRIT"
    }

    group_totals = {"HRS": 0, "ITS": 0, "CMN": 0, "TFP": 0, "HRIT": 0}

    for row in rows:
        group = group_map.get(row.service)
        if group:
            group_totals[group] += row.total or 0

    # CLR value
    clr_conditions = ["status NOT IN ('Dropped', 'Onboarding','Onboarded','Arrived')", "so_created = 0"]
    if from_date and to_date:
        clr_conditions.append("posting_date BETWEEN %(from_date)s AND %(to_date)s")

    clr_result = frappe.db.sql(f"""
        SELECT SUM(
            COALESCE(client_payment_company_currency, 0) +
            COALESCE(candidate_payment_company_currenc, 0) +
            COALESCE(custom_associate_payment_company_currency, 0)
        ) AS payment
        FROM `tabClosure`
        WHERE {' AND '.join(clr_conditions)}
    """, {'from_date': from_date, 'to_date': to_date}, as_dict=True)

    # CLN value
    cln_conditions = ["status NOT IN ('On Hold', 'To Deliver', 'Closed', 'Cancelled', 'Completed')", "docstatus = 1"]
    if from_date and to_date:
        cln_conditions.append("transaction_date BETWEEN %(from_date)s AND %(to_date)s")

    cln_result = frappe.db.sql(f"""
        SELECT SUM(base_grand_total - (base_grand_total * (per_billed / 100) + advance_paid)) AS grand
        FROM `tabSales Order`
        WHERE {' AND '.join(cln_conditions)}
    """, {'from_date': from_date, 'to_date': to_date}, as_dict=True)

    overall_total = sum(group_totals.values())

    return {
        "total": overall_total,
        "groups": group_totals,
        "clr": clr_result[0].payment or 0,
        "cln": cln_result[0].grand or 0
    }


# card -3
@frappe.whitelist()
def card_2(from_date=None, to_date=None):
    conditions = ""
    params = ()

    if from_date and to_date:
        conditions = "AND transaction_date BETWEEN %s AND %s"
        params = (from_date, to_date)

    card_two = frappe.db.sql(f"""
        SELECT
       SUM(base_grand_total-(base_grand_total*(per_billed/100)+advance_paid)) AS `grand`
  FROM `tabSales Order`
 WHERE STATUS NOT IN ('On Hold', 'To Deliver', 'Closed', 'Cancelled','Completed')
   AND docstatus = 1
            {conditions}
    """, params, as_dict=True)

    return card_two[0].grand or 0


# card -4
@frappe.whitelist()
def card_3(from_date=None, to_date=None):
    conditions = "status NOT IN ('Dropped', 'Onboarding','Onboarded','Arrived') AND so_created = 0"
    
    if from_date and to_date:
        conditions += f" AND posting_date BETWEEN '{from_date}' AND '{to_date}'"

    card_three = frappe.db.sql(f"""
        SELECT SUM(
            COALESCE(client_payment_company_currency, 0) + 
            COALESCE(candidate_payment_company_currenc, 0)+
            COALESCE(custom_associate_payment_company_currency, 0)
        ) AS payment
        FROM `tabClosure`
        WHERE {conditions}
    """, as_dict=True)

    return card_three[0].payment or 0



#table
@frappe.whitelist()
def s_table(from_date=None, to_date=None, service=None, am=None, dm=None):

    all_services = frappe.db.sql("""
        SELECT DISTINCT services
        FROM `tabSales Invoice`
        WHERE services IS NOT NULL
    """, as_dict=True)

    service_totals = {row.services: 0 for row in all_services}

    
    conditions = """
        status NOT IN ('Return', 'Credit Note Issued', 'Paid', 'Cancelled')
        AND docstatus = 1
    """
    params = []

    if from_date and to_date:
        conditions += " AND posting_date BETWEEN %s AND %s"
        params += [from_date, to_date]
    if service:
        conditions += " AND services = %s"
        params.append(service)

    if am:
        conditions += " AND account_manager = %s"
        params.append(am)

    if dm:
        conditions += " AND delivery_manager = %s"
        params.append(dm)

    query = f"""
        SELECT 
            services,
            SUM(outstanding_amount) AS total
        FROM `tabSales Invoice`
        WHERE {conditions}
        GROUP BY services
    """

    results = frappe.db.sql(query, params, as_dict=True)

    for row in results:
        if row.services in service_totals:
            service_totals[row.services] = row.total or 0

   
    filtered_services = {s: v for s, v in service_totals.items() if v > 0}

    html = """
    <div style='max-height: 340px;'>
        <div style='min-width: 400px;'>
            <table class='table table-bordered' style='width: 100%; border-collapse: collapse;margin-top:10px;'>
                <thead>
                    <tr>
                        <th style="background: #002060; color: white; text-align: center;">S.No</th>
                        <th style="background: #002060; color: white; text-align: center;">Services</th>
                        <th style="background: #002060; color: white; text-align: center;">Outstanding</th>
                    </tr>
                </thead>
                <tbody>
    """

    grand_total = 0

    for idx, (service, value) in enumerate(filtered_services.items(), 1):
        grand_total += value

     
        if value >= 10000000:
            formatted = f"{value / 10000000:.2f} Cr"
        elif value >= 100000:
            formatted = f"{value / 100000:.2f} L"
        elif value >= 1000:
            formatted = f"{value / 1000:.2f} K"
        else:
            formatted = f"{value:.2f}"

        html += f"""
            <tr>
                <td style="text-align: center;">{idx}</td>
                <td style="white-space: nowrap;">{service}</td>
                <td style='text-align:right;'>{formatted}</td>
            </tr>
        """

    
    if filtered_services:
        if grand_total >= 10000000:
            grand_formatted = f"{grand_total / 10000000:.2f} Cr"
        elif grand_total >= 100000:
            grand_formatted = f"{grand_total / 100000:.2f} L"
        elif grand_total >= 1000:
            grand_formatted = f"{grand_total / 1000:.2f} K"
        else:
            grand_formatted = f"{grand_total:.2f}"

        html += f"""
            <tr>
                <td colspan="2" style="text-align: right; font-weight: bold;">Total</td>
                <td style="text-align: right; font-weight: bold;">{grand_formatted}</td>
            </tr>
        """

    html += """
                </tbody>
            </table>
        </div>
    </div>
    """
    if not filtered_services:
       
        html = """
        <div style="padding: 15px; text-align: center; background-color: #e1f0ff; color: #002060; border-radius: 8px; font-weight: bold; font-size: 16px;">
            Collection Outstanding ₹ 0.00
        </div>
        """

    return html


# table_2
# @frappe.whitelist()
# def so_billing(from_date=None, to_date=None,service=None,am=None,dm=None):

#     all_services = frappe.db.sql("""
#         SELECT DISTINCT service
#         FROM `tabSales Order`
#         WHERE service IS NOT NULL
#     """, as_dict=True)
#     service_totals = {row.service: 0 for row in all_services}

#     conditions = """
#         status NOT IN ('On Hold', 'To Deliver', 'Closed', 'Cancelled', 'Completed')
#         AND docstatus = 1
#     """
#     params = []
#     if from_date and to_date:
#         conditions += " AND transaction_date BETWEEN %s AND %s"
#         params += [from_date, to_date]
#     if service:
#         conditions += " AND service = %(service)s"
#         filters["service"] = service

#     if am:
#         conditions += " AND account_manager = %(am)s"
#         filters["am"] = am

#     query = f"""
#         SELECT 
#             service,
#             SUM(base_grand_total * ((100 - per_billed) / 100)) AS to_be_billed
#         FROM `tabSales Order`
#         WHERE {conditions}
#         GROUP BY service
#     """
#     results = frappe.db.sql(query, params, as_dict=True)
#     for row in results:
#         if row.service in service_totals:
#             service_totals[row.service] = row.to_be_billed or 0
#     filtered_services = {s: v for s, v in service_totals.items() if v > 0}
#     html = """
#     <div style='max-height: 340px;'>
#         <div style='min-width: 400px;'>
#             <table class='table table-bordered' style='width: 100%; border-collapse: collapse;'>
#                 <thead>
#                     <tr>
#                         <th style="background: #002060; color: white; text-align: center;">S.No</th>
#                         <th style="background: #002060; color: white; text-align: center;">Services</th>
#                         <th style="background: #002060; color: white; text-align: center;">To Be Billed</th>
#                     </tr>
#                 </thead>
#                 <tbody>
#     """

#     grand_total = 0

#     for idx, (service, value) in enumerate(filtered_services.items(), 1):
#         grand_total += value

        
#         if value >= 10000000:
#             formatted = f"{value / 10000000:.2f} Cr"
#         elif value >= 100000:
#             formatted = f"{value / 100000:.2f} L"
#         elif value >= 1000:
#             formatted = f"{value / 1000:.2f} K"
#         else:
#             formatted = f"{value:.2f}"

#         html += f"""
#             <tr>
#                 <td style="text-align: center;">{idx}</td>
#                 <td style="white-space: nowrap;">{service}</td>
#                 <td style='text-align:right;'>{formatted}</td>
#             </tr>
#         """

   
#     if filtered_services:
#         if grand_total >= 10000000:
#             total_formatted = f"{grand_total / 10000000:.2f} Cr"
#         elif grand_total >= 100000:
#             total_formatted = f"{grand_total / 100000:.2f} L"
#         elif grand_total >= 1000:
#             total_formatted = f"{grand_total / 1000:.2f} K"
#         else:
#             total_formatted = f"{grand_total:.2f}"

#         html += f"""
#             <tr>
#                 <td colspan="2" style="text-align: center; font-weight: bold;">Grand Total</td>
#                 <td style="text-align: right; font-weight: bold;">{total_formatted}</td>
#             </tr>
#         """

#         html += """
#                 </tbody>
#             </table>
#         </div>
#     </div>
#         """

#     else:
        
#         html = """
#         <div style="padding: 15px; text-align: center; background-color: #ffe0cc; color: #c04100; border-radius: 8px; font-weight: bold; font-size: 16px;">
#             SO To Be Billed ₹ 0.00
#         </div>
#         """

#     return html

@frappe.whitelist()
def so_billing(from_date=None, to_date=None, service=None, am=None, dm=None):

    all_services = frappe.db.sql("""
        SELECT DISTINCT service
        FROM `tabSales Order`
        WHERE service IS NOT NULL
    """, as_dict=True)

    service_totals = {row.service: 0 for row in all_services}

    conditions = """
        status NOT IN ('On Hold', 'To Deliver', 'Closed', 'Cancelled', 'Completed')
        AND docstatus = 1
    """

    values = {}

    if from_date and to_date:
        conditions += " AND transaction_date BETWEEN %(from_date)s AND %(to_date)s"
        values["from_date"] = from_date
        values["to_date"] = to_date

    if service:
        conditions += " AND service = %(service)s"
        values["service"] = service

    if am:
        conditions += " AND account_manager = %(am)s"
        values["am"] = am

    if dm:
        conditions += " AND delivery_manager = %(dm)s"
        values["dm"] = dm

    query = f"""
        SELECT 
            service,
            SUM(base_grand_total * ((100 - per_billed) / 100)) AS to_be_billed
        FROM `tabSales Order`
        WHERE {conditions}
        GROUP BY service
    """

    results = frappe.db.sql(query, values, as_dict=True)

    for row in results:
        if row.service in service_totals:
            service_totals[row.service] = row.to_be_billed or 0

    filtered_services = {s: v for s, v in service_totals.items() if v > 0}

    html = """
    <div style='max-height: 340px;'>
        <div style='min-width: 400px;'>
            <table class='table table-bordered' style='width: 100%; border-collapse: collapse;'>
                <thead>
                    <tr>
                        <th style="background: #002060; color: white; text-align: center;">S.No</th>
                        <th style="background: #002060; color: white; text-align: center;">Services</th>
                        <th style="background: #002060; color: white; text-align: center;">To Be Billed</th>
                    </tr>
                </thead>
                <tbody>
    """

    grand_total = 0

    for idx, (service, value) in enumerate(filtered_services.items(), 1):
        grand_total += value

        if value >= 10000000:
            formatted = f"{value / 10000000:.2f} Cr"
        elif value >= 100000:
            formatted = f"{value / 100000:.2f} L"
        elif value >= 1000:
            formatted = f"{value / 1000:.2f} K"
        else:
            formatted = f"{value:.2f}"

        html += f"""
            <tr>
                <td style="text-align: center;">{idx}</td>
                <td style="white-space: nowrap;">{service}</td>
                <td style='text-align:right;'>{formatted}</td>
            </tr>
        """

    if filtered_services:
        if grand_total >= 10000000:
            total_formatted = f"{grand_total / 10000000:.2f} Cr"
        elif grand_total >= 100000:
            total_formatted = f"{grand_total / 100000:.2f} L"
        elif grand_total >= 1000:
            total_formatted = f"{grand_total / 1000:.2f} K"
        else:
            total_formatted = f"{grand_total:.2f}"

        html += f"""
            <tr>
                <td colspan="2" style="text-align: center; font-weight: bold;">Grand Total</td>
                <td style="text-align: right; font-weight: bold;">{total_formatted}</td>
            </tr>
        """

    else:
        return """
        <div style="padding: 15px; text-align: center; background-color: #ffe0cc; color: #c04100; border-radius: 8px; font-weight: bold; font-size: 16px;">
            SO To Be Billed ₹ 0.00
        </div>
        """

    html += """
                </tbody>
            </table>
        </div>
    </div>
    """

    return html



@frappe.whitelist()
def so_collection(from_date=None, to_date=None):
    from frappe.utils import getdate

    filters = ""
    if from_date and to_date:
        filters = f"AND transaction_date BETWEEN '{getdate(from_date)}' AND '{getdate(to_date)}'"

    data = frappe.db.sql(f"""
        SELECT 
            service AS Services,
            SUM(base_grand_total - (base_grand_total * (per_billed / 100)) - advance_paid) AS raw_to_be_billed,
            CASE 
                WHEN SUM(base_grand_total - (base_grand_total * (per_billed / 100)) - advance_paid) >= 10000000 
                    THEN CONCAT(FORMAT(SUM(base_grand_total - (base_grand_total * (per_billed / 100)) - advance_paid) / 10000000, 2), 'Cr')
                WHEN SUM(base_grand_total - (base_grand_total * (per_billed / 100)) - advance_paid) >= 100000 
                    THEN CONCAT(FORMAT(SUM(base_grand_total - (base_grand_total * (per_billed / 100)) - advance_paid) / 100000, 2), 'L')
                WHEN SUM(base_grand_total - (base_grand_total * (per_billed / 100)) - advance_paid) >= 1000 
                    THEN CONCAT(FORMAT(SUM(base_grand_total - (base_grand_total * (per_billed / 100)) - advance_paid) / 1000, 2), 'K')
                ELSE FORMAT(SUM(base_grand_total - (base_grand_total * (per_billed / 100)) - advance_paid), 2)
            END AS To_Billed
        FROM `tabSales Order`
        WHERE status NOT IN ('On Hold', 'To Deliver', 'Closed', 'Cancelled', 'Completed')
            AND docstatus = 1
            {filters}
        GROUP BY service WITH ROLLUP
    """, as_dict=True)

    total_to_be_billed = 0
    content_rows = ""

    for idx, row in enumerate(data, 1):
        if row.Services is None:
            total_to_be_billed = row.raw_to_be_billed or 0
            continue
        if not row.Services:
            continue
        content_rows += f"""
            <tr>
                <td style="text-align: center;">{idx}</td>
                <td style="white-space: nowrap;">{row.Services}</td>
                <td style="text-align: right;">{row.To_Billed}</td>
            </tr>
        """

    if not content_rows:
        return f"""
            <div style="padding: 15px; text-align: center; background-color: #ffe0cc; color: #c04100; border-radius: 8px; font-weight: bold; font-size: 16px;">
            Sales Order - Billing & Payment Outstanding   ₹ 0.00
        </div>
        """

    if total_to_be_billed >= 10000000:
        formatted_total = f"{total_to_be_billed / 10000000:.2f}Cr"
    elif total_to_be_billed >= 100000:
        formatted_total = f"{total_to_be_billed / 100000:.2f}L"
    elif total_to_be_billed >= 1000:
        formatted_total = f"{total_to_be_billed / 1000:.2f}K"
    else:
        formatted_total = f"{total_to_be_billed:.2f}"

    html = f"""
    <div style='max-height: 340px;'>
        <div style='min-width: 400px;'>
            <table class='table table-bordered' style='width: 100%; border-collapse: collapse;'>
                <thead>
                    <tr>
                        <th style="background: #002060; color: white; text-align: center;">S.No</th>
                        <th style="background: #002060; color: white; text-align: center;">Services</th>
                        <th style="background: #002060; color: white; text-align: center;">To Be Billed</th>
                    </tr>
                </thead>
                <tbody>
                    {content_rows}
                    <tr>
                        <td colspan="2" style="text-align: center;">Grand Total</td>
                        <td style="text-align: right;">{formatted_total}</td>
                    </tr>
                </tbody>
            </table>
        </div>
    </div>
    """
    return html



@frappe.whitelist()
def so_payment(service=None, am=None,dm=None):

    from frappe.utils import today

    conditions = """
        status NOT IN ('Dropped', 'Onboarding', 'Onboarded', 'Arrived')
        AND so_created = 0
    """

    params = []

    if service:
        conditions += " AND service = %s"
        params.append(service)

    if am:
        conditions += " AND account_manager = %s"
        params.append(am)
    if dm:
        conditions+="AND candidate_owner=%s"
        params.append(dm)

    query = f"""
        SELECT 
            IFNULL(payment, 'Total') AS payment,

            SUM(
                COALESCE(client_payment_company_currency, 0)
                + COALESCE(candidate_payment_company_currenc, 0)
                + COALESCE(custom_associate_payment_company_currency, 0)
            ) AS total_amount

        FROM `tabClosure`

        WHERE {conditions}

        GROUP BY payment WITH ROLLUP
    """

    data = frappe.db.sql(query, params, as_dict=True)

    html = """
    <div style='max-height: 340px; overflow-y: auto;'>
        <div style='min-width: 400px;'>
            <table class='table table-bordered'
                style='width: 100%; border-collapse: collapse;'>

                <thead>
                    <tr>
                        <th style="background:#002060;color:white;text-align:center;">
                            S.No
                        </th>

                        <th style="background:#002060;color:white;text-align:center;">
                            Payment Type
                        </th>

                        <th style="background:#002060;color:white;text-align:center;">
                            Payment
                        </th>
                    </tr>
                </thead>

                <tbody>
    """

    total_row = None
    row_no = 1

    for row in data:

        if row.payment == "Total":
            total_row = row
            continue

        value = row.total_amount or 0

        if value >= 10000000:
            formatted = f"{value / 10000000:.2f} Cr"

        elif value >= 100000:
            formatted = f"{value / 100000:.2f} L"

        elif value >= 1000:
            formatted = f"{value / 1000:.2f} K"

        else:
            formatted = f"{value:.2f}"

        html += f"""
            <tr>

                <td style="text-align:center;">
                    {row_no}
                </td>

                <td style="white-space:nowrap;">
                    {row.payment}
                </td>

                <td style="text-align:right;">
                    {formatted}
                </td>

            </tr>
        """

        row_no += 1

    if total_row:

        total_value = total_row.total_amount or 0

        if total_value >= 10000000:
            total_formatted = f"{total_value / 10000000:.2f} Cr"

        elif total_value >= 100000:
            total_formatted = f"{total_value / 100000:.2f} L"

        elif total_value >= 1000:
            total_formatted = f"{total_value / 1000:.2f} K"

        else:
            total_formatted = f"{total_value:.2f}"

        html += f"""
            <tr>

                <td colspan="2"
                    style="text-align:center;font-weight:bold;">

                    Total

                </td>

                <td style="text-align:right;font-weight:bold;">

                    {total_formatted}

                </td>

            </tr>
        """

    html += """
                </tbody>
            </table>
        </div>
    </div>
    """

    return html

#card

@frappe.whitelist()
def in_bank(from_date=None, to_date=None):
    conditions = ""
    if from_date and to_date:
        conditions = f"posting_date BETWEEN '{from_date}' AND '{to_date}'"
    else:
        conditions = "posting_date <= CURDATE()"

    sql = f"""
        SELECT 
            SUM((IFNULL(g.opening_debit, 0) + IFNULL(sq.total_debit, 0)) - 
                (IFNULL(g.opening_credit, 0) + IFNULL(sq.total_credit, 0))) AS closing_balance
        FROM `tabCompany` c
        LEFT JOIN (
            SELECT 
                company,
                SUM(CASE WHEN is_opening = 'Yes' THEN debit ELSE 0 END) AS opening_debit,
                SUM(CASE WHEN is_opening = 'Yes' THEN credit ELSE 0 END) AS opening_credit
            FROM `tabGL Entry`
            WHERE {conditions}
              AND account IN (
                  '777705160983 - ICICI Bank - THIS',
                  '50200054611436 - HDFC - THIS',
                  '50200082906246-HDFC - This',
                  '777705755022 - ICICI Bank - TGTP',
                  '50200059117831 - HDFC Bank - TFP'
              )
              AND is_cancelled = 0
            GROUP BY company
        ) g ON c.name = g.company
        LEFT JOIN (
            SELECT 
                company,
                SUM(debit_in_account_currency) AS total_debit,
                SUM(credit_in_account_currency) AS total_credit
            FROM `tabGL Entry`
            WHERE {conditions}
              AND account IN (
                  '777705160983 - ICICI Bank - THIS',
                  '50200054611436 - HDFC - THIS',
                  '50200082906246-HDFC - This',
                  '777705755022 - ICICI Bank - TGTP',
                  '50200059117831 - HDFC Bank - TFP'
              )
              AND is_opening = 'No'
              AND is_cancelled = 0
            GROUP BY company
        ) sq ON c.name = sq.company
    """

    result = frappe.db.sql(sql, as_dict=1)
    return result[0].closing_balance if result else 0



#card_cash
@frappe.whitelist()
def in_cash(from_date=None, to_date=None):
    conditions = ""
    if from_date and to_date:
        conditions = f"posting_date BETWEEN '{from_date}' AND '{to_date}'"
    else:
        conditions = "posting_date <= CURDATE()"

    result = frappe.db.sql(f"""
        SELECT 
            SUM((IFNULL(g.opening_debit, 0) + IFNULL(sq.debit, 0)) - 
                (IFNULL(g.opening_credit, 0) + IFNULL(sq.credit, 0))) AS `Closing_Balance`
        FROM `tabCompany` c
        LEFT JOIN (
            SELECT company,
                   SUM(CASE WHEN is_opening = 'Yes' THEN debit ELSE 0 END) AS opening_debit,
                   SUM(CASE WHEN is_opening = 'Yes' THEN credit ELSE 0 END) AS opening_credit
              FROM `tabGL Entry`
             WHERE {conditions}
               AND account = 'Cash - THIS'
               AND is_cancelled = 0
             GROUP BY company
        ) g ON c.name = g.company
        LEFT JOIN (
            SELECT company,
                   SUM(debit_in_account_currency) AS debit,
                   SUM(credit_in_account_currency) AS credit
              FROM `tabGL Entry`
             WHERE {conditions}
               AND account IN ('Cash - THIS', 'Cash - TFP', 'Cash - TGTP')
               AND is_opening = 'No'
               AND is_cancelled = 0
             GROUP BY company
        ) sq ON c.name = sq.company
    """, as_dict=1)

    return result[0].Closing_Balance if result else 0

@frappe.whitelist()
def sfd(from_date=None, to_date=None):
    conditions = ""
    if from_date and to_date:
        conditions = f"AND posting_date BETWEEN '{from_date}' AND '{to_date}'"

    result = frappe.db.sql(f"""
        SELECT
            SUM(
                IFNULL(g.opening_debit, 0) + IFNULL(sq.debit, 0)
                - IFNULL(g.opening_credit, 0) - IFNULL(sq.credit, 0)
            ) AS `Closing_Balance`
        FROM `tabCompany` c
        LEFT JOIN (
            SELECT company,
                   SUM(CASE WHEN is_opening = 'Yes' THEN debit ELSE 0 END) AS opening_debit,
                   SUM(CASE WHEN is_opening = 'Yes' THEN credit ELSE 0 END) AS opening_credit
            FROM `tabGL Entry`
            WHERE account IN ('Fixed Deposits - THIS','Fixed Deposits - TGTP')
              AND is_cancelled = 0
              {conditions}
            GROUP BY company
        ) g ON c.name = g.company
        LEFT JOIN (
            SELECT company,
                   SUM(debit_in_account_currency) AS debit,
                   SUM(credit_in_account_currency) AS credit
            FROM `tabGL Entry`
            WHERE account IN ('Fixed Deposits - THIS','Fixed Deposits - TGTP')
              AND is_opening = 'No'
              AND is_cancelled = 0
              {conditions}
            GROUP BY company
        ) sq ON c.name = sq.company
    """, as_dict=1)

    return result[0].Closing_Balance or 0
@frappe.whitelist()
def lfd(from_date=None, to_date=None):
    conditions = ""
    if from_date and to_date:
        conditions = f"AND posting_date BETWEEN '{from_date}' AND '{to_date}'"

    result = frappe.db.sql(f"""
        SELECT
            SUM(
                IFNULL(g.opening_debit, 0) + IFNULL(sq.debit, 0)
                - IFNULL(g.opening_credit, 0) - IFNULL(sq.credit, 0)
            ) AS `Closing_Balance`
        FROM `tabCompany` c
        LEFT JOIN (
            SELECT company,
                   SUM(CASE WHEN is_opening = 'Yes' THEN debit ELSE 0 END) AS opening_debit,
                   SUM(CASE WHEN is_opening = 'Yes' THEN credit ELSE 0 END) AS opening_credit
            FROM `tabGL Entry`
            WHERE account IN ('Fixed Deposits(Res) - THIS')
              AND is_cancelled = 0
              {conditions}
            GROUP BY company
        ) g ON c.name = g.company
        LEFT JOIN (
            SELECT company,
                   SUM(debit_in_account_currency) AS debit,
                   SUM(credit_in_account_currency) AS credit
            FROM `tabGL Entry`
            WHERE account IN ('Fixed Deposits(Res) - THIS')
              AND is_opening = 'No'
              AND is_cancelled = 0
              {conditions}
            GROUP BY company
        ) sq ON c.name = sq.company
    """, as_dict=1)

    return result[0].Closing_Balance or 0
# @frappe.whitelist()
# def po(from_date=None, to_date=None):
    conditions = ""
    if from_date and to_date:
        conditions = f"AND transaction_date BETWEEN '{from_date}' AND '{to_date}'"

    result = frappe.db.sql(f"""
        SELECT 
            SUM(base_grand_total - (base_grand_total * (per_billed / 100))) AS `To_Billed`
        FROM `tabPurchase Order`
        WHERE status NOT IN ('On Hold', 'To Receive', 'Closed', 'Cancelled', 'Completed')
          AND docstatus = 1
          {conditions}
    """, as_dict=1)

    return result[0].To_Billed or 0


@frappe.whitelist()
def fund_card(from_date=None, to_date=None):
    def get_conditions(col="posting_date", with_and=False):
        if from_date and to_date:
            prefix = "AND " if with_and else ""
            return f"{prefix}{col} BETWEEN '{from_date}' AND '{to_date}'"
        return f"{col} <= CURDATE()" if not with_and else ""

    # Bank
    bank = frappe.db.sql(f"""
        SELECT SUM((IFNULL(g.opening_debit,0)+IFNULL(sq.total_debit,0))-
                   (IFNULL(g.opening_credit,0)+IFNULL(sq.total_credit,0))) AS val
        FROM `tabCompany` c
        LEFT JOIN (
            SELECT company,
                SUM(CASE WHEN is_opening='Yes' THEN debit ELSE 0 END) AS opening_debit,
                SUM(CASE WHEN is_opening='Yes' THEN credit ELSE 0 END) AS opening_credit
            FROM `tabGL Entry`
            WHERE {get_conditions()} AND account IN (
                '777705160983 - ICICI Bank - THIS','50200054611436 - HDFC - THIS',
                '50200082906246-HDFC - This','777705755022 - ICICI Bank - TGTP',
                '50200059117831 - HDFC Bank - TFP')
            AND is_cancelled=0 GROUP BY company
        ) g ON c.name=g.company
        LEFT JOIN (
            SELECT company,
                SUM(debit_in_account_currency) AS total_debit,
                SUM(credit_in_account_currency) AS total_credit
            FROM `tabGL Entry`
            WHERE {get_conditions()} AND account IN (
                '777705160983 - ICICI Bank - THIS','50200054611436 - HDFC - THIS',
                '50200082906246-HDFC - This','777705755022 - ICICI Bank - TGTP',
                '50200059117831 - HDFC Bank - TFP')
            AND is_opening='No' AND is_cancelled=0 GROUP BY company
        ) sq ON c.name=sq.company
    """, as_dict=1)

    # Cash
    cash = frappe.db.sql(f"""
        SELECT SUM((IFNULL(g.opening_debit,0)+IFNULL(sq.debit,0))-
                   (IFNULL(g.opening_credit,0)+IFNULL(sq.credit,0))) AS val
        FROM `tabCompany` c
        LEFT JOIN (
            SELECT company,
                SUM(CASE WHEN is_opening='Yes' THEN debit ELSE 0 END) AS opening_debit,
                SUM(CASE WHEN is_opening='Yes' THEN credit ELSE 0 END) AS opening_credit
            FROM `tabGL Entry`
            WHERE {get_conditions()} AND account='Cash - THIS' AND is_cancelled=0
            GROUP BY company
        ) g ON c.name=g.company
        LEFT JOIN (
            SELECT company,
                SUM(debit_in_account_currency) AS debit,
                SUM(credit_in_account_currency) AS credit
            FROM `tabGL Entry`
            WHERE {get_conditions()} AND account IN ('Cash - THIS','Cash - TFP','Cash - TGTP')
            AND is_opening='No' AND is_cancelled=0 GROUP BY company
        ) sq ON c.name=sq.company
    """, as_dict=1)

    # SFD
    sfd = frappe.db.sql(f"""
        SELECT SUM(IFNULL(g.opening_debit,0)+IFNULL(sq.debit,0)
                  -IFNULL(g.opening_credit,0)-IFNULL(sq.credit,0)) AS val
        FROM `tabCompany` c
        LEFT JOIN (
            SELECT company,
                SUM(CASE WHEN is_opening='Yes' THEN debit ELSE 0 END) AS opening_debit,
                SUM(CASE WHEN is_opening='Yes' THEN credit ELSE 0 END) AS opening_credit
            FROM `tabGL Entry`
            WHERE account IN ('Fixed Deposits - THIS','Fixed Deposits - TGTP')
            AND is_cancelled=0 {get_conditions(with_and=True)} GROUP BY company
        ) g ON c.name=g.company
        LEFT JOIN (
            SELECT company,
                SUM(debit_in_account_currency) AS debit,
                SUM(credit_in_account_currency) AS credit
            FROM `tabGL Entry`
            WHERE account IN ('Fixed Deposits - THIS','Fixed Deposits - TGTP')
            AND is_opening='No' AND is_cancelled=0 {get_conditions(with_and=True)} GROUP BY company
        ) sq ON c.name=sq.company
    """, as_dict=1)

    # LFD
    lfd = frappe.db.sql(f"""
        SELECT SUM(IFNULL(g.opening_debit,0)+IFNULL(sq.debit,0)
                  -IFNULL(g.opening_credit,0)-IFNULL(sq.credit,0)) AS val
        FROM `tabCompany` c
        LEFT JOIN (
            SELECT company,
                SUM(CASE WHEN is_opening='Yes' THEN debit ELSE 0 END) AS opening_debit,
                SUM(CASE WHEN is_opening='Yes' THEN credit ELSE 0 END) AS opening_credit
            FROM `tabGL Entry`
            WHERE account IN ('Fixed Deposits(Res) - THIS')
            AND is_cancelled=0 {get_conditions(with_and=True)} GROUP BY company
        ) g ON c.name=g.company
        LEFT JOIN (
            SELECT company,
                SUM(debit_in_account_currency) AS debit,
                SUM(credit_in_account_currency) AS credit
            FROM `tabGL Entry`
            WHERE account IN ('Fixed Deposits(Res) - THIS')
            AND is_opening='No' AND is_cancelled=0 {get_conditions(with_and=True)} GROUP BY company
        ) sq ON c.name=sq.company
    """, as_dict=1)

    bank_val  = bank[0].val  or 0
    cash_val  = cash[0].val  or 0
    sfd_val   = sfd[0].val   or 0
    lfd_val   = lfd[0].val   or 0

    return {
        "bank": bank_val,
        "cash": cash_val,
        "sfd":  sfd_val,
        "lfd":  lfd_val,
        "total": bank_val + cash_val
    }



@frappe.whitelist()
def po(from_date=None, to_date=None):
    conditions = ""
    if from_date and to_date:
        conditions = f"AND transaction_date BETWEEN '{from_date}' AND '{to_date}'"

    result = frappe.db.sql(f"""
        SELECT 
            SUM(base_grand_total - (base_grand_total * (per_billed / 100))) AS To_Billed
        FROM `tabPurchase Order`
        WHERE status NOT IN ('On Hold', 'To Receive', 'Closed', 'Cancelled', 'Completed')
          AND docstatus = 1
          {conditions}
    """, as_dict=1)

    cln_result = frappe.db.sql(f"""
        SELECT SUM(base_grand_total - advance_paid) AS billed
        FROM `tabPurchase Order`
        WHERE status NOT IN ('On Hold', 'To Receive', 'Closed', 'Cancelled', 'Completed')
          AND docstatus = 1
          {conditions}
    """, as_dict=1)

    service_rows = frappe.db.sql(f"""
        SELECT 
            custom_service,
            SUM(base_grand_total - (base_grand_total * (per_billed / 100))) AS total
        FROM `tabPurchase Order`
        WHERE status NOT IN ('On Hold', 'To Receive', 'Closed', 'Cancelled', 'Completed')
          AND docstatus = 1
          AND custom_service IS NOT NULL
          {conditions}
        GROUP BY custom_service
    """, as_dict=1)

    group_map = {
        "BCS": "HRS", "REC-I": "HRS", "Payroll": "HRS", "SEP": "HRS", "REC-D": "HRS",
        "IT-SW": "ITS", "IT-IS": "ITS",
        "R&S": "CMN", "TGT": "CMN", "EMS": "CMN", "CMN": "CMN", "NL": "CMN",
        "TFP": "TFP",
        "HRIT": "HRIT"
    }

    groups = {"HRS": 0, "ITS": 0, "CMN": 0, "TFP": 0, "HRIT": 0}
    for row in service_rows:
        grp = group_map.get(row.custom_service)
        if grp:
            groups[grp] += row.total or 0

    return {
        "total": result[0].To_Billed or 0,
        "cln": cln_result[0].billed or 0,
        "groups": groups
    }

@frappe.whitelist()
def po_payment(from_date=None, to_date=None):
    conditions = ""
    if from_date and to_date:
        conditions = f"AND transaction_date BETWEEN '{from_date}' AND '{to_date}'"

    result = frappe.db.sql(f"""
        SELECT 
       SUM(base_grand_total-advance_paid) AS `billed`
  FROM `tabPurchase Order`
 WHERE STATUS NOT IN ('On Hold', 'To Receive', 'Closed', 'Cancelled','Completed')
   AND docstatus = 1
          {conditions}
    """, as_dict=1)

    return result[0].billed or 0
# @frappe.whitelist()
# def po_out(from_date=None, to_date=None):
#     conditions = ""
#     if from_date and to_date:
#         conditions = f"AND posting_date BETWEEN '{from_date}' AND '{to_date}'"

#     result = frappe.db.sql(f"""
#         SELECT 
#             SUM(outstanding_amount) AS Outstanding
#         FROM `tabPurchase Invoice`
#         WHERE status NOT IN ('Return', 'Debit Note Issued', 'Paid', 'Cancelled')
#           AND docstatus = 1
#           {conditions}
#     """, as_dict=1)

#     return result[0].Outstanding or 0


@frappe.whitelist()
def po_out(from_date=None, to_date=None):
    conditions = ""
    if from_date and to_date:
        conditions = f"AND posting_date BETWEEN '{from_date}' AND '{to_date}'"

    result = frappe.db.sql(f"""
        SELECT 
            SUM(outstanding_amount) AS Outstanding
        FROM `tabPurchase Invoice`
        WHERE status NOT IN ('Return', 'Debit Note Issued', 'Paid', 'Cancelled')
          AND docstatus = 1
          {conditions}
    """, as_dict=1)

    service_rows = frappe.db.sql(f"""
        SELECT 
            services,
            SUM(outstanding_amount) AS total
        FROM `tabPurchase Invoice`
        WHERE status NOT IN ('Return', 'Debit Note Issued', 'Paid', 'Cancelled')
          AND docstatus = 1
          AND services IS NOT NULL
          {conditions}
        GROUP BY services
    """, as_dict=1)

    group_map = {
        "BCS": "HRS", "REC-I": "HRS", "Payroll": "HRS", "SEP": "HRS", "REC-D": "HRS",
        "IT-SW": "ITS", "IT-IS": "ITS",
        "R&S": "CMN", "TGT": "CMN", "EMS": "CMN", "CMN": "CMN", "NL": "CMN",
        "TFP": "TFP",
        "HRIT": "HRIT"
    }

    groups = {"HRS": 0, "ITS": 0, "CMN": 0, "TFP": 0, "HRIT": 0}
    for row in service_rows:
        grp = group_map.get(row.services)
        if grp:
            groups[grp] += row.total or 0

    return {
        "total": result[0].Outstanding or 0,
        "groups": groups
    }


#table

@frappe.whitelist()
def in_table():
    from frappe.utils import today, getdate, nowdate, fmt_money
    from datetime import datetime

    fiscal_year = frappe.db.get_value(
        "Fiscal Year",
        filters={"year_start_date": ["<=", today()], "year_end_date": [">=", today()]},
        fieldname=["year_start_date", "year_end_date"],
        as_dict=True
    )

    from_date = fiscal_year["year_start_date"]
    to_date = fiscal_year["year_end_date"]

    filters = {
        'from_date': from_date,
        'to_date': to_date
    }

    data = frappe.db.sql("""
       SELECT c.abbr AS Company,
       ((IFNULL(g_hand.opening_debit, 0) + IFNULL(SUM(sq_hand.debit), 0)) - (IFNULL(g_hand.opening_credit, 0) + IFNULL(SUM(sq_hand.credit), 0))) AS CIH_value,
       ((IFNULL(g_bank.opening_debit, 0) + IFNULL(SUM(sq_bank.debit), 0)) - (IFNULL(g_bank.opening_credit, 0) + IFNULL(SUM(sq_bank.credit), 0))) AS CIB_value
  FROM `tabCompany` c
  LEFT JOIN (
        SELECT company,
               SUM(CASE WHEN is_opening = 'Yes' THEN debit ELSE 0 END) AS opening_debit,
               SUM(CASE WHEN is_opening = 'Yes' THEN credit ELSE 0 END) AS opening_credit
          FROM `tabGL Entry`
         WHERE posting_date <= NOW()
           AND account = 'Cash - THIS'
           AND is_cancelled = 0
         GROUP BY company
       ) g_hand ON c.name = g_hand.company
  LEFT JOIN (
        SELECT company,
               SUM(debit_in_account_currency) AS debit,
               SUM(credit_in_account_currency) AS credit
          FROM `tabGL Entry`
         WHERE posting_date <= NOW()
           AND account IN ('Cash - THIS', 'Cash - TFP', 'Cash - TGTP')
           AND is_opening = 'No'
           AND is_cancelled = 0
         GROUP BY company
       ) sq_hand ON c.name = sq_hand.company
  LEFT JOIN (
        SELECT company,
               SUM(CASE WHEN is_opening = 'Yes' THEN debit ELSE 0 END) AS opening_debit,
               SUM(CASE WHEN is_opening = 'Yes' THEN credit ELSE 0 END) AS opening_credit
          FROM `tabGL Entry`
         WHERE posting_date <= NOW()
           AND account IN ('777705160983 - ICICI Bank - THIS', '50200054611436 - HDFC - THIS', '50200082906246-HDFC - This', '777705755022 - ICICI Bank - TGTP', '50200059117831 - HDFC Bank - TFP')
           AND is_cancelled = 0
         GROUP BY company
       ) g_bank ON c.name = g_bank.company
  LEFT JOIN (
        SELECT company,
               SUM(debit_in_account_currency) AS debit,
               SUM(credit_in_account_currency) AS credit
          FROM `tabGL Entry`
         WHERE posting_date <= NOW()
           AND account IN ('777705160983 - ICICI Bank - THIS', '50200054611436 - HDFC - THIS', '50200082906246-HDFC - This', '777705755022 - ICICI Bank - TGTP', '50200059117831 - HDFC Bank - TFP')
           AND is_opening = 'No'
           AND is_cancelled = 0
         GROUP BY company
       ) sq_bank ON c.name = sq_bank.company
 WHERE c.name NOT IN ("TEAMPRO General Trading", "TEAMPRO Saudi Arabia")
 GROUP BY c.abbr
    """, filters, as_dict=True)

    # Calculate totals
    total_cih = sum(row['CIH_value'] or 0 for row in data)
    total_cib = sum(row['CIB_value'] or 0 for row in data)

    def format_amount(value):
        if value >= 10000000:
            return f"{value / 10000000:.2f}Cr"
        elif value >= 100000:
            return f"{value / 100000:.2f}L"
        elif value >= 1000:
            return f"{value / 1000:.2f}K"
        else:
            return f"{value:.2f}"

    html = """
    <div style='max-height: 340px; overflow-x: auto;'>
        <div style='min-width: 400px;'>
            <table class='table table-bordered' style='width: 100%; border-collapse: collapse;'>
                <thead>
                    <tr>
                        <th style="background: #002060; color: white; text-align: center;">S.No</th>
                        <th style="background: #002060; color: white; text-align: center;">Company</th>
                        <th style="background: #002060; color: white; text-align: center;">CIH</th>
                        <th style="background: #002060; color: white; text-align: center;">CIB</th>
                    </tr>
                </thead>
                <tbody>
    """

    for idx, row in enumerate(data, 1):
        html += f"""
            <tr>
                <td style="text-align: center;">{idx}</td>
                <td style="white-space: nowrap;">{row.Company or ""}</td>
                <td style="text-align: right;">{format_amount(row.CIH_value or 0)}</td>
                <td style="text-align: right;">{format_amount(row.CIB_value or 0)}</td>
            </tr>
        """

    # Grand total row
    html += f"""
        <tr style="background: #f0f0f0; font-weight: bold;">
            <td colspan="2" style="text-align: center;">Total</td>
            <td style="text-align: right;">{format_amount(total_cih)}</td>
            <td style="text-align: right;">{format_amount(total_cib)}</td>
        </tr>
    """

    html += """
                </tbody>
            </table>
        </div>
    </div>
    """

    return html



@frappe.whitelist()
def table_two():
    from frappe.utils import today, flt
    from datetime import datetime

    fiscal_year = frappe.db.get_value(
        "Fiscal Year",
        filters={"year_start_date": ["<=", today()], "year_end_date": [">=", today()]},
        fieldname=["year_start_date", "year_end_date"],
        as_dict=True
    )

    from_date = fiscal_year["year_start_date"]
    to_date = fiscal_year["year_end_date"]

    filters = {
        'from_date': from_date,
        'to_date': to_date
    }

    data = frappe.db.sql("""
        SELECT 
            IFNULL(custom_service, 'Total') AS Service,
            CASE 
                WHEN SUM(base_grand_total - (base_grand_total * (per_billed / 100))) >= 10000000 
                    THEN CONCAT(FORMAT(SUM(base_grand_total - (base_grand_total * (per_billed / 100))) / 10000000, 2), 'Cr')
                WHEN SUM(base_grand_total - (base_grand_total * (per_billed / 100))) >= 100000 
                    THEN CONCAT(FORMAT(SUM(base_grand_total - (base_grand_total * (per_billed / 100))) / 100000, 2), 'L')
                WHEN SUM(base_grand_total - (base_grand_total * (per_billed / 100))) >= 1000 
                    THEN CONCAT(FORMAT(SUM(base_grand_total - (base_grand_total * (per_billed / 100))) / 1000, 2), 'K')
                ELSE FORMAT(SUM(base_grand_total - (base_grand_total * (per_billed / 100))), 2)
            END AS `to_billed`
        FROM `tabPurchase Order`
        WHERE `status` NOT IN ('On Hold', 'To Receive', 'Closed', 'Cancelled', 'Completed')
          AND `docstatus` = 1
        GROUP BY `custom_service` WITH ROLLUP
    """, filters, as_dict=True)

    total = 0
    html = """
    <div style='max-height: 340px; overflow-y: auto;'>
        <div style='min-width: 400px;'>
            <table class='table table-bordered' style='width: 100%; border-collapse: collapse;'>
                <thead>
                    <tr>
                        <th style="background: #002060; color: white; text-align: center;">S.No</th>
                        <th style="background: #002060; color: white; text-align: center;">Service</th>
                        <th style="background: #002060; color: white; text-align: center;">To Be Billed</th>
                    </tr>
                </thead>
                <tbody>
    """

    for idx, row in enumerate(data):
        if row.Service == 'Total':
            continue  

        raw_value = 0
        display_value = row.to_billed or "0"

    
        try:
            if 'Cr' in display_value:
                raw_value = flt(display_value.replace('Cr', '')) * 10000000
            elif 'L' in display_value:
                raw_value = flt(display_value.replace('L', '')) * 100000
            elif 'K' in display_value:
                raw_value = flt(display_value.replace('K', '')) * 1000
            else:
                raw_value = flt(display_value)
        except:
            raw_value = 0

        total += raw_value

        html += f"""
            <tr>
                <td style="text-align: center;">{idx + 1}</td>
                <td style="white-space: nowrap;">{row.Service}</td>
                <td style="text-align: right;">{display_value}</td>
            </tr>
        """

 
    if total >= 10000000:
        total_formatted = f"{total / 10000000:.2f}Cr"
    elif total >= 100000:
        total_formatted = f"{math.floor((total / 100000) * 100) / 100:.2f}L"

    elif total >= 1000:
        total_formatted = f"{total / 1000:2f}K"
    else:
        total_formatted = f"{total:.2f}"


    html += f"""
        <tr>
            <td colspan="2" style="text-align: center;">Total</td>
            <td style="text-align: right;">{total_formatted}</td>
        </tr>
    """

    html += """
                </tbody>
            </table>
        </div>
    </div>
    """

    return html




@frappe.whitelist()
def table_three():
    from frappe.utils import today, fmt_money
    from datetime import datetime

    
    fiscal_year = frappe.db.get_value(
        "Fiscal Year",
        filters={"year_start_date": ["<=", today()], "year_end_date": [">=", today()]},
        fieldname=["year_start_date", "year_end_date"],
        as_dict=True
    )

    from_date = fiscal_year["year_start_date"]
    to_date = fiscal_year["year_end_date"]

    filters = {
        'from_date': from_date,
        'to_date': to_date
    }


    data = frappe.db.sql("""
       SELECT custom_service AS Services,
       sum(base_grand_total) AS base_grand_total,
       sum(advance_paid) AS advance_paid,
       sum(per_billed) AS per_billed,
       sum(base_grand_total-advance_paid) AS to_be_billed,
       CASE WHEN SUM(base_grand_total-advance_paid) >= 10000000 THEN CONCAT(FORMAT(SUM(base_grand_total-advance_paid) / 10000000, 2), 'Cr')
            WHEN SUM(base_grand_total-advance_paid) >= 100000   THEN CONCAT(FORMAT(SUM(base_grand_total-advance_paid) / 100000, 2), 'L')
            WHEN SUM(base_grand_total-advance_paid) >= 1000     THEN CONCAT(FORMAT(SUM(base_grand_total-advance_paid) / 1000, 2), 'K')
            ELSE FORMAT(SUM(base_grand_total-advance_paid), 2)
             END AS `To Be Billed`
  FROM `tabPurchase Order`
 WHERE STATUS NOT IN ('On Hold', 'To Receive', 'Closed', 'Cancelled','Completed')
   AND docstatus = 1
 GROUP BY custom_service WITH ROLLUP
    """, filters, as_dict=True)

 
    data = [row for row in data if row.Services is not None]

  
    total_grand_total = sum([row.base_grand_total or 0 for row in data])
    total_advance = sum([row.advance_paid or 0 for row in data])
    total_per_billed = sum([row.per_billed or 0 for row in data])
    total_to_be_billed = sum([row.to_be_billed or 0 for row in data])

   
    if total_to_be_billed >= 10000000:
        formatted_total_to_be_billed = f"{total_to_be_billed / 10000000:.2f}Cr"
    elif total_to_be_billed >= 100000:
        formatted_total_to_be_billed = f"{total_to_be_billed / 100000:.2f}L"
    elif total_to_be_billed >= 1000:
        formatted_total_to_be_billed = f"{total_to_be_billed / 1000:.2f}K"
    else:
        formatted_total_to_be_billed = f"{total_to_be_billed:.2f}"

    html = """
    <div style='max-height: 340px; overflow-y:auto;'>
        <div style='min-width: 800px;'>
            <table class='table table-bordered' style='width: 100%; border-collapse: collapse;'>
                <thead>
                    <tr>
                        <th style="background: #002060; color: white; text-align: center;">S.No</th>
                        <th style="background: #002060; color: white; text-align: center;">Service</th>
                        <th style="background: #002060; color: white; text-align: center;">Base Grand Total</th>
                        <th style="background: #002060; color: white; text-align: center;">Advance Paid</th>
                        <th style="background: #002060; color: white; text-align: center;">Pre Billed</th>
                        <th style="background: #002060; color: white; text-align: center;">To Be Billed</th>
                    </tr>
                </thead>
                <tbody>
    """

    for idx, row in enumerate(data, 1):
        html += f"""
            <tr>
                <td style="text-align: center;">{idx}</td>
                <td style="white-space: nowrap;">{row.Services or ""}</td>
                <td style='text-align:right;'>{fmt_money(row.base_grand_total)}</td>
                <td style='text-align:right;'>{fmt_money(row.advance_paid)}</td>
                <td style='text-align:right;'>{round(row.per_billed or 0, 2)}</td>
                <td style='text-align:right;'>{row['To Be Billed']}</td>
            </tr>
        """

   
    html += f"""
        <tr style="background: #f0f0f0; font-weight: bold;">
            <td colspan="2" style="text-align: center;">Total</td>
            <td style="text-align: right;">{fmt_money(total_grand_total)}</td>
            <td style="text-align: right;">{fmt_money(total_advance)}</td>
            <td style="text-align: right;">{round(total_per_billed, 2)}</td>
            <td style="text-align: right;">{formatted_total_to_be_billed}</td>
        </tr>
    """

    html += """
                </tbody>
            </table>
        </div>
    </div>
    """

    return html




@frappe.whitelist()
def out_amount():
    from frappe.utils import today, fmt_money

 
    fiscal_year = frappe.db.get_value(
        "Fiscal Year",
        filters={"year_start_date": ["<=", today()], "year_end_date": [">=", today()]},
        fieldname=["year_start_date", "year_end_date"],
        as_dict=True
    )

    from_date = fiscal_year["year_start_date"]
    to_date = fiscal_year["year_end_date"]

  
    data = frappe.db.sql("""
        SELECT services AS Service,
               SUM(outstanding_amount) AS raw_outstanding
        FROM `tabPurchase Invoice`
        WHERE status NOT IN ('Return', 'Debit Note Issued', 'Paid', 'Cancelled')
          AND docstatus = 1
        GROUP BY services WITH ROLLUP
    """, as_dict=True)

 
    data = [row for row in data if row.Service is not None]

   
    for row in data:
        amount = row.raw_outstanding or 0
        row.formatted_outstanding = f"{amount:,.2f}"
        if amount >= 1e7:
            row.abbr = f"{amount / 1e7:.2f}Cr"
        elif amount >= 1e5:
            row.abbr = f"{amount / 1e5:.2f}L"
        elif amount >= 1e3:
            row.abbr = f"{amount / 1e3:.2f}K"
        else:
            row.abbr = f"{amount:.2f}"

  
    total_amount = sum(row.raw_outstanding or 0 for row in data)
    formatted_total = f"{total_amount:,.2f}"
    if total_amount >= 1e7:
        abbr_total = f"{total_amount / 1e7:.2f}Cr"
    elif total_amount >= 1e5:
        abbr_total = f"{total_amount / 1e5:.2f}L"
    elif total_amount >= 1e3:
        abbr_total = f"{total_amount / 1e3:.2f}K"
    else:
        abbr_total = f"{total_amount:.2f}"

    
    html = """
    <div style='max-height: 340px;'>
        <div style='min-width: 400px;'>
            <table class='table table-bordered' style='width: 100%; border-collapse: collapse;'>
                <thead>
                    <tr>
                        <th style="background: #002060; color: white; text-align: center;">S.No</th>
                        <th style="background: #002060; color: white; text-align: center;">Service</th>
                        <th style="background: #002060; color: white; text-align: center;">Outstanding</th>
                    </tr>
                </thead>
                <tbody>
    """

    for idx, row in enumerate(data, 1):
        html += f"""
            <tr>
                <td style="text-align: center;">{idx}</td>
                <td>{row.Service}</td>
                <td style="text-align: right;">{row.abbr}</td>
            </tr>
        """

    html += f"""
        <tr style="background: #f0f0f0; font-weight: bold;">
            <td colspan="2" style="text-align: center;">Total</td>
            <td style="text-align: right;">{abbr_total}</td>
        </tr>
    """

    html += """
                </tbody>
            </table>
        </div>
    </div>
    """

    return html


# app_name/app/api/dashboard.py

@frappe.whitelist()
def chart():
    data = frappe.db.sql("""
        SELECT 
            service AS Services,
            SUM(base_grand_total - (base_grand_total * (per_billed / 100))) AS to_be_billed,
            CASE 
                WHEN SUM(base_grand_total - (base_grand_total * (per_billed / 100))) >= 10000000 
                    THEN CONCAT(FORMAT(SUM(base_grand_total - (base_grand_total * (per_billed / 100))) / 10000000, 2), 'Cr')
                WHEN SUM(base_grand_total - (base_grand_total * (per_billed / 100))) >= 100000   
                    THEN CONCAT(FORMAT(SUM(base_grand_total - (base_grand_total * (per_billed / 100))) / 100000, 2), 'L')
                WHEN SUM(base_grand_total - (base_grand_total * (per_billed / 100))) >= 1000     
                    THEN CONCAT(FORMAT(SUM(base_grand_total - (base_grand_total * (per_billed / 100))) / 1000, 2), 'K')
                ELSE FORMAT(SUM(base_grand_total - (base_grand_total * (per_billed / 100))), 2)
            END AS `To Be Billed`
        FROM `tabSales Order`
        WHERE status NOT IN ('On Hold', 'To Deliver', 'Closed', 'Cancelled', 'Completed')
        AND docstatus = 1
        GROUP BY service
    """, as_dict=True)
    return data

@frappe.whitelist()
def chart_2():
    data = frappe.db.sql("""
        SELECT service AS Services,
       SUM(base_grand_total - (base_grand_total * (per_billed / 100)) - advance_paid) AS to_be_billed,
       CASE WHEN SUM(base_grand_total - (base_grand_total * (per_billed / 100)) - advance_paid) >= 10000000 THEN CONCAT(FORMAT(SUM(base_grand_total - (base_grand_total * (per_billed / 100)) - advance_paid) / 10000000, 2), 'Cr')
            WHEN SUM(base_grand_total - (base_grand_total * (per_billed / 100)) - advance_paid) >= 100000   THEN CONCAT(FORMAT(SUM(base_grand_total - (base_grand_total * (per_billed / 100)) - advance_paid) / 100000, 2), 'L')
            WHEN SUM(base_grand_total - (base_grand_total * (per_billed / 100)) - advance_paid) >= 1000     THEN CONCAT(FORMAT(SUM(base_grand_total - (base_grand_total * (per_billed / 100)) - advance_paid) / 1000, 2), 'K')
            ELSE FORMAT(SUM(base_grand_total - (base_grand_total * (per_billed / 100)) - advance_paid), 2)
             END AS `To Be Billed`
  FROM `tabSales Order`
 WHERE STATUS NOT IN ('On Hold', 'To Deliver', 'Closed', 'Cancelled', 'Completed')
   AND docstatus = 1
 GROUP BY service
    """, as_dict=True)
    return data

@frappe.whitelist()
def chart_3():
    data = frappe.db.sql("""
       SELECT custom_service AS Services,
       sum(base_grand_total) AS base_grand_total,
       sum(advance_paid) AS advance_paid,
       sum(per_billed) AS per_billed,
       sum(base_grand_total-(base_grand_total*(per_billed/100)+advance_paid)) AS to_be_billed,
       CASE WHEN SUM(base_grand_total-(base_grand_total*(per_billed/100)+advance_paid)) >= 10000000 THEN CONCAT(FORMAT(SUM(base_grand_total-(base_grand_total*(per_billed/100)+advance_paid)) / 10000000, 2), 'Cr')
            WHEN SUM(base_grand_total-(base_grand_total*(per_billed/100)+advance_paid)) >= 100000   THEN CONCAT(FORMAT(SUM(base_grand_total-(base_grand_total*(per_billed/100)+advance_paid)) / 100000, 2), 'L')
            WHEN SUM(base_grand_total-(base_grand_total*(per_billed/100)+advance_paid)) >= 1000     THEN CONCAT(FORMAT(SUM(base_grand_total-(base_grand_total*(per_billed/100)+advance_paid)) / 1000, 2), 'K')
            ELSE FORMAT(SUM(base_grand_total-(base_grand_total*(per_billed/100)+advance_paid)), 2)
             END AS `To Be Billed`
  FROM `tabPurchase Order`
 WHERE STATUS NOT IN ('On Hold', 'To Receive', 'Closed', 'Cancelled','Completed')
   AND docstatus = 1
 GROUP BY custom_service

    """, as_dict=True)
    return data

@frappe.whitelist()
def chart_4():
    res = frappe.db.sql("""
        SELECT 
            IFNULL(custom_service, 'Others') AS Service,
            SUM(base_grand_total - (base_grand_total * (per_billed / 100))) AS to_be_billed_value,
            CASE 
                WHEN SUM(base_grand_total - (base_grand_total * (per_billed / 100))) >= 100000 
                    THEN CONCAT(FORMAT(SUM(base_grand_total - (base_grand_total * (per_billed / 100))) / 100000, 2), 'L')
                WHEN SUM(base_grand_total - (base_grand_total * (per_billed / 100))) >= 1000 
                    THEN CONCAT(FORMAT(SUM(base_grand_total - (base_grand_total * (per_billed / 100))) / 1000, 2), 'K')
                ELSE FORMAT(SUM(base_grand_total - (base_grand_total * (per_billed / 100))), 2)
            END AS `To Be Billed`
        FROM `tabPurchase Order`
        WHERE status NOT IN ('On Hold', 'To Receive', 'Closed', 'Cancelled', 'Completed')
          AND docstatus = 1
        GROUP BY IFNULL(custom_service, 'Others')
    """, as_dict=True)

    return res

@frappe.whitelist()
def chart_5():
    res = frappe.db.sql("""
      SELECT IFNULL(TRIM(services), 'No Service') AS Service,
       SUM(outstanding_amount) AS outstanding_amount,
       CASE WHEN SUM(outstanding_amount) >= 100000 THEN CONCAT(FORMAT(SUM(outstanding_amount) / 100000, 2), 'L')
            WHEN SUM(outstanding_amount) >= 1000   THEN CONCAT(FORMAT(SUM(outstanding_amount) / 1000, 2), 'K')
            ELSE FORMAT(SUM(outstanding_amount), 2)
       END AS Outstanding
FROM `tabPurchase Invoice`
WHERE status NOT IN ('Return', 'Debit Note Issued', 'Paid', 'Cancelled')
  AND docstatus = 1
GROUP BY TRIM(services)

    """, as_dict=True)

    return res


@frappe.whitelist()
def epnc_table(from_date=None, to_date=None):
    from frappe.utils import get_first_day, get_last_day, today, getdate

    if not from_date or not to_date:
        from_date = get_first_day(today())
        to_date = get_last_day(today())

    active_emps = frappe.get_all("Employee",filters={"status": "Active"},pluck="name")

    if not active_emps:
        return "<p>No active employees found.</p>"

 
    data = frappe.db.sql("""
        SELECT
            emp,
            MAX(emp_name) AS emp_name,
            MAX(department) AS department,
            SUM(energy_score) AS ep_score,
            SUM(nc_score) AS nc_score
        FROM `tabEnergy Point And Non Conformity`
        WHERE
            docstatus = 1
            AND emp IN %(emp_list)s
            AND DATE(creation) BETWEEN %(start)s AND %(end)s
        GROUP BY emp
        ORDER BY emp_name
    """, {
        "emp_list": tuple(active_emps),
        "start": from_date,
        "end": to_date
    }, as_dict=True)



    html = """
    <div style='max-height: 340px; overflow-y: auto; overflow-x: auto;'>
        <div style='min-width: 500px;'>
            <table class='table table-bordered' style='width: 100%; border-collapse: collapse;'>
                <thead>
                    <tr>
                        <th style="position: sticky; top: 0; background: #002060; color: white; text-align: center;">S.No</th>
                        <th style="position: sticky; top: 0; background: #002060; color: white; text-align: center; white-space: nowrap;">Employee ID</th>
                        <th style="position: sticky; top: 0; background: #002060; color: white; text-align: center;">Employee Name</th>
                        <th style="position: sticky; top: 0; background: #002060; color: white; text-align: center;">EP Score</th>
                        <th style="position: sticky; top: 0; background: #002060; color: white; text-align: center; white-space: nowrap;">NP Score</th>
                        <th style="position: sticky; top: 0; background: #002060; color: white; text-align: center;">ENS</th>
                    </tr>
                </thead>
                <tbody>
    """

    for i, row in enumerate(data, start=1):
        total = (row.ep_score or 0) + (row.nc_score or 0)
        html += f"""
        <tr>
            <td style="text-align:center;">{i}</td>
            <td>{row.emp}</td>
            <td>{row.emp_name}</td>
            <td style="text-align:center; color:{'green' if row.ep_score > 0 else 'black'};">{row.ep_score or 0}</td>
            <td style="text-align:center; color:{'red' if row.nc_score < 0 else 'black'};">{row.nc_score or 0}</td>
            <td style="text-align:center; color:{'green' if total > 0 else 'red' if total < 0 else 'black'};">{total}</td>
        </tr>
        """

    html += """
                </tbody>
            </table>
        </div>
    </div>
    """

    return html

@frappe.whitelist()
def vm_sales():

    from frappe.utils import today, formatdate
    def format_inr(value):
        value = float(value or 0)

        if value >= 10000000:
            return f"₹ {value / 10000000:.2f} Cr"
        elif value >= 100000:
            return f"₹ {value / 100000:.2f} L"
        elif value >= 1000:
            return f"₹ {value / 1000:.2f} K"
        else:
            return f"₹ {value:,.2f}"

    # Fiscal year
    fiscal_year = frappe.db.get_value(
        "Fiscal Year",
        filters={
            "year_start_date": ["<=", today()],
            "year_end_date": [">=", today()]
        },
        fieldname=["year_start_date", "year_end_date"],
        as_dict=True
    )

    from_date = fiscal_year.year_start_date
    to_date = fiscal_year.year_end_date

    # Get warehouses under vending machines
    warehouses = frappe.db.sql("""
        SELECT name, disabled, warehouse_name
        FROM `tabWarehouse`
        WHERE parent_warehouse = 'LSVM - Vending Machines - TFP'
    """, as_dict=True)

    result = []

    for wh in warehouses:

        warehouse_name = wh.name
        display_name = wh.warehouse_name
        status = "Active" if not wh.disabled else "Inactive"
        stock = frappe.db.sql("""
            SELECT SUM(actual_qty) AS qty
            FROM `tabBin`
            WHERE warehouse = %s
        """, (warehouse_name,), as_dict=True)[0].qty or 0
        pos_profile=frappe.db.get_value("POS Profile",{"disabled":0,"warehouse":warehouse_name},["name"])
        sales = frappe.db.sql("""
            SELECT
                SUM(CASE 
                        WHEN posting_date BETWEEN %s AND %s 
                        THEN grand_total ELSE 0 
                    END) AS ytd,

                SUM(CASE 
                        WHEN MONTH(posting_date) = MONTH(CURDATE())
                        AND YEAR(posting_date) = YEAR(CURDATE())
                        THEN grand_total ELSE 0 
                    END) AS mtd,

                SUM(grand_total) AS overall
            FROM `tabSales Invoice`
            WHERE docstatus = 1
              AND status NOT IN ('Return', 'Credit Note Issued', 'Cancelled')
              AND pos_profile = %s
        """, (from_date, to_date, pos_profile), as_dict=True)[0]

        last_pos = frappe.db.sql("""
            SELECT posting_date
            FROM `tabPOS Invoice`
            WHERE docstatus = 1
              AND pos_profile = %s
            ORDER BY posting_date DESC
            LIMIT 1
        """, (pos_profile,), as_dict=True)

        last_pos_date = ""
        if last_pos and last_pos[0].posting_date:
            last_pos_date = formatdate(last_pos[0].posting_date, "dd-MM-yyyy")

        result.append({
            "vm_id": warehouse_name,
            "vm_name": display_name,
            "status": status,
            "stock": stock,
           "mtd": format_inr(sales.mtd),
            "ytd": format_inr(sales.ytd),
            "overall": format_inr(sales.overall),
            "last_pos": last_pos_date
        })

   
    html = """
    <div style="max-height:400px; overflow:auto;">
        <table class="table table-bordered">
            <thead style="position: sticky; top: 0; z-index: 2; background:#002060; color:white;">
                <tr style="background:#002060;color:white;text-align:center;">
                    <th>VM_ID</th>
                    <th>Vending Machine</th>
                    <th>Status</th>
                    <th>Cr. Stock</th>
                    <th>MTD</th>
                    <th>YTD</th>
                    <th>Overall</th>
                    <th>Last POS</th>
                </tr>
            </thead>
            <tbody>
    """

    for r in result:
        html += f"""
        <tr>
            <td>{r['vm_id']}</td>
            <td>{r['vm_name']}</td>
            <td>{r['status']}</td>
            <td style="text-align:right;">{r['stock']:.2f}</td>
           <td style="text-align:right;">{r['mtd']}</td>
            <td style="text-align:right;">{r['ytd']}</td>
            <td style="text-align:right;">{r['overall']}</td>
            <td>{r['last_pos']}</td>
        </tr>
        """

    html += "</tbody></table></div>"

    return html


@frappe.whitelist()
def retail_shops():

    from frappe.utils import today, formatdate

    # -------------------------
    # Financial Year
    # -------------------------
    fiscal_year = frappe.db.get_value(
        "Fiscal Year",
        filters={
            "year_start_date": ["<=", today()],
            "year_end_date": [">=", today()]
        },
        fieldname=["year_start_date", "year_end_date"],
        as_dict=True
    )

    from_date = fiscal_year.year_start_date
    to_date = fiscal_year.year_end_date

  
    shops = frappe.get_all(
        "Shop RC",
        filters={"workflow_state": "Approved"},
        fields=["name", "customer_name", "city"],
        order_by="name asc"
    )

    result = []

    def format_inr(v):
        v = float(v or 0)
        if v >= 10000000:
            return f"₹ {v/10000000:.2f} Cr"
        elif v >= 100000:
            return f"₹ {v/100000:.2f} L"
        elif v >= 1000:
            return f"₹ {v/1000:.2f} K"
        return f"₹ {v:,.2f}"

    for shop in shops:

      
        customer_id = frappe.db.get_value(
            "Customer",
            {"custom_retail_customer": shop.name},
            "name"
        )

        if not customer_id:
            continue

        warehouse = frappe.db.get_value(
            "Warehouse",
            {"custom_retail_customer": shop.name, "disabled": 0},
            "name"
        )

        if not warehouse:
            continue

      
        stock = frappe.db.sql("""
            SELECT SUM(actual_qty) AS qty
            FROM `tabBin`
            WHERE warehouse = %s
              AND actual_qty > 0
        """, warehouse, as_dict=True)[0].qty or 0

    
        outstanding = frappe.db.sql("""
            SELECT
                SUM(total_amount) AS outstanding_bill,
                SUM(total_amount - paid_amount) AS outstanding_payment
            FROM `tabRS Invoice`
            WHERE customer = %s
              AND docstatus = 1
        """, customer_id, as_dict=True)[0]
        sales = frappe.db.sql("""
            SELECT
                SUM(CASE 
                        WHEN posting_date BETWEEN %s AND %s 
                        THEN grand_total ELSE 0 
                    END) AS ytd,

                SUM(CASE 
                        WHEN MONTH(posting_date) = MONTH(CURDATE())
                        AND YEAR(posting_date) = YEAR(CURDATE())
                        THEN grand_total ELSE 0 
                    END) AS mtd,

                SUM(grand_total) AS overall
            FROM `tabSales Invoice`
            WHERE docstatus = 1
              AND status NOT IN ('Return','Credit Note Issued','Cancelled')
              AND customer = %s
           
        """, (from_date, to_date, customer_id), as_dict=True)[0]

        last_rsi = frappe.db.sql("""
            SELECT invoice_date
            FROM `tabRS Invoice`
            WHERE customer = %s
              AND docstatus = 1
            ORDER BY invoice_date DESC
            LIMIT 1
        """, customer_id, as_dict=True)

        last_rsi_date = ""
        if last_rsi and last_rsi[0].invoice_date:
            last_rsi_date = formatdate(last_rsi[0].invoice_date, "dd-MM-yyyy")

     
        result.append({
            "shop_id": shop.name,
            "shop_name": shop.customer_name,
            "status": "Active",

            "stock": stock,
            "outstanding_bill": format_inr(outstanding.outstanding_bill),
            "outstanding_payment": format_inr(outstanding.outstanding_payment),

            "mtd": format_inr(sales.mtd),
            "ytd": format_inr(sales.ytd),
            "overall": format_inr(sales.overall),

            "last_rsi": last_rsi_date
        })

    
    html = """
    <div style="max-height:320px; overflow:auto;">
        <table class="table table-bordered" style="margin-bottom:0;">
            <thead style="position: sticky; top: 0; z-index: 2; background:#002060; color:white;">
                <tr style="background:#002060;color:white;text-align:center;">
                    <th>Shop ID</th>
                    <th>Shop Name</th>
                    <th>Status</th>
                    <th>Cr. Stock</th>
                    <th>Outstanding Bill</th>
                    <th>Outstanding Payment</th>
                    <th>MTD</th>
                    <th>YTD</th>
                    <th>Overall</th>
                    <th>Last RSI</th>
                </tr>
            </thead>
            <tbody>
    """

    for r in result:
        html += f"""
        <tr>
            <td style="white-space:nowrap">{r['shop_id']}</td>
            <td style="white-space:nowrap">{r['shop_name']}</td>
            <td style="white-space:nowrap">{r['status']}</td>
            <td style="text-align:right;white-space:nowrap">{r['stock']:.2f}</td>
            <td style="text-align:right;white-space:nowrap">{r['outstanding_bill']}</td>
            <td style="text-align:right;white-space:nowrap">{r['outstanding_payment']}</td>
            <td style="text-align:right;white-space:nowrap">{r['mtd']}</td>
            <td style="text-align:right;white-space:nowrap">{r['ytd']}</td>
            <td style="text-align:right;white-space:nowrap">{r['overall']}</td>
            <td style="white-space:nowrap">{r['last_rsi']}</td>
        </tr>
        """

    html += "</tbody></table></div>"

    return html




@frappe.whitelist()
def target_vs_achievement():

    today = getdate(nowdate())

    target_data = frappe.db.sql("""
        SELECT
            name,
            employee,
            employee_name,
            target_based_unit,
            annual_ct,
            total_ct_yta,
            total_ct,
            custom_total_target_point,
            total_ct_achieved,
            custom_total_achieved_point
        FROM `tabTarget Manager`
        WHERE custom_year_start_date <= %s
        AND custom_year_end_date >= %s
        ORDER BY employee_name
    """, (today, today), as_dict=True)
    rows = ""

    for idx, d in enumerate(target_data):

        bg_color = "#e7e6ec" if idx % 2 == 0 else "#ffffff"

        rows += f"""
        <tr style="background-color:{bg_color};">
            <td>{d.name}</td>
            <td>{d.employee or ''}</td>
            <td>{d.employee_name or ''}</td>
            <td>{d.target_based_unit or ''}</td>

            <td style="text-align:right;">{frappe.utils.fmt_money(d.annual_ct or 0, precision=2)}</td>
            <td style="text-align:right;">{frappe.utils.fmt_money(d.total_ct_yta or 0, precision=2)}</td>
            <td style="text-align:right;">{frappe.utils.fmt_money(d.total_ct or 0, precision=2)}</td>
            <td style="text-align:right;">{d.custom_total_target_point or 0:.2f}%</td>
            <td style="text-align:right;">{frappe.utils.fmt_money(d.total_ct_achieved or 0, precision=2)}</td>
            <td style="text-align:right;">{d.custom_total_achieved_point or 0:.2f}%</td>
        </tr>
        """

    return rows


from io import BytesIO
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill
from openpyxl.utils import get_column_letter
import frappe
from frappe.utils import getdate, nowdate
from openpyxl.styles import Font, PatternFill, Border, Side

@frappe.whitelist()
def download_target_achievement():

    today = getdate(nowdate())

    data = frappe.db.sql("""
        SELECT
            name,
            employee,
            employee_name,
            target_based_unit,
            annual_ct,
            total_ct_yta,
            total_ct,
            custom_total_target_point,
            total_ct_achieved,
            custom_total_achieved_point
        FROM `tabTarget Manager`
        WHERE custom_year_start_date <= %s
        AND custom_year_end_date >= %s
        ORDER BY employee_name
    """, (today, today), as_list=True)

    wb = Workbook()
    ws = wb.active
    ws.title = "Target Achievement"
    # Border Style
    thin_border = Border(
        left=Side(style='thin', color='000000'),
        right=Side(style='thin', color='000000'),
        top=Side(style='thin', color='000000'),
        bottom=Side(style='thin', color='000000')
    )

    headers = [
        "ID",
        "Employee",
        "Employee Name",
        "Target Based On",
        "Target (INR)",
        "Total YTA Target (INR)",
        "Total Target (INR)",
        "Total Target (Point)",
        "Total Achieved (INR)",
        "Total Achieved (Point)"
    ]

    # Header
    header_fill = PatternFill("solid", fgColor="002060")
    header_font = Font(color="FFFFFF", bold=True)

    for col_num, header in enumerate(headers, 1):
        cell = ws.cell(row=1, column=col_num)
        cell.value = header
        cell.fill = header_fill
        cell.font = header_font
        cell.border =thin_border

    # Column Width
    widths = [20, 20, 25, 20, 18, 20, 18, 18, 20, 20]
    for i, width in enumerate(widths, 1):
        ws.column_dimensions[get_column_letter(i)].width = width

    # Row Colors
    odd_fill = PatternFill("solid", fgColor="E7E6EC")
    even_fill = PatternFill("solid", fgColor="FFFFFF")
    currency_columns = [5, 6, 7, 9]
    percentage_columns = [8, 10]



    for row_idx, row_data in enumerate(data, start=2):

        for col_idx, value in enumerate(row_data, start=1):
            cell = ws.cell(row=row_idx, column=col_idx)
            cell.value = value

            if row_idx % 2 == 0:
                cell.fill = even_fill
            else:
                cell.fill = odd_fill
            
            cell.border = thin_border

            # Currency Format
            if col_idx in currency_columns and value is not None:
                cell.number_format = '₹#,##0.00'

            if col_idx in percentage_columns and value is not None:
                cell.number_format = '0.00%'

    output = BytesIO()
    wb.save(output)

    frappe.response["filename"] = "Target_vs_Achievement.xlsx"
    frappe.response["filecontent"] = output.getvalue()
    frappe.response["type"] = "binary"

    