import frappe
from frappe.utils.csvutils import read_csv_content
from frappe.utils import get_first_day, get_last_day, format_datetime, get_url_to_form
from frappe.utils import cint
from frappe.utils.data import date_diff, now_datetime, nowdate, today, add_days
import datetime
from frappe import _
from frappe import _
from frappe.model.document import Document
from frappe.utils import getdate, nowdate
from frappe import throw, msgprint
import frappe
from datetime import date
from frappe import throw, _
from frappe.utils import getdate, today
today = date.today()
from frappe.model.document import Document
import datetime
import frappe,erpnext
from frappe.utils import cint
from frappe.utils import validate_email_address
import json
from frappe.utils import date_diff, add_months,today,add_days,add_years,nowdate,flt
from frappe.model.mapper import get_mapped_doc
from frappe.utils.file_manager import get_file
from frappe.utils.csvutils import UnicodeWriter, read_csv_content
import datetime
from datetime import date,datetime,timedelta
import openpyxl
from openpyxl import Workbook
import openpyxl
import xlrd
import re
from frappe.utils import today
from openpyxl import Workbook
from openpyxl.styles import PatternFill, Font, Alignment
from openpyxl.utils import get_column_letter
from datetime import datetime
from io import BytesIO
from openpyxl.styles import Font, Alignment, Border, Side
from openpyxl import load_workbook
from openpyxl.utils import get_column_letter
from openpyxl.styles import GradientFill, PatternFill
import pandas as pd
from frappe.utils import formatdate
from frappe.utils import now
from erpnext.setup.utils import get_exchange_rate
from datetime import date
from six import BytesIO, string_types
from frappe.utils import time_diff
from frappe.utils.csvutils import read_csv_content
from erpnext.buying.doctype.purchase_order.purchase_order import update_status
from frappe.utils.file_manager import get_file
from urllib.parse import urlencode
from frappe.model.rename_doc import rename_doc
from datetime import datetime
from frappe.utils import today

@frappe.whitelist()
def update_case_status_billed(doc,method):
	if doc.services=="BCS" and doc.items:
		for i in doc.items:
			if i.item_code:
				if frappe.db.exists("Case",{"name":i.item_code}):
					frappe.db.set_value("Case",i.item_code,"case_status","Billed")
					frappe.db.set_value("Case",i.item_code,"billed_date",today())

@frappe.whitelist()
def clear_payment_table_si(doc,method):
	doc.payment_schedule=[]

@frappe.whitelist()
def calc_cut_off_prize(doc,method):
	for f in doc.items:
		tfp_item = frappe.db.get_value(
				"Item",
				f.item_code,
				["tfp", "custom_is_retail_item"],
				as_dict=1
				)

		tfp = tfp_item.tfp or 0
		retail_item = tfp_item.custom_is_retail_item or 0
		# tfp_item = frappe.db.sql("""select tfp from `tabItem` where name = '%s' """%(f.item_code),as_dict=1)[0]
		# tfp = tfp_item['tfp']
		if tfp == 1 and retail_item == 0:
			price_list = frappe.db.sql("""select price_list_rate from `tabItem Price` where price_list = 'Cut Off Price' and item_code = '%s' """%(f.item_code),as_dict=1)
			for p in price_list:
				if f.uom == 'Gram':
					item_price = (p.price_list_rate / 1000)
					item_rate = round((item_price),2)
					if f.rate < item_rate:
						frappe.throw(_(' %s Rate is lesser than Cut Off Price')%(f.item_name))
				elif f.uom == 'Kg':
					if f.rate < p.price_list_rate:
						frappe.throw(_(' %s Rate is lesser than Cut Off Price')%(f.item_name))

import frappe
@frappe.whitelist()
def calculate_advances_invoice(doc, method):
	total_advance_doc_currency = 0
	for d in doc.advances:
		total_advance_doc_currency += d.custom_advance_amount_doc_currency or 0

	doc.custom_total_advance_doc_currency = total_advance_doc_currency
	doc.custom_outstanding_amount_doc_currency = (doc.grand_total or 0) - total_advance_doc_currency

@frappe.whitelist()
def validate_maintain_stok_si(doc,method):
	if not doc.custom_rc_invoice_:
		if doc.pos_profile not in ["Main Store","VM1_Precision","VM2_INFAC - TFP"]:
			if doc.items:
				for i in doc.items:
					stock=frappe.db.get_value("Item",i.item_code,"is_stock_item")
					if stock and not i.delivery_note:
						frappe.throw(_("Row {0}:Stock Item '{1}' requires a Delivery Note.Kindly create Invoice from Delivery Note.").format(i.idx, i.item_code))
