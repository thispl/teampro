
import openpyxl
import frappe
import re
from openpyxl.styles import Font, Alignment, Border, Side
from openpyxl import load_workbook
from openpyxl.utils import get_column_letter
from openpyxl.styles import GradientFill, PatternFill
from six import BytesIO, string_types
import openpyxl.styles as styles
from frappe.utils import flt, fmt_money

@frappe.whitelist()
def download(name):
    frappe.errprint('Hii')
    try:
        filename = 'IRS_Report'
        build_xlsx_response(filename, name)
    except Exception as e:
        frappe.log_error(frappe.get_traceback(), "Excel Download Error")
        raise e

def make_xlsx(filename, name, sheet_name=None, wb=None, column_widths=None):
    column_widths = column_widths or []
    if wb is None:
        wb = openpyxl.Workbook()
        wb.remove(wb.active) 
         
    ws = wb.create_sheet(sheet_name, 0)

    header_fill = PatternFill(start_color="000080", end_color="000080", fill_type="solid")
    header_font = Font(color="FFFFFF", bold=True)
    
    thin_border = Border(
        left=Side(style='thin', color='000000'),
        right=Side(style='thin', color='000000'),
        top=Side(style='thin', color='000000'),
        bottom=Side(style='thin', color='000000')
    )

    ws.column_dimensions['A'].width = 15
    ws.column_dimensions['B'].width = 12
    ws.column_dimensions['C'].width = 33
    ws.column_dimensions['D'].width = 10

    ws.append(['','','',''])
    ws.append(["ERP Implementation - Master Check List ", " ", " ", " "])
    ws.merge_cells(start_row=2, start_column=1, end_row=2, end_column=4)
    
    for col in range(1, 5):
        ws.cell(row=2, column=col).border = thin_border

    ws.append(["Type", "Department", "Document", "DEV"])
    current_header_row = ws.max_row
    for col in range(1, 5):
        cell = ws.cell(row=current_header_row, column=col)
        cell.fill = header_fill
        cell.font = header_font
        cell.border = thin_border

    data1 = get_site_data(name)
    for row in data1:
        ws.append(row)
        for col in range(1, 5):
            ws.cell(row=ws.max_row, column=col).border = thin_border

    ws.append(['','','','','']) 

    ws.append(["Type", "Department", "Document", "DEV"])
    current_header_row = ws.max_row
    for col in range(1, 5):
        cell = ws.cell(row=current_header_row, column=col)
        cell.fill = header_fill
        cell.font = header_font
        cell.border = thin_border

    data2 = get_doctype_data(name)
    if data2:
        for row_data in data2:
            current_row = ws.max_row + 1
            
            # Populate and style Type and Department
            for i in range(2):
                c = ws.cell(row=current_row, column=i+1)
                c.value = row_data[i]
                c.border = thin_border
            
            # Style Hyperlink cell
            cell_doc = ws.cell(row=current_row, column=3)
            cell_doc.value = row_data[2]
            cell_doc.hyperlink = row_data[4]
            cell_doc.style = "Hyperlink"
            cell_doc.border = thin_border # Re-apply border as style can overwrite it
            
            # Style DEV cell
            cell_dev = ws.cell(row=current_row, column=4)
            cell_dev.value = row_data[3]
            cell_dev.border = thin_border
    ws.append(['','','','','']) 

    ws.append(["Type", "Department", "Document", "DEV"])
    current_header_row = ws.max_row
    for col in range(1, 5):
        cell = ws.cell(row=current_header_row, column=col)
        cell.fill = header_fill
        cell.font = header_font
        cell.border = thin_border

    data2 = get_report_data(name)
    if data2:
        for row_data in data2:
            current_row = ws.max_row + 1
            
            # Populate and style Type and Department
            for i in range(2):
                c = ws.cell(row=current_row, column=i+1)
                c.value = row_data[i]
                c.border = thin_border
            
            # Style Hyperlink cell
            cell_doc = ws.cell(row=current_row, column=3)
            cell_doc.value = row_data[2]
            cell_doc.hyperlink = row_data[4]
            cell_doc.style = "Hyperlink"
            cell_doc.border = thin_border 
            
            # Style DEV cell
            cell_dev = ws.cell(row=current_row, column=4)
            cell_dev.value = row_data[3]
            cell_dev.border = thin_border
    xlsx_file = BytesIO()
    wb.save(xlsx_file)
    return xlsx_file

def get_site_data(name):
    project = frappe.get_doc('Project', name)
    return [
        [i.document_type, i.department, i.document_name, i.production]
        for i in project.custom_irs_table
    ]

def get_doctype_data(name):
    project = frappe.get_doc('Project', name)
    base_url = project.custom_site_url or ""
    if base_url and not base_url.endswith('/'):
        base_url += '/'
    
    data = []
    for i in project.custom_doctype_list:
        formatted_name = i.document_name.lower().replace(" ", "-")
        full_url = f"{base_url}app/{formatted_name}"
        
        # We pass the URL as the 5th item in the list
        data.append([
            i.document_type, 
            i.department, 
            i.document_name, 
            i.production,
            full_url 
        ])
    return data

def get_report_data(name):
    project = frappe.get_doc('Project', name)
    base_url = project.custom_site_url or ""
    if base_url and not base_url.endswith('/'):
        base_url += '/'
    
    data = []
    for i in project.custom_report_list:
        formatted_name = i.document_name.replace(" ", "%20")
        full_url = f"{base_url}app/query-report/{formatted_name}"
        
        data.append([
            i.document_type, 
            i.department, 
            i.document_name, 
            i.production,
            full_url 
        ])
    return data

def build_xlsx_response(filename, name):
    xlsx_file = make_xlsx(filename, name)
    
    # Set headers so the browser knows it is a file download
    frappe.response['filename'] = filename + '.xlsx'
    frappe.response['filecontent'] = xlsx_file.getvalue()
    frappe.response['type'] = 'download' # Changed from 'binary' to 'download'
