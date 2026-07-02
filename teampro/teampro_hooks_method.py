import frappe
from frappe.utils.csvutils import read_csv_content
from frappe.utils import get_first_day, get_last_day, format_datetime, get_url_to_form
from frappe.utils import cint
from frappe.utils.data import date_diff, now_datetime, nowdate, today, add_days
import datetime
from frappe import _
from frappe import _
from frappe.model.document import Document
from frappe.utils import getdate, nowdate
from frappe import throw, msgprint
import frappe
from datetime import date
from frappe import throw, _
from frappe.utils import getdate, today
today = date.today()
from frappe.model.document import Document
import datetime
import frappe,erpnext
from frappe.utils import cint
from frappe.utils import validate_email_address
import json
from frappe.utils import date_diff, add_months,today,add_days,add_years,nowdate,flt
from frappe.model.mapper import get_mapped_doc
from frappe.utils.file_manager import get_file
from frappe.utils.csvutils import UnicodeWriter, read_csv_content
import datetime
from datetime import date,datetime,timedelta
import openpyxl
from openpyxl import Workbook
import openpyxl
import xlrd
import re
from frappe.utils import today
from openpyxl import Workbook
from openpyxl.styles import PatternFill, Font, Alignment
from openpyxl.utils import get_column_letter
from datetime import datetime
from io import BytesIO
from openpyxl.styles import Font, Alignment, Border, Side
from openpyxl import load_workbook
from openpyxl.utils import get_column_letter
from openpyxl.styles import GradientFill, PatternFill
import pandas as pd
from frappe.utils import formatdate
from frappe.utils import now
from erpnext.setup.utils import get_exchange_rate
from datetime import date
from six import BytesIO, string_types
from frappe.utils import time_diff
from frappe.utils.csvutils import read_csv_content
from erpnext.buying.doctype.purchase_order.purchase_order import update_status
from frappe.utils.file_manager import get_file
from urllib.parse import urlencode
from frappe.model.rename_doc import rename_doc
from datetime import datetime
from frappe.utils import today



@frappe.whitelist()
def set_quotation(doc,method):
    frappe.db.set_value("Opportunity",doc.opportunity,"custom_quotation",doc.name)

@frappe.whitelist()
def update_custodian(doc,method):
    if doc.status == "Left":
        asset = frappe.get_all("Asset",{"custodian":doc.name},["name"])
        for i in asset:
            cust = frappe.get_doc("Asset",i.name)
            cust.custodian = ''
            cust.custodian_name = ''
            cust.department = ''
            cust.save(ignore_permissions = True)

@frappe.whitelist()
def inactive_employee(doc,method):
    if doc.status=="Active":
        if doc.relieving_date:
            throw(_("Please remove the relieving date for the Active Employee."))

@frappe.whitelist()
def issue_status(doc,method):
    if doc.service == 'IT-SW':
        if doc.issue is not None:
            if doc.status in ["Open","Working"]:
                current_status = frappe.db.get_value("Issue", doc.issue, "status")
                if current_status != "Replied":
                    frappe.db.set_value("Issue", doc.issue, {
                        "status": "Replied",
                        "task": doc.name,
                        "assigned_to": doc.completed_by,
                        "project": doc.project,
                    })
            elif doc.status == "Pending Review":
                current_status = frappe.db.get_value("Issue", doc.issue, "status")
                if current_status != "Resolved":
                    frappe.db.set_value("Issue", doc.issue, {
                        "status": "Resolved",
                        "assigned_to": doc.completed_by,
                        "project": doc.project,
                    })
            elif doc.status == "Completed":
                current_status = frappe.db.get_value("Issue", doc.issue, "status")
                if current_status != "Closed":
                    frappe.db.set_value("Issue", doc.issue, {
                        "status": "Closed",
                        "assigned_to": doc.completed_by,
                        "project": doc.project,
                    })



from frappe.utils import getdate
import frappe




# @frappe.whitelist()
# def update_dm(doc, method=None):

#     if not doc.custom_production_date:
#         return

#     if not doc.custom_dev_team:
#         return

#     if not doc.custom_sprint:
#         return
#     if doc.service != "IT-SW":
#         return
#     dm = frappe.get_all(
#         "Daily Monitor",
#         filters={
#             "docstatus": ["not in", [1, 2]],
#             "custom_dm_production_date": doc.custom_production_date,
#             "dev_team": doc.custom_dev_team,
#             "sprint": doc.custom_sprint
#         },
#         fields=["name"],
#         limit=1
#     )

#     if not dm:
#         return

#     dm_name = dm[0].name

#     # Check if task already exists in the child table — single SQL query instead of loading the full doc
#     exists = frappe.db.exists("Allocated Tasks", {"parent": dm_name, "id": doc.name})
#     if exists:
#         return

#     # Insert the child table row directly via SQL to avoid triggering Daily Monitor's
#     # validate hook (update_sprint_avl_time) which performs O(k) SQL queries per unique cb.
#     idx = frappe.db.sql(
#         "SELECT COALESCE(MAX(idx), 0) + 1 FROM `tabAllocated Tasks` WHERE parent = %s",
#         dm_name,
#     )[0][0]

#     row_name = frappe.generate_hash("Allocated Tasks", 10)
#     frappe.db.sql(
#         """INSERT INTO `tabAllocated Tasks`
#             (name, parent, parenttype, parentfield, idx, docstatus,
#              creation, modified, modified_by, owner, id, today_rt)
#            VALUES (%s, %s, %s, %s, %s, 0,
#                    NOW(), NOW(), %s, %s, %s, %s)""",
#         (row_name, dm_name, "Daily Monitor", "task_details", idx,
#          frappe.session.user, frappe.session.user, doc.name, doc.rt),
#     )



@frappe.whitelist()
def update_dm(doc, method=None):

    if not doc.custom_production_date:
        return

    if not doc.custom_dev_team:
        return

    if not doc.custom_sprint:
        return
    if doc.service != "IT-SW":
        return
    dm = frappe.get_all(
        "Daily Monitor",
        filters={
            "docstatus": ["not in", [1, 2]],
            "custom_dm_production_date": doc.custom_production_date,
            "dev_team": doc.custom_dev_team,
            "sprint": doc.custom_sprint
        },
        fields=["name"],
        limit=1
    )

    if not dm:
        return

    dm_name = dm[0].name

    # Check if task already exists in the child table — single SQL query instead of loading the full doc
    exists = frappe.db.exists(
        "Allocated Tasks",
        {"parent": dm_name, "id": doc.name}
    )

    task = frappe.get_doc("Task", doc.name)

    if exists:
        frappe.db.set_value(
            "Allocated Tasks",
            exists,
            {
                "project_name": task.project,
                "subject": task.subject,
                "cb": task.cb,
                "status": task.status,
                "revisions": task.revisions,
                "at": task.actual_time,
                "rt": task.rt,
                "et": task.expected_time,
                "priority": task.priority,
                "allocated_on": task.custom_allocated_on,
                "current_status": task.status,
                "spot_task": task.custom_spot_task,
                "remark": task.custom_remarks,
                "et_vs_at_remark": task.custom_et_vs_at_remark,
                "allocated_to": task.custom_allocated_to,
                "kt_confirmed": task.kt_confirmed,
                "is_confirmed": task.is_confirmed,
                "production_date_count": task.custom_production_date_count
            }
        )
        return

    # Insert the child table row directly via SQL to avoid triggering Daily Monitor's
    # validate hook (update_sprint_avl_time) which performs O(k) SQL queries per unique cb.
    idx = frappe.db.sql(
        "SELECT COALESCE(MAX(idx), 0) + 1 FROM `tabAllocated Tasks` WHERE parent = %s",
        dm_name,
    )[0][0]

    row_name = frappe.generate_hash("Allocated Tasks", 10)
    task = frappe.get_doc("Task", doc.name)
    frappe.db.sql("""
        INSERT INTO `tabAllocated Tasks`
        (
            name, parent, parenttype, parentfield, idx, docstatus,
            creation, modified, modified_by, owner,
            id, today_rt, project_name, subject, cb, status,
            revisions, at, rt, et, priority,
            allocated_on, current_status, spot_task,
            remark, et_vs_at_remark,
            allocated_to, cb_against_task,
            kt_confirmed, is_confirmed,
            production_date_count
        )
        VALUES (
            %s, %s, %s, %s, %s, 0,
            NOW(), NOW(), %s, %s,
            %s, %s, %s, %s, %s, %s,
            %s, %s, %s, %s, %s,
            %s, %s, %s,
            %s, %s,
            %s, %s,
            %s, %s,
            %s
        )
    """,
    (
        row_name,
        dm_name,
        "Daily Monitor",
        "task_details",
        idx,
        frappe.session.user,
        frappe.session.user,
        doc.name,
        doc.rt,
        task.project,
        task.subject,
        task.cb,
        doc.status,
        task.revisions,
        task.actual_time,
        task.rt,
        task.expected_time,
        task.priority,
        task.custom_allocated_on,
        task.status,
        task.custom_spot_task,
        task.custom_remarks,
        task.custom_et_vs_at_remark,
        task.custom_allocated_to,
        task.cb,
        task.kt_confirmed,
        task.is_confirmed,
        task.custom_production_date_count
    ))




import frappe
import requests
import json
from datetime import datetime

@frappe.whitelist()
def update_issue_wonjin(doc,method):
    if not doc.is_new():
        if doc.project == 'Wonjin_ERP_19.12.2023':
            doc_task = frappe.get_doc("Task",doc.name)
            subject = f"{doc.subject} - {doc_task.creation.strftime('%d-%m-%Y')}"
            creation = doc_task.creation.strftime('%d-%m-%Y')

            params = {
                'creation': creation,
                'name': doc.name,
                'subject': subject,
                'description': doc.description,
                'priority': doc.priority,
                'status': doc.status,
                'pr_remarks': doc.custom_taskissue_action_taken,
                'proof': doc.custom_proof_of_closure_review,
                'issue_id': doc.issue,
                'allocated_to': doc.custom_allocated_to,
            }

            frappe.enqueue(
                "teampro.teampro_hooks_method.send_wonjin_issue_update",
                queue="default",
                timeout=30,
                now=False,
                params=params,
            )

        return "No matching task found or it's new"


def send_wonjin_issue_update(params):
    import requests, json
    url = "https://erp.onegeneindia.in/api/method/onegene.www.update_issue.update_issue_from_teampro"
    headers = {
        'Content-Type': 'application/json',
        'Authorization': 'token 7503af112f2692c:812bd60c48b22ed'
    }
    try:
        response = requests.post(url, headers=headers, json=params, verify=False, timeout=15)
        response.raise_for_status()
        return response.json()
    except Exception:
        frappe.log_error(frappe.get_traceback(), "Wonjin Issue Update Failed")


@frappe.whitelist()
def task_status_complete_wonjin(task_id):
    if not frappe.db.exists("Task", task_id):
        return {"status": "error", "message": "Task not found"}

    frappe.db.set_value("Task", task_id, "status", "Completed")
    frappe.db.commit()

    return {
        "status": "success",
        "task_id": task_id
    }
    
    
@frappe.whitelist()
def update_issue_type(doc,method):
    if not doc.issue:
        return

    issue_status_map = {
        "Open": "Open",
        "Overdue": "Open",
        "Hold": "On Hold",
        "Working": "Replied",
        "Pending Review": "Resolved",
        "Client Review": "Resolved",
        "Completed": "Closed",
        "Cancelled": "Closed",
    }

    values = {
        "custom_issue_status": doc.status,
        "issue_type": doc.custom_issue_type,
    }

    mapped_status = issue_status_map.get(doc.status)
    if mapped_status:
        values["status"] = mapped_status

    frappe.db.set_value("Issue", doc.issue, values)

@frappe.whitelist()
def update_issue_typein_issue(doc,method):
    if not doc.issue:
        return
    frappe.db.set_value("Issue",doc.issue,"task",doc.name)

@frappe.whitelist()
def update_country_flag(doc, method):
    mobile_no = frappe.db.get_value("Employee", {"user_id":doc.custom_allocated_to}, ["company_mobile_number"])
    if mobile_no:
        if doc.service in ["REC-I", "REC-D"]:
            if doc.territory:
                flag_url = frappe.db.get_value("Territory", {"name": doc.territory}, ["custom_country_flag"])
                doc.custom_country_flag = flag_url
            if doc.custom_allocated_to:
                doc.custom_recruiter_contact = mobile_no

# method to update the criteria table during the task creation
@frappe.whitelist()
def update_criteria_table(doc, method):
    # pass
    if doc.service == 'REC-I':
        proj = frappe.get_doc("Project", doc.project)
        doc.set("custom_criteria_table", [])
        for row in proj.custom_criteria_table:
            doc.append("custom_criteria_table", {
                "scheduling_criteria": row.scheduling_criteria,
                "scheduling_parameter": row.scheduling_parameter
            })
        doc.save()

@frappe.whitelist()
def update_project_issue(doc,method):
    if not doc.project:
        frappe.db.set_value("Issue",doc.name,"project","Internal ERP - TEAMPRO V15")

@frappe.whitelist()
def update_issueid_wonjin(doc,method):
    # doc = frappe.get_doc("Issue",doc.name)
    if doc.raised_by == 'wonjin_corporate@onegeneindia.in':
        params = {
            'name': doc.name,
            'subject': doc.subject,
            'status': doc.status,
        }

        url = "https://erp.onegeneindia.in/api/method/onegene.www.update_issue.update_issueid_from_teampro"
        headers = {
            'Content-Type': 'application/json',
            'Authorization': 'token 7503af112f2692c:812bd60c48b22ed'
        }


        try:
            response = requests.post(url, headers=headers, json=params, verify=False)
            response.raise_for_status()  # raises exception for 4xx/5xx errors

            res = response.json()
            return res

        except requests.exceptions.RequestException as e:
            frappe.throw(f"HTTP error: {str(e)}")
        except json.JSONDecodeError:
            frappe.throw("Failed to decode JSON response from server")

    return "No matching task found or it's new"

@frappe.whitelist()
def update_service_tm(doc,method):
    doc = frappe.get_doc("Target Manager",doc.name)
    total_ct = 0
    total_ft = 0
    doc.target_child=[]
    doc.monthly_ft_allocation=[]
    if doc.service_list:
        for i in doc.service_list:
            total_ct += i.ct
            total_ft += i.ft
    doc.annual_ct = total_ct
    doc.annual_ft = total_ft
    months = ['Apr', 'May', 'Jun', 'Jul', 'Aug', 'Sep', 'Oct', 'Nov', 'Dec', 'Jan', 'Feb', 'Mar']
    month_no = {'Apr':'12','May':'11','Jun':'10','Jul':'9','Aug':'8','Sep':'7','Oct':'6','Nov':'5','Dec':'4','Jan':'3','Feb':'2','Mar':'1'}
    value =doc.annual_ct / 12
    value_ft =doc.annual_ft / 12
    for month in months:
        doc.append("target_child", {
            'month': month,
            'month_nos': month_no[month],
            'ct': value
        })
        doc.append("monthly_ft_allocation", {
            'month': month,
            'month_nos': month_no[month],
            'ft': value_ft
        })

    doc.save()
    frappe.db.commit()

@frappe.whitelist()
def create_project_completion_task(doc, method):
    if doc.service == 'IT-SW':
        task = frappe.db.exists('Task', {'project': doc.name, "subject": (
            "like", 'Project Completion Certificate')})
        if not task:
            task_id = frappe.new_doc('Task')
            task_id.update({
                "subject": "Project Completion Certificate",
                "customer": doc.customer,
                "project": doc.name,
                "service": 'IT-SW',
                "description":"Project Completion Certificate",
                "type":"Others",
                "priority":"Low",
                "custom_dev_team":"Others",
                "account_manager": doc.account_manager,
                "project_manager": doc.project_manager,

            })
            task_id.save(ignore_permissions=True)


@frappe.whitelist()
def update_color_grade(doc,method):
    if doc.name and doc.custom_color_code_grade:
        tasks=frappe.db.get_all("Task",{"project":doc.name},["name"])
        if tasks:
            for i in tasks:
                frappe.db.set_value("Task",i.name,"custom_color_code_grade",doc.custom_color_code_grade)

@frappe.whitelist()
def update_color_grade_specification(doc,method):
    if doc.name and doc.custom_color_code_grade and doc.has_value_changed("custom_color_code_grade"):
        tasks=frappe.db.get_all("Task",{"project":doc.name},["name"])
        if tasks:
            for i in tasks:
                frappe.db.set_value("Task",i.name,"custom_color_code_grade",doc.custom_color_code_grade)

@frappe.whitelist()
def update_sfp_remarks(doc,method):
    if doc.status in ['Open','Overdue','Enquiry'] and doc.customer:
        status=frappe.db.get_value('Customer',{'name':doc.customer},['disabled'])
        sfp=frappe.db.get_all("Sales Follow Up",{'party_from':'Customer','party_name':doc.customer,'service':doc.service},['active','name'])
        if status==0 and sfp:
            for s in sfp:
                if s.active==0:
                    frappe.db.set_value("Sales Follow Up",s.name,'active',True)


@frappe.whitelist()
def update_cost_center(doc, method):
    for row in doc.accounts:
        row.cost_center = doc.custom_cost_center

@frappe.whitelist()
def fetch_start_time(doc,method):
    min_time = frappe.db.sql("""select min(from_time) as min_time from `tabTimesheet Detail` where parent='%s'""" % doc.name,as_dict=1)[0]
    if not doc.start_time:
        frappe.db.set_value("Timesheet",doc.name,"start_time",min_time['min_time'])

@frappe.whitelist()
def validate_timesheet(doc,method):
    if doc.department=="IT. Development - THIS" and doc.total_hours <6 and not doc.custom_or_remarks:
        frappe.throw("Your Total hours is less than 6 hours.Kindly fill the OR Remarks")
    if doc.timesheet_summary and doc.department=="IT. Development - THIS":
        for i in doc.timesheet_summary:
            if i.status=="Working" and not i.remarks:
                frappe.throw(
                        f"Row #{i.idx}:Kindly fill the Working Remarks."
                    )
                
@frappe.whitelist()
def update_working_remarks(doc,method):
    if doc.timesheet_summary:
        for i in doc.timesheet_summary:
            if i.status=="Working" and i.remarks:
                frappe.db.set_value("Task",i.id,"custom_remarks",i.remarks)

@frappe.whitelist()
def update_batch_status(doc,method):
    if doc.get("batch"):
        batch_doc = frappe.get_doc("Batch", doc.get("batch"))
        batch_doc.batch_status = "Proposed SO"
        batch_doc.save()

@frappe.whitelist()
def update_so_priority_on_submit(doc, method):
    sales_orders = frappe.get_all(
        "Sales Order",
        filters={
            "service": "TFP",
            "status": "To Deliver and Bill",
            "docstatus":1
        },
        fields=["name", "custom_packing_on","custom_priority"]
    )

    from collections import defaultdict
    date_groups = defaultdict(list)

    for so in sales_orders:
        if so.custom_packing_on:
            date_groups[so.custom_packing_on].append(so.name)
    for priority, packing_date in enumerate(sorted(date_groups.keys()), start=1):
        for so_name in date_groups[packing_date]:
            frappe.db.set_value("Sales Order", so_name, "custom_priority", priority)

@frappe.whitelist()
def update_pi_workflow(doc,method):
    pi=frappe.get_all("Purchase Invoice",{'sales_order':doc.name},['workflow_state','name'])
    for i in pi:
        frappe.db.sql("update `tabPurchase Invoice` set workflow_state='Cancelled' where name=%s",(i.name))

@frappe.whitelist()
def update_lead_contacts_sfp(doc,method):
    if doc.party_from=="Lead" and doc.party_name:
        lead_contact = frappe.get_doc("Lead", doc.party_name)
        for i in lead_contact.lead_contacts:
            doc.append("contacts", {
                'person_name': i.person_name,
                'mobile': i.mobile,
                'is_primary': i.is_primary,
                'has_whatsapp': i.has_whatsapp,
                'email_id': i.email_id,
                'is_primaryemail': i.is_primaryemail,
                'service':i.service
            })
            doc.append("custom_contact_details", {
                'person_name': i.person_name,
                'mobile': i.mobile,
                'is_primary': i.is_primary,
                'has_whatsapp': i.has_whatsapp,
                'email_id': i.email_id,
                'is_primaryemail': i.is_primaryemail,
                'service':i.service
            })
    elif doc.party_from=="Customer" and doc.party_name:
        lead_contact = frappe.get_doc("Customer", doc.party_name)
        for i in lead_contact.customer_contact:
            doc.append("customer_contacts", {
                "person_name": i.person_name or '',
                "mobile": i.mobile or '',
                "is_primary": i.is_primary or False,
                "has_whatsapp": i.has_whatsapp or False,
                "email_id": i.email_id or '',
                "is_primaryemail": i.is_primaryemail or False,
                "service":i.service or ''
            })
    doc.save()
    frappe.db.commit()

@frappe.whitelist()
def update_spf_details_lead(doc,method):
    created_on=now_datetime()
    if doc.party_from=="Lead":
        lead=frappe.get_doc("Lead",doc.party_name)
        lead.append("custom_sfp_details", {"sfp_id": doc.name,"sfp_owner":doc.account_manager_lead_owner,"created_on":created_on,"service":doc.service})
        lead.save()
        frappe.db.commit()
    if doc.party_from=="Customer":
        customer=frappe.get_doc("Customer",doc.party_name)
        customer.append("custom_sfp_details", {"sfp_id": doc.name,"sfp_owner":doc.next_contact_by,"created_on":created_on,"service":doc.service})
        customer.save()
        frappe.db.commit()

@frappe.whitelist()
def update_lead_status(doc,method):
    if doc.status=="Lost":
        frappe.db.set_value("Lead", doc.party_name, {"disabled": 0, "docstatus": 0})

@frappe.whitelist()
def update_sfp_opportunity(doc,method):
    if doc.custom_sales_follow_up and doc.custom_quotation and doc.status=="Lost":
        frappe.db.set_value("Sales Follow Up",doc.custom_sales_follow_up,"status","Replied")

@frappe.whitelist()
def validate_permission_request(doc,method):
    start_date = get_first_day(doc.permission_date)
    end_date = get_last_day(doc.permission_date)
    permission=0
    permission_list = frappe.db.get_all("Attendance Permission",{"name": ("!=", doc.name),"employee": doc.employee,"permission_date": ("between", [start_date, end_date]),"docstatus":("!=",2)},["*"])
    if permission_list:
        for i in permission_list:
            permission+=int(i.total_time)
            if permission>=2:
                frappe.throw("Only 2 hours Permission is allowed for the month")
            elif permission<2:
                permission+=int(doc.total_time)
                if permission>2:
                    frappe.throw("Already applied for 1 hour permission.You are only  allow to apply additionaly 1 hour")

@frappe.whitelist()
def update_permission_req_in_att(doc,method):
    attendance=frappe.get_doc("Attendance",{"attendance_date":doc.permission_date,"docstatus":("!=",2),"employee":doc.employee})
    hours=0
    if attendance:
        attendance.custom_attendance_permission=doc.name
        if attendance.bt_difference:
            diff=(attendance.bt_difference)
            hours = diff+int(doc.total_time)
        if hours>=8:
            attendance.status="Present"
        elif hours>=4 and hours < 8:
            attendance.status="Half Day"
        else:
            attendance.status="Absent"
    attendance.save()
    frappe.db.commit()

@frappe.whitelist()
def update_permission_req_in_att_cancel(doc,method):
    attendance=frappe.get_doc("Attendance",{"attendance_date":doc.permission_date,"docstatus":("!=",2),"employee":doc.employee})
    hours=0
    if attendance:
        attendance.custom_attendance_permission=""
        if attendance.bt_difference:
            diff=(attendance.bt_difference)
            hours = diff
        if hours>=8:
            attendance.status="Present"
        elif hours>=4 and hours < 8:
            attendance.status="Half Day"
        else:
            attendance.status="Absent"
    attendance.save()
    frappe.db.commit()

@frappe.whitelist()
def update_wh_att(doc,method):
    if frappe.db.exists("Attendance",{"attendance_request":doc.name,"employee":doc.employee,'docstatus':['!=',2]}):
        att=frappe.db.get_value("Attendance",{"attendance_request":doc.name,"employee":doc.employee,'docstatus':['!=',2]},['name'])
        if att:
            frappe.db.set_value("Attendance",att,'attendance_request','')

@frappe.whitelist()
def update_workflow_state(doc,method):
    if doc.workflow_state:
        frappe.db.sql("""update `tabPurchase Invoice` set custom_status = %s where name = %s""",(doc.workflow_state,doc.name))

@frappe.whitelist()
def calc_cost_prize(doc,method):
    if doc.workflow_state!="Approved":
        for f in doc.items:
            tfp_item = frappe.db.sql("""select tfp from `tabItem` where name = '%s' """%(f.item_code),as_dict=1)[0]
            tfp = tfp_item['tfp']
            if tfp == 1:
                price_list = frappe.db.sql("""select price_list_rate from `tabItem Price` where price_list = 'Cost Price TFP' and item_code = '%s' """%(f.item_code),as_dict=1)
                for p in price_list:
                    if f.uom == 'Gram':
                        item_price = (p.price_list_rate / 1000)
                        item_rate = round((item_price),2)
                        if f.rate > item_rate:
                            frappe.msgprint(_(' %s Rate is Greater than Cost Price')%(f.item_name))
                    elif f.uom == 'Kg':
                        if f.rate > p.price_list_rate:
                            frappe.msgprint(_(' %s Rate is Greater than Cost Price')%(f.item_name))

@frappe.whitelist()
def update_ordered_qty(doc,method):
    material_request = frappe.db.get_value("Purchase Order Item", {"parent": doc.name}, "material_request")
    if frappe.db.exists("Material Request", material_request):
        for i in doc.items:
            mr_name = frappe.db.get_value("Material Request Clubbed Item",{'parent':i.material_request,'item_code':i.item_code},['name'])
            if mr_name:
                frappe.db.set_value("Material Request Clubbed Item",mr_name,'po_qty',i.qty)

@frappe.whitelist()
def update_material_request_status_on_submit(doc,method):
    material_request = frappe.db.get_value("Purchase Order Item", {"parent": doc.name}, "material_request")
    if frappe.db.exists("Material Request", material_request):
        doc = frappe.get_doc("Material Request", material_request)
        row_count = 0
        conditions_satisfied_count = 0
        for row in doc.custom_merged_items:
            row_count += 1
            if row.po_qty >= row.purchase_qty:
                conditions_satisfied_count += 1
        if row_count > 0 and row_count == conditions_satisfied_count:
            frappe.db.set_value("Material Request", material_request, "status", "Ordered")

@frappe.whitelist()
def update_ordered_qty_on_cancel(doc,method):
    material_request = frappe.db.get_value("Purchase Order Item", {"parent": doc.name}, "material_request")
    if frappe.db.exists("Material Request", material_request):
        for i in doc.items:
            mr_name = frappe.db.get_value("Material Request Clubbed Item",{'parent':i.material_request,'item_code':i.item_code},['name'])
            if mr_name:
                po_qty = frappe.db.get_value("Material Request Clubbed Item",{'parent':i.material_request,'item_code':i.item_code},['po_qty'])
                if po_qty and po_qty > 0:
                    frappe.db.set_value("Material Request Clubbed Item",mr_name,'po_qty',po_qty - i.qty)
        doc = frappe.get_doc("Material Request", material_request)
        row_count = 0
        conditions_satisfied_count = 0
        for row in doc.custom_merged_items:
            row_count += 1
            if row.po_qty >= row.purchase_qty:
                conditions_satisfied_count += 1
        if row_count > 0 and  row_count == conditions_satisfied_count:
            frappe.db.set_value("Material Request", material_request, "status", "Pending")

@frappe.whitelist()
def update_employer_pf(doc, method):
    if doc.earnings:
        for i in doc.earnings:
            if i.salary_component=='Provident Fund-Employer':
                doc.custom_employer_pf=i.amount

@frappe.whitelist()
def update_month_cycle(doc,method):
    from_date_obj = datetime.strptime(doc.start_date, "%Y-%m-%d")
    month_year = from_date_obj.strftime("%b %Y")
    frappe.db.set_value("Appraisal Cycle",doc.name,"custom_cycle_month",month_year)

import frappe
from frappe.utils import getdate
import calendar
from datetime import datetime

@frappe.whitelist()
def update_ep_nc_appraisal(doc, method):
    month_year_str = doc.custom_appraisal_cycle_month.strip()
    month_year_dt = datetime.strptime(month_year_str, "%b %Y")

    first_day = month_year_dt.replace(day=1)
    last_day = month_year_dt.replace(
        day=calendar.monthrange(month_year_dt.year, month_year_dt.month)[1]
    )
    ep_nc_records = frappe.db.sql("""
        SELECT name, action, total, total_nc,reason_of_ep
        FROM `tabEnergy Point And Non Conformity`
        WHERE emp = %s
        AND creation BETWEEN %s AND %s AND docstatus=1
    """, (doc.employee, first_day, last_day), as_dict=True)

    doc.custom_details = []
    total_sum = 0
    for row in ep_nc_records:
        ep_score = row.total or 0
        nc_score = -(row.total_nc or 0)
        doc.append("custom_details", {
            "ep__nc": row.name,
            "reason":row.reason_of_ep,
            "action": row.action,
            "ep_score": ep_score,
            "nc_score": nc_score
        })
        total_sum += ep_score + nc_score
    doc.custom_total_ens =total_sum

@frappe.whitelist()
def update_grade(doc,method):
    if doc.total_score:
        if doc.total_score==5:
            doc.custom_grade='A+'
        elif 4 <= doc.total_score < 5:
            doc.custom_grade='A'
        elif 3.5 <= doc.total_score < 4:
            doc.custom_grade='B+'
        elif 3 <= doc.total_score < 3.5:
            doc.custom_grade='B'
        elif 2 <= doc.total_score < 3:
            doc.custom_grade='C'
        elif 1 <= doc.total_score < 2:
            doc.custom_grade='D'
        else:
            doc.custom_grade='E'
    else:
        doc.custom_grade=''

@frappe.whitelist()
def validate_reviewer_remark(doc,method):
    if not doc.custom_reviewer_remark:
        frappe.throw("Kindly enter the Reviewer Remark before submit")

@frappe.whitelist()
def update_lead_as_qualified(doc,method):
    if doc.lead_name:
        if frappe.db.exists("Lead",doc.lead_name):
            lead=frappe.get_doc("Lead",doc.lead_name)
            lead.status='Converted'
            lead.save(ignore_permissions=True)

@frappe.whitelist()
def update_spf_status(doc,method):
    if doc.lead_name:
        sfp_name = frappe.db.get_all("Sales Follow Up", {"party_name": doc.lead_name}, "name")
        if sfp_name:
            for i in sfp_name:
                sfp_doc = frappe.get_doc("Sales Follow Up", i.name)
                sfp_doc.status = "Converted"
                sfp_doc.party_from="Customer"
                sfp_doc.party_name=doc.name
                sfp_doc.save()
                sfp_doc.reload()
                frappe.db.commit()

@frappe.whitelist()
def update_project_dates(doc, method):
    if doc.custom_sla_details:
        for row in doc.custom_sla_details:
            if row.project and row.sla_from_date and row.sla_to_date:
                frappe.db.set_value("Project", row.project, {
                    "expected_start_date": row.sla_from_date,
                    "expected_end_date": row.sla_to_date
                })
                frappe.db.commit()

import frappe
from frappe.model.naming import make_autoname
from frappe.utils import now_datetime

@frappe.whitelist()
def set_customer_id(doc, method):
    if doc.is_new():
        if not doc.customer_id:
            last_customer = frappe.db.sql("""
                SELECT customer_id FROM `tabCustomer`
                WHERE customer_id REGEXP '^CUST-[0-9]{6}$'
                ORDER BY CAST(SUBSTRING(customer_id, 6) AS UNSIGNED) DESC
                LIMIT 1
            """, as_dict=True)

            if last_customer:
                last_id_num = int(last_customer[0]["customer_id"].split("-")[1])
                new_id_num = last_id_num + 1
            else:
                new_id_num = 1
            new_customer_id = f"CUST-{new_id_num:06d}"
            doc.customer_id = new_customer_id

@frappe.whitelist()
def update_check_existing_lead(doc,method):
    if doc.custom_check_existing:
        leads=frappe.get_doc("Existing Leads",doc.custom_check_existing)
        leads.lead_id=doc.name
        leads.save()
        leads.reload()
        frappe.db.commit()

@frappe.whitelist()
def update_task_count(doc, method):
    if doc.task:
        task_status = frappe.db.get_value('Task', doc.task, 'status')
        if task_status in("Open","Working","Overdue","Pending Review"):
            submit_spoc = frappe.db.count(
                'Candidate', {'task': doc.task, 'pending_for': 'Submit(SPOC)'}) or 0
            submit_client = frappe.db.count(
                'Candidate', {'task': doc.task, 'pending_for': 'Submitted(Client)'}) or 0
            interviewed = frappe.db.count(
                'Candidate', {'task': doc.task, 'pending_for': 'Interviewed'}) or 0
            frappe.db.set_value('Task', doc.task, 'fp',(submit_spoc + interviewed + submit_client))

        psl = frappe.db.count('Candidate', {'task': doc.task, 'pending_for': (
            'in', ('Client Offered', 'Proposed PSL'))}) or 0
        shortlisted = frappe.db.count(
            'Candidate', {'task': doc.task, 'pending_for':'Shortlisted'}) or 0
        linedup = frappe.db.count(
            'Candidate', {'task': doc.task, 'pending_for':('in', ('Linedup','Linedup Confirmed'))}) or 0
        result_pending =frappe.db.count('Candidate',{'task':doc.task,'pending_for':'Result Pending'}) or 0
        frappe.db.set_value('Task', doc.task, 'psl', psl)
        frappe.db.set_value('Task',doc.task,'custom_rp',result_pending)
        frappe.db.set_value('Task', doc.task, 'sl', shortlisted)
        frappe.db.set_value('Task', doc.task, 'custom_lp',linedup)


        if task_status in ('Completed', 'Cancelled'):
            frappe.db.set_value('Task', doc.task, 'sp', 0)
            frappe.db.set_value('Task', doc.task, 'fp',0)

        else:
            submit_spoc_list = 0
            submit_client_list = 0
            interviewed_list = 0
            if task_status in("Open","Working","Overdue","Pending Review"):
                submit_spoc_list = frappe.db.count(
                    'Candidate', {'task': doc.task, 'pending_for': 'Submit(SPOC)'}) or 0
                submit_client_list = frappe.db.count(
                    'Candidate', {'task': doc.task, 'pending_for': 'Submitted(Client)'}) or 0
                interviewed_list = frappe.db.count(
                    'Candidate', {'task': doc.task, 'pending_for': 'Interviewed'}) or 0
            vac = frappe.db.get_value('Task', doc.task, 'vac')
            prop = frappe.db.get_value('Task', doc.task, 'prop')
            pps = (vac - psl) * prop - (submit_spoc_list + submit_client_list+
                                        interviewed_list + shortlisted +linedup)
            frappe.db.set_value('Task', doc.task, 'sp', pps)

@frappe.whitelist()
def update_sams_by(doc,method):
    if doc.sa_agent:
        frappe.db.set_value("Candidate",doc.name,'custom_sourced_by',"SAMS")
    else:
        frappe.db.set_value("Candidate",doc.name,'custom_sourced_by',"Normal")
       
@frappe.whitelist()     
def on_creation_of_psl_mail(doc,method):
    acc=''
    acc_manager=''
    spoc=''
    if doc.candidate:
        acc=frappe.db.get_value("Candidate",{"name":doc.candidate},['task'])
        acc_manager=frappe.db.get_value("Task",{'name':acc},['account_manager'])
        spoc=frappe.db.get_value("Task",{'name':acc},['spoc'])
    frappe.sendmail(
        # recipients=["divya.p@groupteampro.com"],
        recipients=["sangeetha.s@groupteampro.com","dc@groupteampro.com",acc_manager,spoc],
        subject = "New PSL Created -  %s" % nowdate(),
        message="""   
        Dear Sir/Mam,<br>
        <p><b>Closure ID: </b>%s  <b> Status :</b>%s  <b> Customer Name :</b>%s -a new PSL added in Closure for your further action.</p><br>
        
            Thanks & Regards<br>TEAM ERP<br>"This email has been automatically generated. Please do not reply"
    """ % (doc.name,doc.status,doc.customer)
        ) 
    
@frappe.whitelist()
def update_score_in_epnc_review(doc, method):
    from datetime import datetime
    creation_date_result = frappe.db.sql(
        """
        SELECT date(creation) as creation from `tabEnergy Point And Non Conformity` where creation=%s AND name=%s
        """, (doc.creation,doc.name), as_dict=True)
    creation_date = creation_date_result[0]['creation'] if creation_date_result else None
    first_date=get_first_day(creation_date)
    documents=frappe.db.get_all("Monthly EP NC Review",{"start_date":first_date,'employee':doc.emp},["name"])
    for i in documents:
        doc=frappe.get_doc("Monthly EP NC Review",i)
        ep = frappe.db.sql(
            """
            SELECT sum(total) as total from `tabEnergy Point And Non Conformity` where emp=%s AND action='Energy Point(EP)' AND docstatus=1 AND date(creation) BETWEEN %s AND %s
            """, (doc.employee, doc.start_date, doc.end_date), as_dict=True)
        nc = frappe.db.sql(
            """
            SELECT sum(total_nc) as total_nc from `tabEnergy Point And Non Conformity` where emp=%s AND action='Non Conformity(NC)' AND docstatus=1 AND date(creation) BETWEEN %s AND %s
            """, (doc.employee, doc.start_date, doc.end_date), as_dict=True)
        doc.total_ep = ep[0]['total'] if ep and ep[0]['total'] is not None else 0
        doc.total_nc = nc[0]['total_nc'] if nc and nc[0]['total_nc'] is not None else 0
        total = 100 - doc.total_nc + doc.total_ep
        doc.total_score = total
        doc.save(ignore_permissions=True)

    frappe.db.commit()

@frappe.whitelist()
def batch_status_update_in_batch(doc,method):
    cases=frappe.db.get_all("Case",{"batch":doc.batch},["*"])
    case_sts=[]
    tat_status=[]
    batch_status=''
    for i in cases:
        case_sts.append(i.case_status)
        tat_status.append(i.tat_monitor)
        if any(status == "Draft" for status in case_sts):
            if any(tat=="In TAT" for tat in tat_status):
                batch_status="Open"
            elif any(tat=="Out TAT" for tat in tat_status):
                batch_status="Overdue"
            else:
                batch_status="Open"
        elif  any(status == "Entry-Insuff" for status in case_sts):
            if any(tat=="Out TAT" for tat in tat_status):
                batch_status="Overdue with Insuff"
            elif any(tat=="In TAT" for tat in tat_status):
                batch_status="Open with Insuff"
        elif any(status == "Execution-Insuff" for status in case_sts):
            if any(tat=="In TAT" for tat in tat_status):
                batch_status="Open with Insuff"
            elif any(tat=="Out TAT" for tat in tat_status):
                batch_status="Overdue with Insuff"
        elif any(status == "Entry-QC" for status in case_sts):
            if any(tat=="In TAT" for tat in tat_status):
                batch_status="Open"
            elif any(tat=="Out TAT" for tat in tat_status):
                batch_status="Overdue"
        elif any(status == "Entry Completed" for status in case_sts):
            if any(tat=="Out TAT" for tat in tat_status):
                batch_status="Overdue"
            elif any(tat=="In TAT" for tat in tat_status):
                batch_status="Open"
        elif any(status == "Execution" for status in case_sts):
            if any(tat=="Out TAT" for tat in tat_status):
                batch_status="Overdue"
            elif any(tat=="In TAT" for tat in tat_status):
                batch_status="Open"
        elif any(status == "Case Report Completed" or status == "Billed" or status == "Case Completed" or status == "To be Billed" or status == "SO Created" or status == "Drop" or status=="Generate Report" for status in case_sts):
            batch_status="Completed"
    frappe.db.set_value("Batch",doc.batch,"batch_status",batch_status)

import frappe

@frappe.whitelist()
def update_case_status_in_batch(doc, method):
    batch = frappe.get_doc("Batch", {"name": doc.batch})
    batch.casewise_status = []
    cases = frappe.db.get_all("Case", filters={"batch": doc.batch}, fields=["name", "case_status"])
    completed = 0
    insuff = 0
    pending = 0
    drop=0
    for case in cases:
        batch.append("casewise_status", {
            "case_id": case["name"],
            "case_status": case["case_status"]
        })
        if case["case_status"] in ["Case Completed","To be Billed","SO Created","Case Report Completed","Generate Report","Billed"]:
            completed += 1
        elif case["case_status"] in ["Entry-Insuff","Execution-Insuff"]:
            insuff += 1
        elif case["case_status"] in ["Drop"]:
            drop+=1
        else:
            pending += 1
        batch.comp=completed
        batch.insuff=insuff
        batch.pending=pending
        batch.custom_drop=drop
    batch.save()
    frappe.db.commit()

@frappe.whitelist()
def auto_submit_stock(doc,method):
    if doc.stock_entry_type=="Material Transfer" and doc.company=="TEAMPRO Food Products" and doc.custom_vm_stock_register:
        if frappe.db.exists("Stock Entry",{"custom_vm_stock_register":doc.custom_vm_stock_register,"docstatus":0,"stock_entry_type":"Material Issue"}):
            stock_entry=frappe.get_doc("Stock Entry",{"custom_vm_stock_register":doc.custom_vm_stock_register,"docstatus":0,"stock_entry_type":"Material Issue"})
            stock_entry.submit()


@frappe.whitelist()
def update_stock_against_vm(doc, method):
    stock_entries = frappe.get_all(
        "Stock Entry",
        filters={"custom_vm_stock_register": doc.name},
        fields=["name", "docstatus"]
    )
    for se in stock_entries:
        stock_entry = frappe.get_doc("Stock Entry", se.name)
        if se.docstatus == 1:
            stock_entry.cancel()
        elif se.docstatus == 0:
            stock_entry.delete()


import frappe
@frappe.whitelist()
def create_packing_issue_stock_entry(doc, method):

    if not (
        doc.company == "TEAMPRO Food Products"
        and doc.stock_entry_type == "Material Transfer"
        and doc.custom_internal_transfer_type == "Retail"
    ):
        return

    issue_items = []

    for d in doc.items:

        if d.custom_primary_packing_cover and (d.custom_covers or 0) > 0:
            issue_items.append({
                "item_code": d.custom_primary_packing_cover,
                "qty": d.custom_covers,
                "uom": d.custom_primary_uom,
                "s_warehouse": "Stores - TFP"
            })

        if d.custom_secondary_packing_bag and (d.custom_bag or 0) > 0:
            issue_items.append({
                "item_code": d.custom_secondary_packing_bag,
                "qty": d.custom_bag,
                "uom": d.custom_secondary_uom,
                "s_warehouse": "Stores - TFP"
            })

        if d.custom_tertiary_packingbox and (d.custom_box or 0) > 0:
            issue_items.append({
                "item_code": d.custom_tertiary_packingbox,
                "qty": d.custom_box,
                "uom": d.custom_tertiary_uom,
                "s_warehouse": "Stores - TFP"
            })

    if not issue_items:
        return

    se = frappe.new_doc("Stock Entry")
    se.stock_entry_type = "Material Issue"
    se.company = doc.company
    se.custom_reference_stock_entry = doc.name

    for item in issue_items:
        se.append("items", {
            "item_code": item["item_code"],
            "qty": item["qty"],
            "uom": item["uom"],
            "s_warehouse": item["s_warehouse"]
        })

    se.insert(ignore_permissions=True)
    se.submit()

import frappe
@frappe.whitelist()
def cancel_packing_issue_stock_entry(doc, method):
    if not (
        doc.company == "TEAMPRO Food Products"
        and doc.stock_entry_type == "Material Issue"
    ):
        return
    stock_entries = frappe.get_all(
        "Stock Entry",
        filters={
            "custom_reference_stock_entry": doc.name,
            "stock_entry_type": "Material Issue"
        },
        fields=["name", "docstatus"]
    )
    for se in stock_entries:

        se_doc = frappe.get_doc("Stock Entry", se.name)

        if se_doc.docstatus == 1:
            se_doc.cancel()

        elif se_doc.docstatus == 0:
            frappe.delete_doc("Stock Entry", se.name)

@frappe.whitelist()
def update_project_image(doc,method):
    if doc.custom_location_imagejob_card:
        task = frappe.db.get_all("Task",{"project":doc.name},["name"])
        if task:
            for i in task:
                if not frappe.db.exists("Task",{"name":i.name,"custom_customer_location_image":doc.custom_location_imagejob_card}):
                    frappe.db.set_value("Task",i.name,"custom_customer_location_image",doc.custom_location_imagejob_card)



@frappe.whitelist()
def update_company_by_employee(doc,method):
    if doc.employee:
        doc.company = frappe.db.get_value("Employee", doc.employee, "company")
    




import frappe
from datetime import datetime, timedelta

@frappe.whitelist()
def submit_previous_month_attendance():

    today = datetime.today()

    first_day_this_month = today.replace(day=1)
    last_day_prev_month = first_day_this_month - timedelta(days=1)
    first_day_prev_month = last_day_prev_month.replace(day=1)

    attendance_list = frappe.get_all(
        "Attendance",
        filters={
            "attendance_date": ["between", [first_day_prev_month, last_day_prev_month]],
            "docstatus": 0   
        },
        fields=["name"]
    )

    count = 0

    for att in attendance_list:
        try:
            doc = frappe.get_doc("Attendance", att.name)
            doc.submit()
            count += 1
        except Exception as e:
            frappe.log_error(frappe.get_traceback(), "Attendance Submit Error")

    return f"{count} Attendance records submitted successfully"


@frappe.whitelist()
def submit_all_attendance ():
    job = frappe.db.exists('Scheduled Job Type','submit_previous_month_attendance')
    if not job:
        task = frappe.new_doc("Scheduled Job Type")
        task.update({
            "method": 'teampro.teampro_hooks_method.submit_previous_month_attendance',
            "frequency": 'Monthly',
            "cron_format": '0 0 3 * *'
        })
        task.save(ignore_permissions=True)





def old_sprint_alert(doc,method):
    if doc.custom_sprint and doc.custom_dev_team:

        sprint = frappe.get_value(
            "Sprint",
            {
                "sprint_id": doc.custom_sprint,
                "team": doc.custom_dev_team
            },
            ["from_date", "to_date"],
            as_dict=True
        )

        if sprint:
            today = getdate(nowdate())

            if not (sprint.from_date <= today <= sprint.to_date):
                frappe.msgprint(
                    "In the task, the selected sprint does not match the latest sprint. Kindly check."
                )



@frappe.whitelist()
def create_user_notification(doc,method):
	if doc.service in ['REC-I','REC-D']:
		un=frappe.new_doc("User Notifications")
		un.subject=doc.subject + " - JOB ALERT🔥"
		un.content="A new Job "+doc.subject+" has been added"
		un.save(ignore_permissions=True)


import frappe
from frappe.utils import nowdate, getdate

@frappe.whitelist()
def get_user_target():
	user = frappe.session.user

	employee = frappe.db.get_value("Employee", {"user_id": user}, "name")
	if not employee:
		return {"has_target": False}

	target_name = frappe.db.get_value("Target Manager", {"employee": employee}, "name")
	if not target_name:
		return {"has_target": False}

	target = frappe.get_doc("Target Manager", target_name)

	today = getdate(nowdate())

	if target.custom_year_start_date and target.custom_year_end_date:
		if not (target.custom_year_start_date <= today <= target.custom_year_end_date):
			return {"has_target": False}

	return {
		"has_target": True,
		"custom_total_achieved_point": target.custom_total_achieved_point,
		"custom_total_target_point": target.custom_total_target_point,
		"custom_sr": target.custom_sr
	}


import frappe
from frappe.utils import getdate, nowdate
from datetime import datetime

@frappe.whitelist()
def get_user_target_month():
	user = frappe.session.user

	# Get employee linked to this user
	employee = frappe.db.get_value("Employee", {"user_id": user}, "name")
	if not employee:
		return {"has_target": False}

	today = getdate(nowdate())
	current_month = datetime.today().strftime("%b").lower()

	# Get the Target Manager where today is within start and end date
	target_name = frappe.db.get_value(
		"Target Manager",
		filters={
			"employee": employee,
			"custom_year_start_date": ("<=", today),
			"custom_year_end_date": (">=", today)
		},
		fieldname="name"
	)

	if not target_name:
		# No target for current date range
		return {"has_target": False}

	target = frappe.get_doc("Target Manager", target_name)

	achieved = 0
	total = 0
	percent = 0

	# Calculate points only for current month
	for row in target.target_child:
		if row.month and row.month.strip().lower() == current_month:
			achieved += row.achieved_point or 0
			total += row.cr_ct_point or 0

	if total:
		percent = (achieved / total) * 100

	return {
		"has_target": True if total else False,
		"month_achieved": achieved,
		"month_target": total,
		"month_percent": percent
	}

@frappe.whitelist()
def emp_short_code_check(doc, method):
    if not doc.short_code:
        return

    
    if doc.is_new():
        if frappe.db.exists("Employee", {"short_code": doc.short_code}):
            frappe.throw("Short Code must be unique")
        return

    
    old_doc = doc.get_doc_before_save()
    if not old_doc:
        return

    if old_doc.short_code != doc.short_code:
        if frappe.db.exists(
            "Employee",
            {
                "short_code": doc.short_code,
                "name": ["!=", doc.name],
            }
        ):
            frappe.throw("Short Code must be unique")

@frappe.whitelist()
def update_cb_bulk(doc,method):
    if doc.service=="IT-SW":
        if not doc.custom_allocated_to:
            return
        employee = frappe.db.get_value(
            "Employee",
            {"user_id": doc.custom_allocated_to},
            ["name", "short_code"],
            as_dict=True
        )

        if not employee:
            return
        if employee.short_code:
            doc.cb = employee.short_code

from frappe import enqueue

@frappe.whitelist()
def send_notification(doc, method):
	device_ids = user_id()
	if device_ids:
		for device_id in device_ids:
			enqueue(
				process_notification,
				queue="default",
				now=False,
				device_id=device_id,
				notification=doc,
			)


@frappe.whitelist()
def process_notification(device_id, notification):
	message = notification.subject
	title = notification.content
	url = "https://fcm.googleapis.com/v1/projects/jobpro-f8bef/messages:send"
	body = {
		"message":{
			"token": device_id.custom_device_id,
			"notification": {"title": message, "body": title},
		}
	}

	server_key = _get_access_token()
	auth = f"Bearer {server_key}"
	req = requests.post(
		url=url,
		data=json.dumps(body),
		headers={
			"Authorization": auth,
			"Content-Type": "application/json",
		},
	)


@frappe.whitelist()
def user_id():
	# user_email = doc.for_user
	user_device_id = frappe.get_all(
		"Candidate", filters={"custom_device_id": ('!=','')}, fields=["custom_device_id"]
	)
	return user_device_id


import google.auth.transport.requests
from google.oauth2 import service_account
import os
SCOPES = ["https://www.googleapis.com/auth/cloud-platform"]
@frappe.whitelist()
def _get_access_token():
	json_path = os.getenv('SERVICE_ACCOUNT_JSON', '/home/frappe/teampro-bench/apps/teampro/teampro/jobpro-f8bef-f2152785d551.json')
	credentials = service_account.Credentials.from_service_account_file(json_path, scopes=SCOPES)
	request = google.auth.transport.requests.Request()
	credentials.refresh(request)
	return credentials.token
