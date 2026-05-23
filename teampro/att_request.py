import frappe
from frappe.utils.csvutils import read_csv_content
from frappe.utils import get_first_day, get_last_day, format_datetime, get_url_to_form
from frappe.utils import cint
from frappe.utils.data import date_diff, now_datetime, nowdate, today, add_days
import datetime
from frappe import _
from frappe import _
from frappe.model.document import Document
from frappe.utils import getdate, nowdate
from frappe import throw, msgprint
import frappe
from datetime import date
from frappe import throw, _
from frappe.utils import getdate, today
today = date.today()
from frappe.model.document import Document
import datetime
import frappe,erpnext
from frappe.utils import cint
from frappe.utils import validate_email_address
import json
from frappe.utils import date_diff, add_months,today,add_days,add_years,nowdate,flt
from frappe.model.mapper import get_mapped_doc
from frappe.utils.file_manager import get_file
from frappe.utils.csvutils import UnicodeWriter, read_csv_content
import datetime
from datetime import date,datetime,timedelta
import openpyxl
from openpyxl import Workbook
import openpyxl
import xlrd
import re
from frappe.utils import today
from openpyxl import Workbook
from openpyxl.styles import PatternFill, Font, Alignment
from openpyxl.utils import get_column_letter
from datetime import datetime
from io import BytesIO
from openpyxl.styles import Font, Alignment, Border, Side
from openpyxl import load_workbook
from openpyxl.utils import get_column_letter
from openpyxl.styles import GradientFill, PatternFill
import pandas as pd
from frappe.utils import formatdate
from frappe.utils import now
from erpnext.setup.utils import get_exchange_rate
from datetime import date
from six import BytesIO, string_types
from frappe.utils import time_diff
from frappe.utils.csvutils import read_csv_content
from erpnext.buying.doctype.purchase_order.purchase_order import update_status
from frappe.utils.file_manager import get_file
from urllib.parse import urlencode
from frappe.model.rename_doc import rename_doc
from datetime import datetime
from frappe.utils import today

import frappe
from frappe.utils import getdate

import frappe
from frappe.utils import getdate, date_diff, add_days
from frappe.utils import get_datetime

@frappe.whitelist()
def validate_att_working_day(doc, method):
    holiday_list = frappe.db.get_value("Employee", doc.employee, "holiday_list")
    if not holiday_list:
        frappe.throw("Holiday List not found for this Employee")
    days = date_diff(doc.to_date, doc.from_date)
    for i in range(days + 1):
        check_date = add_days(doc.from_date, i)
        is_holiday = frappe.db.exists("Holiday", {"parent": holiday_list,"holiday_date": check_date})
        formatted_date = formatdate(check_date, "dd-MM-yyyy")
        if doc.reason == "On Duty Working Day" and is_holiday:
            frappe.throw(f"{formatted_date} is a Holiday. Cannot apply 'On Duty Working Day'")
        if doc.reason == "Comp Off_ On Duty Holiday" and not is_holiday:
            frappe.throw(f"{formatted_date} is not a Holiday. Comp Off allowed only on holidays")
    if doc.reason == "Mispunch":
        if not (doc.custom_in_time or doc.custom_out_time):
            frappe.throw("No punches found for this date.")
        if doc.custom_in_time or doc.custom_out_time:
            formatted_date = formatdate(doc.from_date, "dd-MM-yyyy")
            existing = frappe.db.exists("Attendance Request",{"employee": doc.employee,"from_date": doc.from_date,"reason": "Mispunch","docstatus": ["!=", 2]})
            if existing and existing != doc.name:
                frappe.throw(f"Mispunch already applied for {formatted_date}")
    if doc.reason == "Permission":
        start_date = get_first_day(doc.from_date)
        end_date = get_last_day(doc.from_date)
        permission = 0
        permission_list = frappe.db.get_all("Attendance Request",{"name": ("!=", doc.name),"employee": doc.employee,"from_date": ("between", [start_date, end_date]),"reason": "Permission","docstatus": ("!=", 2)},["custom_total_time"])
        if permission_list:
            for i in permission_list:
                permission+=int(i.custom_total_time)
                if permission>=2:
                    frappe.throw("Only 2 hours Permission is allowed for the month")
                elif permission<2:
                    permission+=int(doc.custom_total_time)
                    if permission>2:
                        frappe.throw("Already applied for 1 hour permission.You are only  allow to apply additionaly 1 hour")
        
        if doc.custom_permission_session == "First Half" and doc.custom_morning_time:
            
            allowed_time = get_allowed_time(doc.custom_morning_time)
            if doc.custom_in_time and allowed_time:
                # in_time = doc.custom_in_time.split(" ")[1][:5] 
                in_time_obj = get_datetime(doc.custom_in_time)
                in_time = in_time_obj.strftime("%H:%M")
                if not in_time:
                    frappe.throw("In Time is required for First Half Permission")
                if in_time > allowed_time:
                    frappe.throw(
                        f"In Time {in_time} exceeds allowed time {allowed_time}"
                    )
        if doc.custom_permission_session == "Second Half" and doc.custom_evening_time:
            allowed_time = get_allowed_time(doc.custom_evening_time)
            if doc.custom_out_time and allowed_time:
                # out_time = doc.custom_out_time.split(" ")[1][:5]
                out_time_obj = get_datetime(doc.custom_out_time)
                out_time = out_time_obj.strftime("%H:%M")
                if not out_time:
                    frappe.throw("Out Time is required for Second Half Permission")
                if out_time < allowed_time:
                    frappe.throw(
                        f"Out Time {out_time} should be after {allowed_time}"
                    )

@frappe.whitelist()
def update_permission_req_in_att_submission(doc,method):
    if doc.workflow_state != "Approved":
        return
    if doc.reason == "Permission":
        attendance=frappe.get_doc("Attendance",{"attendance_date":doc.from_date,"docstatus":("!=",2),"employee":doc.employee})
        hours=0
        if attendance:
            attendance.attendance_request=doc.name
            attendance.custom_session=doc.custom_permission_session
            if attendance.bt_difference:
                diff=(attendance.bt_difference)
                hours = diff+int(doc.custom_total_time)
            if hours>=8:
                attendance.status="Present"
            elif hours>=4 and hours < 8:
                attendance.status="Half Day"
            else:
                attendance.status="Absent"
        attendance.save()
        frappe.db.commit()

@frappe.whitelist()
def update_perm_req_in_att_cancel(doc,method):
    attendance=frappe.get_doc("Attendance",{"attendance_date":doc.from_date,"docstatus":("!=",2),"employee":doc.employee})
    hours=0
    if attendance:
        attendance.attendance_request=""
        attendance.custom_session=""
        if attendance.bt_difference:
            diff=(attendance.bt_difference)
            hours = diff
        if hours>=8:
            attendance.status="Present"
        elif hours>=4 and hours < 8:
            attendance.status="Half Day"
        else:
            attendance.status="Absent"
    attendance.save()
    frappe.db.commit()

@frappe.whitelist()
def validate_and_update_mispunch(doc, method):
    if doc.workflow_state != "Approved":
        return
    if doc.reason == "Mispunch":
        formatted_date = formatdate(doc.from_date, "dd-MM-yyyy")
        attendance = frappe.db.get_value("Attendance",{"employee": doc.employee, "attendance_date": doc.from_date},["name", "in_time", "out_time", "status"])
        if not attendance:
            throw(f"No attendance record found for {formatted_date}.")
        att_doc = frappe.get_doc("Attendance", attendance[0])
        if att_doc.in_time or att_doc.out_time:
            att_doc.status = "Present"
            att_doc.attendance_request=doc.name
        att_doc.save(ignore_permissions=True)
        frappe.msgprint(f"Attendance updated successfully for {formatted_date}.")

@frappe.whitelist()
def update_att_oncancel_mispunch(doc,method):
    mark_att_request(doc.from_date, doc.to_date, doc.employee)

from teampro.mark_attendance import mark_attendance_from_checkin, get_dates
@frappe.whitelist()
def mark_att_request(from_date=None, to_date=None,employee=None):
    from_date = from_date
    to_date = to_date
    dates = get_dates(from_date,to_date)
    for date in dates:
        from_date = date
        to_date = date
        checkins = frappe.db.sql("""select * from `tabEmployee Checkin` where date(time) between '%s' and '%s' and employee = '%s' order by time """%(from_date,to_date,employee),as_dict=1)
        for c in checkins:
            employee = frappe.db.exists('Employee',{'status':'Active','date_of_joining':['<=',from_date],'name':c.employee})
            if employee:  
                mark_attendance_from_checkin(c.name,c.employee,c.time)


def get_allowed_time(slot):
    mapping = {
        "09 30 to 10 30": "10:35",
        "09 30 to 11 30": "11:35",
        "16 30 to 18 30": "16:25",
        "17 30 to 18 30": "17:25",
    }
    return mapping.get(slot)

import frappe
from frappe.utils import nowdate

def create_comp_off_requests():
    today = nowdate()
    attendances = frappe.get_all(
        "Attendance",
        filters={
            "attendance_date": today,
            "docstatus":("!=", 2),
        },
        fields=["name", "employee", "attendance_date", "status", "in_time", "out_time","bt_difference"]
    )

    for att in attendances:
        if not (att.in_time and att.out_time):
            continue

        holiday_list = frappe.db.get_value("Employee", att.employee, "holiday_list")
        if not holiday_list:
            continue

        is_holiday = frappe.db.exists("Holiday", {
            "parent": holiday_list,
            "holiday_date": att.attendance_date
        })

        if not is_holiday:
            continue

        if att.status not in ["Present", "Half Day"]:
            continue
        existing = frappe.db.exists("Attendance Request", {
            "employee": att.employee,
            "from_date": att.attendance_date,
            "to_date": att.attendance_date,
            "reason": "Comp Off_ Present Holiday"
        })
        if existing:
            continue
        doc = frappe.new_doc("Attendance Request")
        doc.employee = att.employee
        doc.from_date = att.attendance_date
        doc.to_date = att.attendance_date
        if att.status=="Half Day":
            doc.half_day =1
            doc.half_day_date = att.attendance_date
        doc.reason = "Comp Off_ Present Holiday"
        formatted_date = formatdate(att.attendance_date, "dd-MM-yyyy")
        doc.explanation=f"Auto-generated for {formatted_date} as it is a holiday and employee was marked {att.status}"
        doc.custom_attendance = att.name 
        doc.custom_in_time = att.in_time
        doc.custom_out_time = att.out_time
        doc.custom_bt_difference=att.bt_difference
        if att.status=="Present":
            doc.total_days=1
        elif att.status=="Half Day":
            doc.total_days=0.5
            doc.custom_session="Second Half" if att.in_time else "First Half"
        doc.insert(ignore_permissions=True)
        doc.submit()


import frappe
from frappe.utils import getdate, add_days, formatdate
from frappe.utils import add_months, add_days

@frappe.whitelist()
def on_submit_attendance_request(doc, method):
    if doc.workflow_state != "Approved":
        return
    if doc.reason in ["Mispunch", "Permission", "On Duty","On Duty Working Day"]:
        return

    if not doc.from_date or not doc.to_date:
        frappe.throw("From Date and To Date are required")

    total_days = doc.total_days

    period_start = getdate(doc.from_date)
    # period_end = add_days(period_start, 29) 
    period_end = add_days(add_months(period_start, 6), -1)
    overlapping_alloc = frappe.db.get_all(
        "Leave Allocation",
        filters={
            "employee": doc.employee,
            "leave_type": "Compensatory Off",
            "docstatus": 1,
            "from_date": ["<=", period_end],
            "to_date": [">=", period_start]
        },
        fields=["name", "from_date", "to_date", "new_leaves_allocated"],
        limit=1
    )

    if overlapping_alloc:
        leave_doc = frappe.get_doc("Leave Allocation", overlapping_alloc[0]["name"])
        leave_doc.new_leaves_allocated += total_days
        leave_doc.to_date = max(getdate(overlapping_alloc[0]["to_date"]), period_end)
        leave_doc.save(ignore_permissions=True)
        frappe.msgprint(f"Existing Comp Off Allocation updated for {doc.employee}")
    else:
        leave_alloc = frappe.new_doc("Leave Allocation")
        leave_alloc.employee = doc.employee
        leave_alloc.leave_type = "Compensatory Off"
        leave_alloc.from_date = period_start
        leave_alloc.to_date = period_end
        leave_alloc.new_leaves_allocated = total_days
        leave_alloc.company = doc.company
        leave_alloc.insert(ignore_permissions=True)
        leave_alloc.submit()
        frappe.msgprint(f"Comp Off Leave Allocation created for {doc.employee} for period {formatdate(period_start)} to {formatdate(period_end)}")

import frappe
from frappe.utils import getdate, add_days
from frappe.utils import add_months, add_days

@frappe.whitelist()
def on_cancel_attendance_request(doc, method):
    if doc.reason in ["Mispunch", "Permission", "On Duty","On Duty Working Day"]:
        return

    if not doc.from_date or not doc.to_date:
        return

    total_days = doc.total_days

    period_start = getdate(doc.from_date)
    # period_end = add_days(period_start, 29)
    period_end = add_days(add_months(period_start, 6), -1)
    overlapping_alloc = frappe.db.get_all(
        "Leave Allocation",
        filters={
            "employee": doc.employee,
            "leave_type": "Compensatory Off",
            "docstatus": 1,
            "from_date": ["<=", period_end],
            "to_date": [">=", period_start]
        },
        fields=["name", "new_leaves_allocated"],
        limit=1
    )

    if overlapping_alloc:
        leave_doc = frappe.get_doc("Leave Allocation", overlapping_alloc[0]["name"])
        leave_doc.new_leaves_allocated -= total_days
        if leave_doc.new_leaves_allocated < 0:
            leave_doc.new_leaves_allocated = 0

        leave_doc.save(ignore_permissions=True)

        frappe.msgprint(f"Comp Off Allocation reduced for {doc.employee}")