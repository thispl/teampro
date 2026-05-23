import frappe
from frappe.utils.csvutils import read_csv_content
from frappe.utils import get_first_day, get_last_day, format_datetime, get_url_to_form
from frappe.utils import cint
from frappe.utils.data import date_diff, now_datetime, nowdate, today, add_days
import datetime
from frappe import _
from frappe import _
from frappe.model.document import Document
from frappe.utils import getdate, nowdate
from frappe import throw, msgprint
import frappe
from datetime import date
from frappe import throw, _
from frappe.utils import getdate, today
today = date.today()
from frappe.model.document import Document
import datetime
import frappe,erpnext
from frappe.utils import cint
from frappe.utils import validate_email_address
import json
from frappe.utils import date_diff, add_months,today,add_days,add_years,nowdate,flt
from frappe.model.mapper import get_mapped_doc
from frappe.utils.file_manager import get_file
from frappe.utils.csvutils import UnicodeWriter, read_csv_content
import datetime
from datetime import date,datetime,timedelta
import openpyxl
from openpyxl import Workbook
import openpyxl
import xlrd
import re
from frappe.utils import today
from openpyxl import Workbook
from openpyxl.styles import PatternFill, Font, Alignment
from openpyxl.utils import get_column_letter
from datetime import datetime
from io import BytesIO
from openpyxl.styles import Font, Alignment, Border, Side
from openpyxl import load_workbook
from openpyxl.utils import get_column_letter
from openpyxl.styles import GradientFill, PatternFill
import pandas as pd
from frappe.utils import formatdate
from frappe.utils import now
from erpnext.setup.utils import get_exchange_rate
from datetime import date
from six import BytesIO, string_types
from frappe.utils import time_diff
from frappe.utils.csvutils import read_csv_content
from erpnext.buying.doctype.purchase_order.purchase_order import update_status
from frappe.utils.file_manager import get_file
from urllib.parse import urlencode
from frappe.model.rename_doc import rename_doc
from datetime import datetime
from frappe.utils import today

# @frappe.whitelist()
# def get_so_item_details(doc,method):
#     if doc.company=="TEAMPRO Food Products":
#         if not doc.sales_order:
#             return

#         for item in doc.items:
#             if not item.item_code:
#                 continue

#             so_item = frappe.db.get_value(
#                 "Sales Order Item",
#                 {
#                     "parent": item.against_sales_order or doc.sales_order,
#                     "item_code": item.item_code,
#                     "idx": item.idx
#                 },
#                 [
#                     "custom_cover_type",
#                     "custom_packing_type",
#                     "custom_tertiary_packingbox",
#                     "custom_bag",
#                     "custom_covers",
#                     "custom_box",
#                     "custom_per_2p",
#                     "custom_per_3p",
#                     "custom_wrd_uom",
#                     "custom_wrd_rate",
#                     "custom_mfg_on",
#                     "custom_name_print",
#                     "custom_wrd_item_name"
#                 ],
#                 as_dict=True
#             )

#             if so_item:
#                 item.custom_cover_type = so_item.custom_cover_type
#                 item.custom_packing_type = so_item.custom_packing_type
#                 item.custom_tertiary_packingbox = so_item.custom_tertiary_packingbox
#                 item.custom_bag = so_item.custom_bag
#                 item.custom_covers = so_item.custom_covers
#                 item.custom_box = so_item.custom_box
#                 item.custom_per_2p = so_item.custom_per_2p
#                 item.custom_per_3p = so_item.custom_per_3p
#                 item.custom_wrd_uom = so_item.custom_wrd_uom
#                 item.custom_wrd_rate = so_item.custom_wrd_rate
#                 item.custom_mfg_on=so_item.custom_mfg_on
#                 item.custom_name_print=so_item.custom_name_print
#                 item.custom_wrd_item_name=so_item.custom_wrd_item_name
#         doc.save(ignore_permissions=True)


@frappe.whitelist()
def get_so_item_details(doc, method):
    if doc.company != "TEAMPRO Food Products":
        return

    if not doc.sales_order:
        return
    if doc.sales_order:
        delivery_date=frappe.db.get_value("Sales Order",doc.sales_order,"delivery_date")
        if delivery_date:
            doc.custom_delivery_date=delivery_date
    for item in doc.items:

        
        if not item.so_detail:
            so_items = frappe.get_all(
                "Sales Order Item",
                filters={
                    "parent": doc.sales_order,
                    "item_code": item.item_code
                },
                fields=["name", "qty", "delivered_qty"]
            )

            
            pending = next((s for s in so_items if s.delivered_qty < s.qty), None)

            if pending:
                item.so_detail = pending.name

        
        if not item.so_detail:
            continue

        
        so_item = frappe.db.get_value(
            "Sales Order Item",
            item.so_detail,
            [
                "custom_cover_type",
                "custom_packing_type",
                "custom_tertiary_packingbox",
                "custom_bag",
                "custom_covers",
                "custom_box",
                "custom_per_2p",
                "custom_per_3p",
                "custom_wrd_uom",
                "custom_wrd_rate",
                "custom_mfg_on",
                "custom_name_print",
                "custom_wrd_item_name"
            ],
            as_dict=True
        )

        
        if so_item:
            for f, val in so_item.items():
                setattr(item, f, val)

    doc.save(ignore_permissions=True)



@frappe.whitelist()
def set_totals_in_delivery_note(doc, method):
    doc = frappe.get_doc("Delivery Note", doc.name) if isinstance(doc, str) else doc
    total_covers = 0
    total_bag = 0
    total_box = 0
    if doc.items:
        for item in doc.items:
            total_covers += flt(item.custom_covers)
            total_bag += flt(item.custom_bag)
            total_box += flt(item.custom_box)

        doc.custom_total_covers = total_covers
        doc.custom_total_bag = total_bag
        doc.custom_total_box = total_box

@frappe.whitelist()
def validate_packing_items_on_dn(doc, method):
    warehouse = 'Stores - TFP'
    if doc.company=="TEAMPRO Food Products":
        for idx, row in enumerate(doc.items, start=1):
            # Primary Packing
            if row.custom_cover_type and row.custom_covers:
                required_qty =row.custom_covers
                available_qty = get_available_balance(row.custom_cover_type, warehouse)
                conversion_factor = get_uom_conversion(row.custom_cover_type, row.custom_primary_uom)
                if conversion_factor:
                    required_qty = required_qty * conversion_factor
                else:
                    frappe.throw(f"UOM Conversion factor not found for Item: {row.custom_cover_type} and UOM: {row.custom_primary_uom}")
                if available_qty is not None and required_qty > available_qty:
                    throw(
                        _("Row #{0}: Item Code {1} is not available in warehouse {2}.")
                        .format(
                            idx,
                            frappe.bold(row.custom_cover_type),
                            frappe.bold(warehouse),
                        ),
                        title=_("Item Unavailable")
                    )


            # Secondary Packing
            if row.custom_packing_type and row.custom_bag:
                available_qty = get_available_balance(row.custom_packing_type, warehouse)
                conversion_factor = get_uom_conversion(row.custom_packing_type, row.custom_secondary_uom)
                required_qty =row.custom_bag
                if conversion_factor:
                    required_qty = required_qty * conversion_factor
                else:
                    frappe.throw(f"UOM Conversion factor not found for Item: {row.custom_packing_type} and UOM: {row.custom_secondary_uom}")
                if available_qty is not None and required_qty > available_qty:
                    throw(
                        _("Row #{0}: Item Code {1} is not available in warehouse {2}.")
                        .format(
                            idx,
                            frappe.bold(row.custom_packing_type),
                            frappe.bold(warehouse)),
                        title=_("Item Unavailable")
                    )

            # Tertiary Packing
            if row.custom_tertiary_packingbox and row.custom_box:
                required_qty =row.custom_box
                available_qty = get_available_balance(row.custom_tertiary_packingbox, warehouse)
                conversion_factor = get_uom_conversion(row.custom_tertiary_packingbox, row.custom_tertiary_uom)
                if conversion_factor:
                    required_qty = required_qty * conversion_factor
                else:
                    frappe.throw(f"UOM Conversion factor not found for Item: {row.custom_tertiary_packingbox} and UOM: {row.custom_tertiary_uom}")
                if available_qty is not None and required_qty > available_qty:
                    throw(
                        _("Row #{0}: Item Code {1} is not available in warehouse {2}.")
                        .format(
                            idx,
                            frappe.bold(row.custom_tertiary_packingbox),
                            frappe.bold(warehouse)
                        ),
                        title=_("Item Unavailable")
                    )

@frappe.whitelist()
def get_available_balance(item_code,warehouse):
    bin_record = frappe.get_list('Bin',
        filters={
            'item_code': item_code,
            'warehouse': warehouse
        },
        fields=['actual_qty'],
        order_by='creation desc',
        limit=1
    )

    return bin_record[0].actual_qty if bin_record else 0

@frappe.whitelist()
def get_uom_conversion(item_code, uom):
    conversion = frappe.db.get_value('UOM Conversion Detail', {
        'parent': item_code,
        'uom': uom
    }, 'conversion_factor')
    return conversion

@frappe.whitelist()
def update_so_priority(doc,method):
    dn_list = frappe.get_all(
        "Delivery Note",
        filters={
            "company": "TEAMPRO Food Products",
            "status": "To Bill",
            "docstatus":1,
            "custom_delivery_status_new":"Schedule"
        },
        fields=["name", "custom_packing_on","custom_priority"]
    )

    from collections import defaultdict
    date_groups = defaultdict(list)

    for so in dn_list:
        if so.custom_packing_on:
            date_groups[so.custom_packing_on].append(so.name)
    for priority, packing_date in enumerate(sorted(date_groups.keys()), start=1):
        for so_name in date_groups[packing_date]:
            frappe.db.set_value("Delivery Note", so_name, "custom_priority", priority)

import frappe
from frappe.model.document import Document
from frappe.utils import flt

@frappe.whitelist()
def create_material_issue(doc, method):
    doc = frappe.get_doc("Delivery Note", doc.name) if isinstance(doc, str) else doc

    if doc.company != "TEAMPRO Food Products":
        return

    stock_entry = None
    for item in doc.items:
        items_to_issue = []

        if flt(item.custom_covers) > 0 and item.custom_cover_type:
            items_to_issue.append({
                "item_code": item.custom_cover_type,
                "qty": item.custom_covers
            })

        if flt(item.custom_bag) > 0 and item.custom_packing_type:
            items_to_issue.append({
                "item_code": item.custom_packing_type,
                "qty": item.custom_bag
            })

        if flt(item.custom_box) > 0 and item.custom_tertiary_packingbox:
            items_to_issue.append({
                "item_code": item.custom_tertiary_packingbox,
                "qty": item.custom_box
            })

        if items_to_issue:
            if not stock_entry:
                stock_entry = frappe.new_doc("Stock Entry")
                stock_entry.stock_entry_type = "Material Issue"
                stock_entry.from_warehouse = "Stores - TFP"
                stock_entry.company = doc.company
                # stock_entry.set_posting_time = 1
                stock_entry.posting_date = doc.posting_date
                stock_entry.posting_time = doc.posting_time
                stock_entry.custom_delivery_note=doc.name
            for i in items_to_issue:
                stock_entry.append("items", {
                    "item_code": i["item_code"],
                    "qty": i["qty"],
                    "uom":"Nos",
                    "s_warehouse": stock_entry.from_warehouse,
                    "cost_center": item.cost_center or frappe.db.get_value("Company", doc.company, "cost_center"),
                    "allow_zero_valuation_rate":1
                })

    if stock_entry:
        stock_entry.insert()
        stock_entry.submit()
        frappe.msgprint(f"Material Issue created: {stock_entry.name}")