import frappe
from frappe.utils import flt, getdate, cstr, get_first_day, get_last_day
from frappe.utils.pdf import get_pdf
from frappe import _

@frappe.whitelist()
def get_party_statement(party_type, party, from_date=None, to_date=None, company=None):
    """Return statement data for a party (Customer/Supplier) with Payment Entry settlements."""
    if not company:
        company = frappe.defaults.get_user_default("Company")
    if not company:
        frappe.throw(_("Company is required"))

    if not from_date:
        # Default to current fiscal year start
        fy = frappe.db.get_value("Fiscal Year", {"disabled": 0}, "name", order_by="year_start_date desc")
        if fy:
            from_date = frappe.db.get_value("Fiscal Year", fy, "year_start_date")
        else:
            from_date = get_first_day(getdate())
    if not to_date:
        to_date = getdate()

    # Resolve party account
    party_account = frappe.db.get_value(
        "Party Account",
        {"parent": party, "parenttype": party_type, "company": company},
        "account"
    )
    if not party_account:
        party_account = frappe.db.get_value(
            "GL Entry",
            {"party_type": party_type, "party": party, "company": company, "is_cancelled": 0},
            "account"
        )
    if not party_account:
        if party_type == "Customer":
            party_account = frappe.db.get_value("Company", company, "default_receivable_account")
        elif party_type == "Supplier":
            party_account = frappe.db.get_value("Company", company, "default_payable_account")

    # Get currency
    currency = frappe.db.get_value("Company", company, "default_currency") or "INR"

    data = []
    opening = 0
    total_debit = 0
    total_credit = 0
    closing_balance = 0

    if party_account:
        # Opening balance
        opening = flt(frappe.db.sql("""
            SELECT SUM(debit - credit)
            FROM `tabGL Entry`
            WHERE company = %s AND account = %s AND party = %s
              AND is_cancelled = 0 AND posting_date < %s
        """, (company, party_account, party, from_date))[0][0])

        # GL entries grouped by voucher
        conditions = "company = %s AND account = %s AND party = %s AND is_cancelled = 0"
        params = [company, party_account, party]
        conditions += " AND posting_date >= %s"
        params.append(from_date)
        conditions += " AND posting_date <= %s"
        params.append(to_date)

        rows = frappe.db.sql("""
            SELECT
                posting_date,
                voucher_type,
                voucher_no,
                MAX(remarks) AS remarks,
                SUM(debit) AS debit,
                SUM(credit) AS credit,
                GROUP_CONCAT(DISTINCT against_voucher) AS against_vouchers
            FROM `tabGL Entry`
            WHERE """ + conditions + """
            GROUP BY voucher_type, voucher_no, posting_date
            ORDER BY posting_date, voucher_no
        """, params, as_dict=True)

        # Batch fetch Payment Entry references
        pe_names = [r.voucher_no for r in rows if r.voucher_type == "Payment Entry"]
        pe_refs = {}
        if pe_names:
            refs = frappe.db.get_all(
                "Payment Entry Reference",
                filters={"parent": ["in", pe_names], "parenttype": "Payment Entry"},
                fields=["parent", "reference_doctype", "reference_name", "allocated_amount",
                        "total_amount", "outstanding_amount", "due_date", "bill_no"],
                order_by="parent, idx",
            )
            for ref in refs:
                pe_refs.setdefault(ref.parent, []).append(ref)

        # Batch fetch bill_no and outstanding for Purchase Invoices (supplier bills)
        pi_names = [r.voucher_no for r in rows if r.voucher_type == "Purchase Invoice"]
        pi_bills = {}
        if pi_names:
            bill_data = frappe.db.get_all(
                "Purchase Invoice",
                filters={"name": ["in", pi_names]},
                fields=["name", "bill_no", "bill_date", "outstanding_amount", "docstatus"],
            )
            for b in bill_data:
                pi_bills[b.name] = b

        # Batch fetch outstanding for Sales Invoices
        si_names = [r.voucher_no for r in rows if r.voucher_type == "Sales Invoice"]
        si_outstanding = {}
        if si_names:
            si_data = frappe.db.get_all(
                "Sales Invoice",
                filters={"name": ["in", si_names]},
                fields=["name", "outstanding_amount", "docstatus"],
            )
            for s in si_data:
                si_outstanding[s.name] = s

        # Opening balance row
        if opening != 0:
            data.append({
                "posting_date": from_date,
                "voucher_type": "",
                "voucher_no": "",
                "remarks": "Opening Balance",
                "debit": opening if opening > 0 else 0,
                "credit": abs(opening) if opening < 0 else 0,
                "balance": opening,
                "settled_against": "",
                "is_opening": 1,
                "settlements": [],
            })

        running = opening
        for r in rows:
            debit = flt(r.debit)
            credit = flt(r.credit)
            running += debit - credit
            total_debit += debit
            total_credit += credit

            settled_str = ""
            settlements = []
            if r.voucher_type == "Payment Entry" and r.voucher_no in pe_refs:
                parts = []
                for ref in pe_refs[r.voucher_no]:
                    label = ref.reference_name or ""
                    if ref.bill_no:
                        label = ref.bill_no + " (" + ref.reference_name + ")"
                    amt = flt(ref.allocated_amount)
                    parts.append(label + " [" + str(amt) + "]")
                    settlements.append({
                        "reference_doctype": ref.reference_doctype,
                        "reference_name": ref.reference_name,
                        "bill_no": ref.bill_no or "",
                        "allocated_amount": flt(ref.allocated_amount),
                        "total_amount": flt(ref.total_amount),
                        "outstanding_amount": flt(ref.outstanding_amount),
                        "due_date": cstr(ref.due_date) if ref.due_date else "",
                    })
                settled_str = ", ".join(parts)

            # Get bill_no for Purchase Invoices
            bill_no = ""
            outstanding = 0
            is_unsettled = 0
            if r.voucher_type == "Purchase Invoice" and r.voucher_no in pi_bills:
                bill_no = pi_bills[r.voucher_no].get("bill_no") or ""
                outstanding = flt(pi_bills[r.voucher_no].get("outstanding_amount"))
                if outstanding > 0:
                    is_unsettled = 1
            elif r.voucher_type == "Sales Invoice" and r.voucher_no in si_outstanding:
                outstanding = flt(si_outstanding[r.voucher_no].get("outstanding_amount"))
                if outstanding > 0:
                    is_unsettled = 1

            data.append({
                "posting_date": cstr(r.posting_date),
                "voucher_type": r.voucher_type,
                "voucher_no": r.voucher_no,
                "bill_no": bill_no,
                "outstanding": outstanding,
                "is_unsettled": is_unsettled,
                "remarks": r.remarks or "",
                "debit": debit,
                "credit": credit,
                "balance": running,
                "settled_against": settled_str,
                "settlements": settlements,
            })

        closing_balance = running

    # For Supplier statements, show from the supplier's perspective:
    # Supplier's invoices (our purchases) = Debit (they billed us)
    # Our payments to supplier = Credit (they received)
    # Positive balance = company owes supplier
    if party_type == "Supplier":
        opening = -opening
        total_debit, total_credit = total_credit, total_debit
        closing_balance = -closing_balance
        for row in data:
            row["debit"], row["credit"] = row.get("credit", 0), row.get("debit", 0)
            row["balance"] = -row.get("balance", 0)

    # Build HTML
    html = _render_statement_html(
        party=party,
        party_type=party_type,
        company=company,
        from_date=from_date,
        to_date=to_date,
        party_account=party_account or "",
        currency=currency,
        opening=opening,
        total_debit=total_debit,
        total_credit=total_credit,
        closing_balance=closing_balance,
        data=data,
    )

    return {"html": html, "data": data}


def _render_statement_html(party, party_type, company, from_date, to_date,
                           party_account, currency, opening, total_debit,
                           total_credit, closing_balance, data):
    """Render the statement as presentable HTML for print/PDF."""

    def fmt(amount):
        if amount is None:
            return ""
        a = flt(amount)
        if a == 0:
            return ""
        # Format with thousands separator
        s = "{:,.2f}".format(a)
        return s

    def fmt0(amount):
        """Always show amount, even if 0."""
        if amount is None:
            amount = 0
        return "{:,.2f}".format(flt(amount))

    def fmt_date(date_str):
        if not date_str:
            return ""
        try:
            parts = cstr(date_str).split("-")
            if len(parts) == 3:
                return "{}-{}-{}".format(parts[2], parts[1], parts[0][-2:])
        except Exception:
            pass
        return cstr(date_str)

    rows_html = ""
    for row in data:
        cls = ""
        if row.get("is_opening"):
            cls = ' class="opening"'
        if row.get("is_total"):
            cls = ' class="total"'
        if row.get("is_unsettled"):
            cls = ' class="unsettled"'

        rows_html += "<tr{}>".format(cls)
        rows_html += "<td>{}</td>".format(fmt_date(row.get("posting_date", "")))
        # Reference column: show voucher type, voucher no, and bill_no (for suppliers)
        ref_html = "<div class='vtype'>{}</div>".format(row.get("voucher_type", ""))
        ref_html += "<div class='vref'>{}</div>".format(row.get("voucher_no", ""))
        bill_no = row.get("bill_no", "")
        if bill_no:
            ref_html += "<div class='billno'>Bill No: {}</div>".format(bill_no)
        if row.get("is_unsettled"):
            outstanding = flt(row.get("outstanding", 0))
            ref_html += "<div class='unsettled-tag'>UNSETTLED - Outstanding: {}</div>".format(fmt0(outstanding))
        rows_html += "<td>{}</td>".format(ref_html)
        debit_val = fmt(row.get("debit"))
        credit_val = fmt(row.get("credit"))
        rows_html += "<td class='right debit-color'>{}</td>".format(debit_val)
        rows_html += "<td class='right credit-color'>{}</td>".format(credit_val)
        rows_html += "<td class='right'>{}</td>".format(fmt(row.get("balance")))
        rows_html += "</tr>"

        # Detail sub-rows for settlements
        settlements = row.get("settlements") or []
        for i, st in enumerate(settlements):
            detail_cls = "detail"
            if i == 0:
                detail_cls += " detail-first"
            if i == len(settlements) - 1:
                detail_cls += " detail-last"
            rows_html += "<tr class='{}'>".format(detail_cls)
            rows_html += "<td colspan='5'>&rdsh;&nbsp; {}: {}".format(
                st.get("reference_doctype", ""), st.get("reference_name", ""))
            if st.get("bill_no"):
                rows_html += " <span style='color:#999;'>(Bill: {})</span>".format(st["bill_no"])
            rows_html += " &nbsp;|&nbsp; Allocated: {}".format(fmt(st.get("allocated_amount")))
            if st.get("total_amount") and flt(st["total_amount"]) != flt(st.get("allocated_amount")):
                rows_html += " &nbsp;|&nbsp; Invoice Total: {}".format(fmt(st["total_amount"]))
            rows_html += "</td></tr>"

    # Total row
    rows_html += "<tr class='total'>"
    rows_html += "<td>{}</td>".format(fmt_date(to_date))
    rows_html += "<td>TOTAL</td>"
    rows_html += "<td class='right'>{}</td>".format(fmt(total_debit))
    rows_html += "<td class='right'>{}</td>".format(fmt(total_credit))
    rows_html += "<td class='right'>{}</td>".format(fmt(closing_balance))
    rows_html += "</tr>"

    html = """<!DOCTYPE html>
<html>
<head>
<meta charset="utf-8">
<style>
  * {{ box-sizing: border-box; }}
  body {{ font-family: Helvetica, Arial, sans-serif; font-size: 12px; color: #333; margin: 0; padding: 20px; }}
  .header {{ text-align: center; margin-bottom: 18px; padding: 20px 15px 16px 15px; background: #1a5276; border: 3px solid #1a5276; color: #fff; }}
  .header .party-name {{ font-size: 24px; font-weight: 700; color: #fff; margin: 0 0 4px 0; }}
  .header .title {{ font-size: 14px; color: #fff; text-transform: uppercase; letter-spacing: 3px; margin: 0 0 6px 0; opacity: 0.9; }}
  .header .meta {{ font-size: 11px; color: #aed6f1; margin: 0; }}
  .summary-box {{ margin: 18px 0; border: 1px solid #d5dbdb; }}
  .summary-box table {{ width: 100%; border-collapse: collapse; }}
  .summary-box td {{ padding: 8px 12px; font-size: 12px; background: #fff; }}
  .summary-box .label {{ color: #5d6d7e; font-weight: 600; }}
  .summary-box .value {{ color: #1a5276; font-weight: 700; text-align: right; }}
  .summary-box .row1 td {{ background: #ebf5fb; }}
  .summary-box .row2 td {{ background: #f4f6f7; }}
  .summary-box .closing td {{ background: #1a5276; color: #fff; font-size: 15px; padding: 10px 12px; border-top: 2px solid #1a5276; }}
  .summary-box .closing .label {{ color: #fff; }}
  .summary-box .closing .value {{ color: #fff; font-size: 15px; }}
  table.statement {{ width: 100%; border-collapse: collapse; margin-top: 12px; }}
  table.statement thead th {{ background: #1a5276; color: #fff; padding: 9px 8px; font-size: 11px; text-align: left; text-transform: uppercase; letter-spacing: 1px; border: 1px solid #1a5276; }}
  table.statement thead th.right {{ text-align: right; }}
  table.statement tbody td {{ padding: 6px 8px; border: 1px solid #d5dbdb; vertical-align: top; }}
  table.statement tbody td.right {{ text-align: right; white-space: nowrap; }}
  table.statement tbody tr:nth-child(even) td {{ background: #f8f9fb; }}
  table.statement tbody tr:nth-child(odd) td {{ background: #fff; }}
  table.statement tbody tr.opening td {{ font-style: italic; background: #fef9e7 !important; color: #7d6608; font-weight: 600; }}
  table.statement tbody tr.unsettled td {{ background: #fdedec !important; border-left: 4px solid #c0392b; }}
  table.statement tbody tr.total td {{ font-weight: 700; background: #eaf2f8 !important; border-top: 2px solid #1a5276; font-size: 13px; color: #1a5276; }}
  table.statement tbody tr.detail td {{ background: #f0f3f4 !important; font-size: 10px; color: #5d6d7e; border-left: 3px solid #2980b9; border-bottom: none; padding: 3px 8px 3px 24px; }}
  table.statement tbody tr.detail-first td {{ border-top: 1px dashed #bdc3c7; }}
  table.statement tbody tr.detail-last td {{ border-bottom: 1px solid #d5dbdb; }}
  .vtype {{ font-size: 10px; color: #2980b9; font-weight: 600; }}
  .vref {{ font-weight: 700; color: #1a5276; }}
  .billno {{ font-size: 10px; color: #e67e22; font-weight: 600; margin-top: 1px; }}
  .unsettled-tag {{ font-size: 10px; color: #c0392b; font-weight: 700; margin-top: 2px; padding: 1px 4px; background: #fadbd8; display: inline-block; }}
  .debit-color {{ color: #c0392b; }}
  .credit-color {{ color: #27ae60; }}
  .footer {{ margin-top: 18px; text-align: right; color: #95a5a6; font-size: 10px; padding-top: 8px; border-top: 1px solid #eee; }}
</style>
</head>
<body>
  <div class="header">
    <p class="party-name">{party}</p>
    <p class="title">Statement of Account</p>
    <p class="meta">{party_type} &nbsp;&bull;&nbsp; {company} &nbsp;&bull;&nbsp; {from_date} to {to_date}</p>
  </div>
  <div class="summary-box">
    <table>
      <tr class="row1">
        <td class="label">Opening Balance</td>
        <td class="value">{opening}</td>
        <td class="label">Total Debit</td>
        <td class="value">{total_debit}</td>
      </tr>
      <tr class="row2">
        <td class="label">Total Credit</td>
        <td class="value">{total_credit}</td>
        <td class="label">Account</td>
        <td class="value" style="text-align:right;">{party_account}</td>
      </tr>
      <tr class="closing">
        <td class="label" colspan="3">Closing Balance</td>
        <td class="value">{closing_balance}</td>
      </tr>
    </table>
  </div>
  <table class="statement">
    <thead>
      <tr>
        <th style="width:12%">Date</th>
        <th style="width:30%">Reference</th>
        <th class="right" style="width:18%">Debit</th>
        <th class="right" style="width:18%">Credit</th>
        <th class="right" style="width:22%">Balance</th>
      </tr>
    </thead>
    <tbody>
      {rows}
    </tbody>
  </table>
  <div class="footer">Printed On {print_date}</div>
</body>
</html>""".format(
        party=party,
        party_type=party_type,
        company=company,
        from_date=fmt_date(from_date),
        to_date=fmt_date(to_date),
        opening=fmt0(opening),
        total_debit=fmt(total_debit),
        total_credit=fmt(total_credit),
        party_account=party_account,
        closing_balance=fmt0(closing_balance),
        rows=rows_html,
        print_date=fmt_date(cstr(getdate())),
    )

    return html


@frappe.whitelist()
def download_party_statement_pdf(party_type, party, from_date=None, to_date=None, company=None):
    """Generate and return a PDF of the party statement."""
    result = get_party_statement(party_type, party, from_date, to_date, company)
    options = {
        "page-size": "A4",
        "margin-top": "15mm",
        "margin-bottom": "20mm",
        "margin-left": "12mm",
        "margin-right": "12mm",
        "encoding": "UTF-8",
        "enable-local-file-access": "",
        "print-media-type": "",
        "footer-center": "Page [page] of [topage]",
        "footer-font-size": "8",
        "footer-spacing": "5",
    }
    pdf = get_pdf(result["html"], options=options)
    filename = "Statement_{}_{}.pdf".format(
        party.replace(" ", "_").replace("/", "_").replace("&", "and"),
        cstr(getdate())
    )
    frappe.local.response.filename = filename
    frappe.local.response.filecontent = pdf
    frappe.local.response.type = "download"


@frappe.whitelist()
def download_party_statement_excel(party_type, party, from_date=None, to_date=None, company=None):
    """Generate and return an Excel file of the party statement."""
    result = get_party_statement(party_type, party, from_date, to_date, company)
    data = result["data"]

    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
    from io import BytesIO

    wb = Workbook()
    ws = wb.active
    ws.title = "Statement"

    # Styles
    header_font = Font(name="Helvetica", size=14, bold=True, color="FFFFFF")
    header_fill = PatternFill(start_color="1A5276", end_color="1A5276", fill_type="solid")
    col_header_font = Font(name="Helvetica", size=11, bold=True, color="FFFFFF")
    col_header_fill = PatternFill(start_color="1A5276", end_color="1A5276", fill_type="solid")
    label_font = Font(name="Helvetica", size=11, bold=True, color="5D6D7E")
    value_font = Font(name="Helvetica", size=11, bold=True, color="1A5276")
    total_font = Font(name="Helvetica", size=11, bold=True, color="1A5276")
    total_fill = PatternFill(start_color="EAF2F8", end_color="EAF2F8", fill_type="solid")
    opening_fill = PatternFill(start_color="FEF9E7", end_color="FEF9E7", fill_type="solid")
    normal_font = Font(name="Helvetica", size=10)
    thin_border = Border(
        left=Side(style="thin", color="D5DBDB"),
        right=Side(style="thin", color="D5DBDB"),
        top=Side(style="thin", color="D5DBDB"),
        bottom=Side(style="thin", color="D5DBDB"),
    )

    def fmt_amt(amount):
        if amount is None:
            return ""
        a = flt(amount)
        if a == 0:
            return ""
        return a

    def fmt_date(d):
        try:
            parts = cstr(d).split("-")
            if len(parts) == 3:
                return "{}-{}-{}".format(parts[2], parts[1], parts[0][-2:])
        except Exception:
            pass
        return cstr(d)

    # Title row
    ws.merge_cells("A1:E1")
    cell = ws["A1"]
    cell.value = party
    cell.font = Font(name="Helvetica", size=16, bold=True, color="1A5276")
    cell.alignment = Alignment(horizontal="center")

    ws.merge_cells("A2:E2")
    cell = ws["A2"]
    cell.value = "Statement of Account"
    cell.font = Font(name="Helvetica", size=12, bold=True, color="2980B9")
    cell.alignment = Alignment(horizontal="center")

    ws.merge_cells("A3:E3")
    cell = ws["A3"]
    cell.value = "{}  |  {}  |  {} to {}".format(party_type, company, fmt_date(from_date), fmt_date(to_date))
    cell.font = Font(name="Helvetica", size=10, color="7F8C8D")
    cell.alignment = Alignment(horizontal="center")

    # Summary section
    row_idx = 5
    # Compute summary from data (already flipped for suppliers)
    opening_val = 0
    t_debit = 0
    t_credit = 0
    closing_val = 0
    for row in data:
        if row.get("is_opening"):
            opening_val = flt(row.get("debit", 0)) - flt(row.get("credit", 0))
        elif not row.get("is_total"):
            t_debit += flt(row.get("debit", 0))
            t_credit += flt(row.get("credit", 0))
    closing_val = opening_val + t_debit - t_credit

    summary = [
        ("Opening Balance", opening_val, "Total Debit", t_debit),
        ("Total Credit", t_credit, "Account", result.get("party_account", "")),
    ]
    for srow in summary:
        ws.cell(row=row_idx, column=1, value=srow[0]).font = label_font
        ws.cell(row=row_idx, column=2, value=srow[1]).font = value_font
        ws.cell(row=row_idx, column=2).alignment = Alignment(horizontal="right")
        ws.cell(row=row_idx, column=3, value=srow[2]).font = label_font
        ws.cell(row=row_idx, column=4, value=srow[3]).font = value_font
        ws.cell(row=row_idx, column=4).alignment = Alignment(horizontal="right")
        row_idx += 1

    # Closing balance row
    ws.cell(row=row_idx, column=1, value="Closing Balance").font = Font(name="Helvetica", size=12, bold=True, color="FFFFFF")
    ws.cell(row=row_idx, column=1).fill = header_fill
    ws.merge_cells(start_row=row_idx, start_column=1, end_row=row_idx, end_column=3)
    ws.cell(row=row_idx, column=4, value=closing_val).font = Font(name="Helvetica", size=12, bold=True, color="FFFFFF")
    ws.cell(row=row_idx, column=4).fill = header_fill
    ws.cell(row=row_idx, column=4).alignment = Alignment(horizontal="right")
    row_idx += 2

    # Column headers
    headers = ["Date", "Reference", "Debit", "Credit", "Balance"]
    for col, h in enumerate(headers, 1):
        cell = ws.cell(row=row_idx, column=col, value=h)
        cell.font = col_header_font
        cell.fill = col_header_fill
        cell.alignment = Alignment(horizontal="right" if col >= 3 else "left")
        cell.border = thin_border
    row_idx += 1

    # Data rows
    for row in data:
        ref = row.get("voucher_type", "") + " " + row.get("voucher_no", "")
        if row.get("bill_no"):
            ref += " (Bill: {})".format(row["bill_no"])
        if row.get("is_opening"):
            ref = "Opening Balance"
        if row.get("is_unsettled"):
            ref += " [UNSETTLED - Outstanding: {}]".format(fmt0(row.get("outstanding", 0)))

        values = [
            fmt_date(row.get("posting_date", "")),
            ref,
            fmt_amt(row.get("debit")),
            fmt_amt(row.get("credit")),
            fmt_amt(row.get("balance")),
        ]
        unsettled_fill = PatternFill(start_color="FDEDEC", end_color="FDEDEC", fill_type="solid")
        for col, val in enumerate(values, 1):
            cell = ws.cell(row=row_idx, column=col, value=val)
            cell.font = normal_font
            cell.border = thin_border
            if col >= 3:
                cell.alignment = Alignment(horizontal="right")
            if row.get("is_opening"):
                cell.fill = opening_fill
                cell.font = Font(name="Helvetica", size=10, italic=True, bold=True, color="7D6608")
            if row.get("is_total"):
                cell.fill = total_fill
                cell.font = total_font
            if row.get("is_unsettled"):
                cell.fill = unsettled_fill
                cell.font = Font(name="Helvetica", size=10, bold=True, color="C0392B")
        row_idx += 1

        # Detail sub-rows
        for st in (row.get("settlements") or []):
            detail_text = "  -> {}: {} | Allocated: {}".format(
                st.get("reference_doctype", ""), st.get("reference_name", ""),
                fmt_amt(st.get("allocated_amount")))
            if st.get("bill_no"):
                detail_text += " (Bill: {})".format(st["bill_no"])
            cell = ws.cell(row=row_idx, column=1, value=detail_text)
            ws.merge_cells(start_row=row_idx, start_column=1, end_row=row_idx, end_column=5)
            cell.font = Font(name="Helvetica", size=9, color="5D6D7E")
            row_idx += 1

    # Column widths
    ws.column_dimensions["A"].width = 14
    ws.column_dimensions["B"].width = 40
    ws.column_dimensions["C"].width = 15
    ws.column_dimensions["D"].width = 15
    ws.column_dimensions["E"].width = 15

    # Write to BytesIO
    output = BytesIO()
    wb.save(output)
    output.seek(0)

    filename = "Statement_{}_{}.xlsx".format(
        party.replace(" ", "_").replace("/", "_").replace("&", "and"),
        cstr(getdate())
    )
    frappe.local.response.filename = filename
    frappe.local.response.filecontent = output.getvalue()
    frappe.local.response.type = "download"
